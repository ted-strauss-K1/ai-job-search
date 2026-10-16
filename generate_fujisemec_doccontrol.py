"""
Generate application docs for Fuji SEMEC — Documentation Control & Technical
Publications Specialist. Boucherville, QC. Hybrid. Full-time, $70k-$100k.
Outputs (in docx/):
  - cv_fujisemec_doccontrol.docx     / .pdf
  - cover_fujisemec_doccontrol.docx  / .pdf

Content is defined ONCE below, then rendered to both docx (python-docx) and
pdf (reportlab) so the two formats never drift.

Targeting note: real match on controlled-document practice (numbering, naming,
revision and approval workflow, retention/classification/access policy,
traceability), on technical writing at volume, and on the training and
continuous-improvement half of the role. Genuine gaps, stated plainly and never
hidden: no ISO 9001 / EN standards work, no CMMI, no experience reading
engineering drawings, and no PLM or authoring-tool experience (Windchill,
Teamcenter, DOORS, FrameMaker, MadCap Flare).

Company facts used in the cover letter were verified independently at
fujisemec.com (not from the posting text): door opening and closing devices for
railcars, light rail and transit; Fuji Electric acquired SEMEC Electromecanique;
sites in Boucherville QC, Plattsburgh NY and Edison NJ.
"""
import os

os.makedirs("docx", exist_ok=True)

# ------------------------------------------------------------------
# SHARED CONTENT
# ------------------------------------------------------------------

NAME = "Ted Strauss"
CV_CONTACT = "Montreal, Quebec, Canada  |  514-475-2374  |  ted.strauss@gmail.com  |  linkedin.com/in/tedstrauss-63b01821"
COVER_CONTACT = "514-475-2374  |  ted.strauss@gmail.com  |  linkedin.com/in/tedstrauss-63b01821"

PROFILE = (
    "Documentation and information governance specialist with over a decade of experience defining how "
    "technical documents are structured, numbered, revised, approved, and found again. Led data "
    "governance at McGill University's McConnell Brain Imaging Centre, where I wrote the metadata "
    "standards, naming conventions, and classification, retention, and access policies that dozens of "
    "research groups worked to, and maintained the institutional catalogue that held the approved "
    "versions. Author of nine peer-reviewed publications, a published documentation template, and a "
    "public knowledge base, with a record of turning specialist material into documents non-specialists "
    "can actually use. Bilingual (EN/FR), Montreal-based."
)

# (bold label, remaining text) — label is emphasised
COMPETENCIES = [
    ("Document control and configuration",
     ": established and enforced controlled-document practice covering numbering and naming conventions, "
     "templates, revision history, review and approval workflow, distribution, and archiving; wrote the "
     "policies governing classification, retention, and access, and maintained the release records showing "
     "who received which version"),
    ("Technical writing and technical publications",
     ": authored manuals, templates, specifications, and training guides for mixed technical and "
     "non-technical readers; nine peer-reviewed publications; a data management plan template published on "
     "Zenodo and reused as a controlled template across an academic field; a public knowledge base written "
     "from scratch"),
    ("Standards, traceability, and audit support",
     ": maintained documentation standards, metadata standards, data definitions, and quality expectations so that documents stayed "
     "accurate, complete, consistent, and traceable; maintained regulatory compliance obligations and the "
     "supporting approval trail on government learning-platform deployments"),
    ("Document management systems",
     ": Collibra as the institutional catalogue, maintaining the business glossary, metadata fields, and "
     "naming conventions; OwnCloud document portal; SharePoint and MS Teams; Drupal and WordPress; Git for "
     "version control; Adobe suite for publication-quality output"),
    ("Subject-matter expert (SME) collaboration",
     ": gathered, clarified, and validated technical content from principal investigators, engineers, and "
     "developers who did not report to me; coordinated documentation between government clients and "
     "development teams; briefed institutional leadership and national committees"),
    ("Process improvement and training",
     ": replaced manual, email-based document requests with a self-serve portal; built a Moodle-based "
     "learning portal and ran hands-on training so that standards were adopted rather than merely issued; "
     "certified in Scientific Community Engagement Fundamentals (CSCCE); evaluate and adopt new tooling, "
     "including AI tools such as Claude Code, for searchability and drafting workflows"),
]

# each: (dates, title, org, location, [bullets])
EXPERIENCE = [
    ("09/2016 – 10/2021", "Manager of Data Resources", "McConnell Brain Imaging Centre, The Neuro, McGill University", "Montreal, QC", [
        "Owned the controlled-document side of a research data function: metadata standards, data definitions, naming and numbering conventions, ownership structures, and the classification, retention, and access policies that governed them, applied across dozens of principal investigators and hundreds of researchers",
        "Maintained Collibra as the institutional catalogue, including the business glossary, metadata fields, and naming conventions, so that every definition had one approved source instead of competing copies circulating by email",
        "Wrote and published a data management plan template (Zenodo, 2021) that was adopted as a reusable controlled template by researchers across the neurosciences",
        "Replaced manual, email-based data-request handling with a self-serve OwnCloud portal, shortening turnaround and producing a consistent record of what was released, to whom, and under which version",
        "Represented the centre on national standards bodies (Portage, CONP, RDA), contributing to shared metadata and documentation standards used across Canadian research institutions",
    ]),
    ("07/2025 – 01/2026", "Moodle Analyst", "Premergency Inc.", "Canada", [
        "Coordinated between government clients and development teams on regulated learning platforms, maintaining the compliance obligations and approval trail those deployments required",
        "Translated client requirements into specifications developers could build from, and translated the result back into language non-technical stakeholders could review and sign off on",
    ]),
    ("08/2020 – 10/2021", "Data Management & Sharing Coordinator (p/t)", "Tanenbaum Open Science Institute, The Neuro, McGill University", "Montreal, QC", [
        "Wrote and maintained a public knowledge base explaining data management and open science concepts to non-specialists, and designed and delivered the training programme that went with it",
    ]),
    ("04/2017 – 08/2021", "Founder and Director", "Sur Place Media", "Montreal, QC", [
        "Founded and ran an arts school, recruiting and supporting 20 instructors across 150 workshop groups; led a full pivot to virtual delivery during the pandemic",
        "Owned the operating and records side end to end: budgeting, invoicing, expense tracking, purchasing, scheduling, and retention of business records",
    ]),
    ("01/2024 – 10/2024", "Full Stack Developer", "KnowledgeOne", "Montreal, QC", [
        "Developed learning platform features for active clients (LAMP++, Moodle, LXD), working from written specifications and documenting the delivered configuration for client handover",
    ]),
    ("02/2015 – 09/2016", "Software Developer", "McGill Centre for Integrative Neuroscience, McGill University", "Montreal, QC", [
        "Programmed the data acquisition pipeline for the national CCNA longitudinal dementia study on the LORIS platform (PHP, MySQL, JavaScript), a multi-site study where every site collected to the same documented protocol",
    ]),
]

ADDITIONAL = [
    "Photonic Knowledge — Database administrator and software developer (2012–2014): built a Drupal ERP system and its underlying data model for the mining exploration sector, including the document and record structures behind it",
    "Moss Rehabilitation Hospital — Research software developer (2007–2010)",
    "University of Connecticut & Columbia University — Research assistant and software developer (2004–2008)",
]

EDUCATION = [
    ("1998–2003", "B.A. (Honours), Cognitive Science", "McGill University", "Montreal, QC",
     ["Minor: Music Technology. Topics: psycholinguistics, cognitive science, computational modeling of language, music technology."]),
    ("2021", "Scientific Community Engagement Fundamentals", "CSCCE", "", []),
]

PUBLICATIONS = [
    "Strauss, T. (2021). Data Management Plan Template: Neuroimaging in the Neurosciences. Zenodo. doi:10.5281/zenodo.4673558",
    "Poline, J.-B., et al. (2023). Data and tools integration in the Canadian open neuroscience platform. Scientific Data, 10(1), 189.",
    "Halchenko, Y. O., et al. (2021). DataLad: distributed system for joint management of code, data, and their relationship. Journal of Open Source Software, 6(63), 3262.",
]
PUB_FOOT = "9 publications | citation count 630 | h-index 9 | Full list: scholar.google.com/citations"

COVER_DATE = "August 7, 2026"
COVER_SUBJECT = "Subject: Application: Documentation Control & Technical Publications Specialist"
COVER_SALUTATION = "Dear Hiring Manager,"
COVER_PARAS = [
    "I am applying for the Documentation Control & Technical Publications Specialist position in "
    "Boucherville. Fuji SEMEC builds door opening and closing devices for railcars, light rail, and "
    "transit across three plants in Boucherville, Plattsburgh, and Edison. Safety-critical hardware "
    "built to customer specification is a documentation problem before it is anything else: the revision "
    "an engineer works from has to be the revision quality signed off on, and both have to be provable a "
    "year later in an audit. Holding that discipline is what I have spent my career doing.",

    "At McGill University's McConnell Brain Imaging Centre I owned the controlled-document side of a "
    "research data function serving dozens of research groups. I wrote the naming and numbering "
    "conventions, the metadata standards and data definitions, the review and approval workflow, and the "
    "classification, retention, and access policies. I maintained Collibra as the institutional "
    "catalogue so definitions had one approved source rather than competing copies moving by email, and "
    "I replaced manual document-request handling with a self-serve portal that recorded what was "
    "released, to whom, and under which version. That is the shape of improvement I would look for here: "
    "better traceability reached by removing a manual step rather than adding a form.",

    "On the publications half of the role, I have written for specialist and general readers for twenty "
    "years: nine peer-reviewed publications, a public knowledge base built from scratch that explains "
    "data management to readers with no background in it, and a documentation template published on "
    "Zenodo that other researchers adopted as their own controlled template. Sitting with a "
    "subject-matter expert, extracting what they know, and returning it as something a technician or a "
    "customer can follow is the skill I have practised most. At Premergency I did the same under "
    "regulatory conditions, keeping the compliance and approval trail government deployments required.",

    "Three things I want to be straight about. My document control experience is in research and "
    "regulated software rather than manufacturing, so I have not worked to ISO 9001 or the EN standards "
    "and have not been inside a CMMI appraisal. I read schematics and specifications, but interpreting "
    "engineering drawings is something I would have to learn. And "
    "my tooling is SharePoint, Collibra, Git, and the Adobe suite rather than Windchill, Teamcenter, "
    "DOORS, FrameMaker, or MadCap Flare. The underlying practice transfers, and I have picked up "
    "enterprise platforms under deadline before.",

    "I am bilingual in French and English at professional working proficiency, based in Montreal and a "
    "short drive from Boucherville, so the hybrid arrangement works well. I would welcome the "
    "opportunity to discuss how I can tighten the documentation practice behind your engineering and "
    "quality teams. Thank you for your consideration.",
]
COVER_CLOSING = "Kind regards,"

# ==================================================================
# DOCX RENDERING
# ==================================================================
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BLUE = RGBColor(0x21, 0x6F, 0xAA)
BLACK = RGBColor(0x00, 0x00, 0x00)
GRAY = RGBColor(0x55, 0x55, 0x55)


def _margins(doc, top=1.5, bottom=1.5, left=2.0, right=2.0):
    for s in doc.sections:
        s.top_margin, s.bottom_margin = Cm(top), Cm(bottom)
        s.left_margin, s.right_margin = Cm(left), Cm(right)


def _name_header(doc, name, contact):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(name); r.bold = True; r.font.size = Pt(28); r.font.color.rgb = BLUE
    p2 = doc.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(contact); r2.font.size = Pt(10); r2.font.color.rgb = GRAY
    p2.paragraph_format.space_after = Pt(6)


def _heading(doc, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title.upper()); r.bold = True; r.font.size = Pt(11); r.font.color.rgb = BLUE
    pPr = p._p.get_or_add_pPr(); pBdr = OxmlElement('w:pBdr'); bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single'); bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1'); bottom.set(qn('w:color'), '216FAA')
    pBdr.append(bottom); pPr.append(pBdr)


def _entry(doc, dates, title, org, location, bullets):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    r1 = p.add_run(title); r1.bold = True; r1.font.size = Pt(10.5); r1.font.color.rgb = BLACK
    r2 = p.add_run(f"  —  {org}"); r2.font.size = Pt(10.5); r2.font.color.rgb = GRAY
    if location:
        r3 = p.add_run(f", {location}"); r3.font.size = Pt(10.5); r3.font.color.rgb = GRAY
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_before = Pt(0); p2.paragraph_format.space_after = Pt(2)
    p2.paragraph_format.keep_with_next = True
    r4 = p2.add_run(dates); r4.italic = True; r4.font.size = Pt(9.5); r4.font.color.rgb = GRAY
    for b in bullets:
        _bullet(doc, b)


def _bullet(doc, text, indent=0.25):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(1); p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.left_indent = Inches(indent)
    for i, part in enumerate(text.split('**')):
        if not part:
            continue
        r = p.add_run(part); r.font.size = Pt(10); r.bold = (i % 2 == 1)


def _para(doc, text, size=10, italic=False, color=None, space_after=4):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(text); r.font.size = Pt(size); r.italic = italic
    if color:
        r.font.color.rgb = color
    return p


def build_docx_cv():
    doc = Document(); _margins(doc)
    _name_header(doc, NAME, CV_CONTACT)
    _para(doc, PROFILE, size=10, space_after=6)
    _heading(doc, "Core Competencies")
    for label, rest in COMPETENCIES:
        _bullet(doc, f"**{label}**{rest}")
    _heading(doc, "Relevant Experience")
    for e in EXPERIENCE:
        _entry(doc, *e)
    _heading(doc, "Additional Experience")
    for a in ADDITIONAL:
        _bullet(doc, a)
    _heading(doc, "Education")
    for e in EDUCATION:
        _entry(doc, *e)
    _heading(doc, "Languages")
    _para(doc, "English (native), French (bilingual, professional working proficiency).", size=10)
    _heading(doc, "Selected Publications and Authored Documentation")
    for pub in PUBLICATIONS:
        _bullet(doc, pub)
    _para(doc, PUB_FOOT, size=9, italic=True, color=GRAY)
    doc.save("docx/cv_fujisemec_doccontrol.docx")
    print("OK cv_fujisemec_doccontrol.docx")


def build_docx_cover():
    doc = Document(); _margins(doc, top=2.0, bottom=2.0, left=2.5, right=2.5)
    _name_header(doc, NAME, COVER_CONTACT)
    _para(doc, COVER_DATE, size=10, space_after=12)
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(6)
    r = p.add_run(COVER_SUBJECT); r.bold = True; r.font.size = Pt(10)
    _para(doc, COVER_SALUTATION, size=10, space_after=6)
    for para in COVER_PARAS:
        pp = doc.add_paragraph(); pp.paragraph_format.space_after = Pt(8)
        rr = pp.add_run(para); rr.font.size = Pt(10)
    _para(doc, COVER_CLOSING, size=10, space_after=24)
    _para(doc, NAME, size=10)
    doc.save("docx/cover_fujisemec_doccontrol.docx")
    print("OK cover_fujisemec_doccontrol.docx")


# ==================================================================
# PDF RENDERING
# ==================================================================
from reportlab.lib.pagesizes import letter
from reportlab.lib.units import cm
from reportlab.lib.colors import HexColor
from reportlab.lib.enums import TA_CENTER
from reportlab.platypus import (
    BaseDocTemplate, PageTemplate, Frame, Paragraph, Spacer,
    HRFlowable, ListFlowable, ListItem, KeepTogether,
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

PDF_BLUE = HexColor("#216FAA")
PDF_GRAY = HexColor("#555555")


def _esc(t):
    return t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _bold_md(t):
    t = _esc(t)
    parts = t.split("**")
    out = ""
    for i, p in enumerate(parts):
        out += f"<b>{p}</b>" if i % 2 == 1 else p
    return out


def _styles():
    styles = getSampleStyleSheet()
    return {
        "name": ParagraphStyle("Name", parent=styles["Normal"], fontName="Helvetica-Bold",
                               fontSize=22, textColor=PDF_BLUE, alignment=TA_CENTER, spaceAfter=2, leading=26),
        "contact": ParagraphStyle("Contact", parent=styles["Normal"], fontName="Helvetica",
                                  fontSize=9, textColor=PDF_GRAY, alignment=TA_CENTER, spaceAfter=8, leading=12),
        "heading": ParagraphStyle("Heading", parent=styles["Normal"], fontName="Helvetica-Bold",
                                  fontSize=11, textColor=PDF_BLUE, spaceBefore=8, spaceAfter=2, leading=13),
        "body": ParagraphStyle("Body", parent=styles["Normal"], fontName="Helvetica", fontSize=10,
                               leading=13, spaceAfter=3),
        "bullet": ParagraphStyle("Bullet", parent=styles["Normal"], fontName="Helvetica",
                                 fontSize=9.5, leading=12, spaceAfter=1),
        "entry_title": ParagraphStyle("EntryTitle", parent=styles["Normal"], fontName="Helvetica",
                                      fontSize=10.5, spaceBefore=5, spaceAfter=0, leading=12),
        "entry_dates": ParagraphStyle("EntryDates", parent=styles["Normal"], fontName="Helvetica",
                                      fontSize=9, textColor=PDF_GRAY, spaceAfter=2, leading=11),
        "letter": ParagraphStyle("Letter", parent=styles["Normal"], fontName="Helvetica",
                                 fontSize=10.5, leading=14, spaceAfter=7),
        "small_italic": ParagraphStyle("SmallItalic", parent=styles["Normal"],
                                       fontName="Helvetica-Oblique", fontSize=8.5,
                                       textColor=PDF_GRAY, spaceAfter=2, leading=11),
    }


def _render(path, story, top=1.5, bottom=1.3, left=2.0, right=2.0):
    doc = BaseDocTemplate(path, pagesize=letter,
                          topMargin=top * cm, bottomMargin=bottom * cm,
                          leftMargin=left * cm, rightMargin=right * cm)
    frame = Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id="main")
    doc.addPageTemplates([PageTemplate(id="all", frames=[frame])])
    doc.build(story)
    print(f"OK {os.path.basename(path)}")


def build_pdf_cv():
    S = _styles()

    def rule():
        return HRFlowable(width="100%", thickness=0.6, color=PDF_BLUE, spaceBefore=1, spaceAfter=3)

    def heading(text):
        return [Paragraph(_esc(text.upper()), S["heading"]), rule()]

    def bullets(items):
        return ListFlowable(
            [ListItem(Paragraph(_bold_md(t), S["bullet"]), leftIndent=10, value="•") for t in items],
            bulletType="bullet", bulletColor=PDF_BLUE, leftIndent=12, bulletFontSize=8,
        )

    def entry(dates, title, org, location, bl):
        loc = f", {_esc(location)}" if location else ""
        head = (f'<b>{_esc(title)}</b>'
                f'<font color="#555555">&nbsp;&nbsp;—&nbsp;&nbsp;{_esc(org)}{loc}</font>')
        header = [Paragraph(head, S["entry_title"]), Paragraph(_esc(dates), S["entry_dates"])]
        if not bl:
            return [KeepTogether(header)]
        # Short entries stay whole; splitting a two-bullet entry across a page
        # break leaves a stranded second bullet at the top of the next page.
        if len(bl) <= 2:
            return [KeepTogether(header + [bullets(bl)])]
        out = [KeepTogether(header + [bullets(bl[:1])])]
        if len(bl) > 1:
            out.append(bullets(bl[1:]))
        return out

    story = [Paragraph(_esc(NAME), S["name"]),
             Paragraph(_esc(CV_CONTACT), S["contact"]),
             Paragraph(_esc(PROFILE), S["body"])]
    story += heading("Core Competencies")
    story.append(bullets([f"**{l}**{r}" for l, r in COMPETENCIES]))
    story += heading("Relevant Experience")
    for e in EXPERIENCE:
        story += entry(*e)
    story += heading("Additional Experience")
    story.append(bullets(ADDITIONAL))
    story += heading("Education")
    for e in EDUCATION:
        story += entry(*e)
    story += heading("Languages")
    story.append(Paragraph("English (native), French (bilingual, professional working proficiency).", S["body"]))
    story += heading("Selected Publications and Authored Documentation")
    story.append(bullets(PUBLICATIONS))
    story.append(Paragraph(_esc(PUB_FOOT), S["small_italic"]))

    _render("docx/cv_fujisemec_doccontrol.pdf", story)


def build_pdf_cover():
    S = _styles()
    story = [Paragraph(_esc(NAME), S["name"]),
             Paragraph(_esc(COVER_CONTACT), S["contact"]),
             Spacer(1, 6),
             Paragraph(_esc(COVER_DATE), S["body"]),
             Spacer(1, 8),
             Paragraph(f"<b>{_esc(COVER_SUBJECT)}</b>", S["body"]),
             Spacer(1, 4),
             Paragraph(_esc(COVER_SALUTATION), S["letter"])]
    for para in COVER_PARAS:
        story.append(Paragraph(_esc(para), S["letter"]))
    story += [Spacer(1, 2),
              KeepTogether([Paragraph(_esc(COVER_CLOSING), S["body"]),
                            Spacer(1, 10),
                            Paragraph(_esc(NAME), S["body"])])]
    _render("docx/cover_fujisemec_doccontrol.pdf", story, top=1.8, bottom=1.1, left=2.5, right=2.5)


if __name__ == "__main__":
    build_docx_cv()
    build_docx_cover()
    build_pdf_cv()
    build_pdf_cover()
    print("\nDone. Files in docx/")
