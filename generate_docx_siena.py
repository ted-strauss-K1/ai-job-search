"""Generate DOCX CV for Siena AI — Head of Education & Community."""
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

os.makedirs("docx", exist_ok=True)

BLUE = RGBColor(0x21, 0x6F, 0xAA)
BLACK = RGBColor(0x00, 0x00, 0x00)
GRAY = RGBColor(0x55, 0x55, 0x55)


def set_margins(doc, top=1.5, bottom=1.5, left=2.0, right=2.0):
    for section in doc.sections:
        section.top_margin = Cm(top)
        section.bottom_margin = Cm(bottom)
        section.left_margin = Cm(left)
        section.right_margin = Cm(right)


def add_name_header(doc, name, contact):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(name)
    run.bold = True
    run.font.size = Pt(28)
    run.font.color.rgb = BLUE
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run2 = p2.add_run(contact)
    run2.font.size = Pt(10)
    run2.font.color.rgb = GRAY
    p2.paragraph_format.space_after = Pt(6)


def add_section_heading(doc, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(title.upper())
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = BLUE
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), '216FAA')
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_job_entry(doc, dates, title, org, location, bullets):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(0)
    r1 = p.add_run(title)
    r1.bold = True
    r1.font.size = Pt(10.5)
    r2 = p.add_run(f"  —  {org}")
    r2.font.size = Pt(10.5)
    r2.font.color.rgb = GRAY
    if location:
        r3 = p.add_run(f", {location}")
        r3.font.size = Pt(10.5)
        r3.font.color.rgb = GRAY
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(2)
    r4 = p2.add_run(dates)
    r4.italic = True
    r4.font.size = Pt(9.5)
    r4.font.color.rgb = GRAY
    for b in bullets:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_before = Pt(1)
        bp.paragraph_format.space_after = Pt(1)
        bp.paragraph_format.left_indent = Inches(0.25)
        parts = b.split('**')
        for i, part in enumerate(parts):
            if not part:
                continue
            r = bp.add_run(part)
            r.font.size = Pt(10)
            r.bold = (i % 2 == 1)


def add_bullet(doc, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.left_indent = Inches(0.25)
    parts = text.split('**')
    for i, part in enumerate(parts):
        if not part:
            continue
        r = p.add_run(part)
        r.font.size = Pt(10)
        r.bold = (i % 2 == 1)


def add_para(doc, text, size=10, italic=False, color=None, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.italic = italic
    if color:
        r.font.color.rgb = color
    return p


def cv_siena():
    doc = Document()
    set_margins(doc)
    add_name_header(
        doc,
        "Ted Strauss",
        "Montreal, Quebec, Canada  |  514-475-2374  |  ted.strauss@gmail.com  |  linkedin.com/in/tedstrauss-63b01821",
    )

    add_para(
        doc,
        "Learning architect and community builder with a decade of experience creating educational programs, "
        "certification paths, and community spaces that make complex technology genuinely accessible. Built "
        "self-serve education portals and public knowledge bases at McGill University's research institutes; "
        "designed and delivered hands-on workshops for hundreds of researchers, faculty, and government clients; "
        "founded and operated a community school with 150 workshop groups and 20 instructors. Actively builds "
        "with AI tools — including daily use of Claude Code for agentic software workflows — and has evaluated "
        "AI adoption pathways in enterprise learning contexts. A storyteller who translates technical depth into "
        "learning experiences that change how people work.",
        size=10,
        space_after=6,
    )

    add_section_heading(doc, "Core Competencies")
    add_bullet(
        doc,
        "**Education & curriculum design**: built Moodle-based education portals with certifications, learning "
        "paths, and onboarding sequences; designed instructional programs for researchers, government clients, "
        "and non-specialist audiences; produced public knowledge bases and open educational content from scratch",
    )
    add_bullet(
        doc,
        "**Community building & engagement**: founded and operated a community learning organization (Sur Place "
        "Media) with 150 workshop groups and 20 instructors; built communities of practice around data governance "
        "at McGill; developed ongoing engagement programs to embed new practices into workflows",
    )
    add_bullet(
        doc,
        "**AI tool building & adoption**: builds daily with Claude Code (agentic workflows, vibe coding); "
        "evaluated and integrated AI tools into enterprise learning contexts at KnowledgeOne; assessed AI "
        "adoption pathways, governance, and usability for non-specialist users",
    )
    add_bullet(
        doc,
        "**Storytelling & knowledge translation**: consistent track record of translating neuroimaging data "
        "practices, research workflows, and technical platforms into accessible language for diverse audiences; "
        "produced video, podcast, and multimedia educational content",
    )
    add_bullet(
        doc,
        "**Shipped from conception**: launched OwnCloud self-serve data portal, Moodle education portal, "
        "public open science knowledge base, and Sur Place Media — each conceived, built, and delivered "
        "end-to-end without prior templates",
    )

    add_section_heading(doc, "Professional Experience")
    add_job_entry(
        doc,
        "07/2025 – 01/2026",
        "Moodle Analyst",
        "Premergency Inc.",
        "Canada",
        [
            "Developed product education sessions, documentation, and onboarding pathways for government LMS clients across multiple market segments",
            "Facilitated customer workshops and recurring training sessions; maintained regulatory compliance documentation throughout",
            "Designed bespoke learning & development programs from conception through delivery, coordinating between government clients and technical teams",
        ],
    )
    add_job_entry(
        doc,
        "01/2024 – 10/2024",
        "Full Stack Developer",
        "KnowledgeOne",
        "Montreal, QC",
        [
            "Evaluated and integrated AI tools into enterprise learning contexts, assessing adoption pathways, "
            "usability, and learning impact for non-specialist users",
            "Developed and shipped LMS product features from specification to production (LAMP++, Moodle, LXD)",
        ],
    )
    add_job_entry(
        doc,
        "09/2016 – 10/2021",
        "Manager of Data Resources",
        "McConnell Brain Imaging Centre, The Neuro, McGill University",
        "Montreal, QC",
        [
            "Conceived and launched a Moodle-based self-serve education portal — curriculum, certifications, "
            "and learning paths for researcher onboarding — reducing time-to-productivity for new staff",
            "Built community of practice around data governance across a multi-PI research environment "
            "supporting dozens of principal investigators and hundreds of researchers",
            "Translated complex neuroimaging data workflows into accessible documentation, training materials, "
            "and platform guides for audiences ranging from graduate students to senior faculty",
            "Created demo libraries and recorded walkthroughs for data management tools and platform workflows",
        ],
    )
    add_job_entry(
        doc,
        "08/2020 – 10/2021",
        "Data Management & Sharing Coordinator (p/t)",
        "Tanenbaum Open Science Institute, The Neuro, McGill University",
        "Montreal, QC",
        [
            "Built a public-facing open science knowledge base — open courses and educational content written "
            "for both specialists and general audiences, shipped from scratch",
            "Designed and delivered training workshops for students, faculty, and research teams; created "
            "curriculum from zero on data management, open access, and research workflows",
        ],
    )
    add_job_entry(
        doc,
        "04/2017 – 08/2021",
        "Founder and Director",
        "Sur Place Media",
        "Montreal, QC",
        [
            "Founded a community learning organization from concept: built curriculum, recruited 20 instructors, "
            "hosted 150 workshop groups across creative writing, media arts, and visual arts",
            "Built and maintained community spaces for learners, instructors, and partners; ran recurring "
            "programming and managed bilingual (EN/FR) learner engagement",
            "Shipped fully virtual delivery during COVID-19 pivot; produced video, podcast, and multimedia "
            "educational content throughout",
        ],
    )
    add_job_entry(
        doc,
        "02/2015 – 09/2016",
        "Software Developer",
        "McGill Centre for Integrative Neuroscience, McGill University",
        "Montreal, QC",
        [
            "Programmed the data acquisition pipeline for the national CCNA longitudinal study (PHP, MySQL, "
            "JavaScript); contributed to user documentation and onboarding materials for research platform",
        ],
    )

    add_section_heading(doc, "Additional Experience")
    add_bullet(doc, "Language Mate — Developer and UX designer (2007–2010)")
    add_bullet(doc, "Moss Rehabilitation Hospital — Research software developer (2007–2010)")
    add_bullet(doc, "University of Connecticut & Columbia University — Research assistant & software developer (2004–2008)")

    add_section_heading(doc, "Education")
    add_job_entry(
        doc,
        "1998–2003",
        "B.A. (Honours), Cognitive Science",
        "McGill University",
        "Montreal, QC",
        ["Minor: Music Technology. Topics: Psycholinguistics, cognitive science, computational modeling of language."],
    )
    add_job_entry(doc, "2021", "Scientific Community Engagement Fundamentals", "CSCCE", "", [])

    add_section_heading(doc, "Languages")
    add_para(doc, "English (native), French (bilingual).", size=10)

    add_section_heading(doc, "Selected Publications")
    add_bullet(
        doc,
        "Poline, J.-B., et al. (2023). Data and tools integration in the Canadian open neuroscience platform. "
        "Scientific Data, 10(1), 189.",
    )
    add_bullet(
        doc,
        "Halchenko, Y. O., et al. (2021). DataLad: distributed system for joint management of code, data, "
        "and their relationship. JOSS, 6(63), 3262.",
    )
    add_bullet(
        doc,
        "Strauss, T. (2021). Data Management Plan Template: Neuroimaging in the Neurosciences. Zenodo. "
        "doi:10.5281/zenodo.4673558",
    )
    add_para(
        doc,
        "9 publications | citation count: 630 | h-index: 9 | scholar.google.com/citations",
        size=9,
        italic=True,
        color=GRAY,
    )

    add_section_heading(doc, "References")
    add_para(doc, "Available upon request.", size=10)

    doc.save("docx/cv_siena.docx")
    print("cv_siena.docx")


cv_siena()
print("Done. File: docx/cv_siena.docx")
