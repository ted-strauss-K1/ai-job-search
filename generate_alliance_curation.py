"""
Generate application docs for the Digital Research Alliance of Canada -
Curation Officer (data curation team; FRDR and active data management support).
Remote, contract to 2028-03-31, $70,550 - $95,450 CAD. Posted 2026-09-29.
Apply via the Alliance's ADP portal. English documents. No deadline stated.
Outputs (in docx/): cv_alliance_curation, cover_alliance_curation,
application_alliance_curation (.docx and .pdf each).

Content is defined ONCE below, then rendered to both docx (python-docx) and
pdf (reportlab) so the two formats never drift.

CANDIDATE-CONFIRMED FACTS (2026-10-02): at the BIC Ted worked on projects that
involved depositing data to or sourcing data from FRDR, Zenodo, and OpenNeuro;
worked directly with DataCite, BIDS, DOI, ORCID, and DMP standards in developing
data repositories; handled protected human-subject data from clinical and
research contexts (privacy and security a daily responsibility); attended
workshops and conferences on Indigenous data sovereignty. He is "familiar with"
other repositories: NOT claimed as hands-on.

GAPS STATED HONESTLY: no library / information-science degree or library
employment (the posting accepts strong applied research experience as a data
manager, or equivalent); contract, not permanent.

Company facts (FRDR is a bilingual, curated repository operated by the Alliance,
reserves DOIs, curation team reviews submissions for FAIR) verified at
frdr-dfdr.ca, not from the posting. The posting's service name for active data
management is not independently verified, so the letter describes it generically.
"""
import os

os.makedirs("docx", exist_ok=True)

# ------------------------------------------------------------------
# SHARED CONTENT
# ------------------------------------------------------------------

NAME = "Ted Strauss"
CV_CONTACT = "Montreal, Quebec, Canada  |  514-475-2374  |  ted.strauss@gmail.com  |  linkedin.com/in/tedstrauss-63b01821"
COVER_CONTACT = "514-475-2374  |  ted.strauss@gmail.com  |  linkedin.com/in/tedstrauss-63b01821"

PROFILE = (
    "Research data manager with five years running the data resources of McGill's McConnell Brain Imaging "
    "Centre (BIC): worked directly with DataCite, BIDS, DOI, ORCID, and DMP standards in developing data "
    "repositories, on projects that deposited data to and sourced data from FRDR, Zenodo, and OpenNeuro, and "
    "supported dozens of principal investigators and hundreds of researchers with data management, "
    "documentation, and training. Responsible daily for the privacy and security of protected human-subject "
    "data from clinical and research settings. Co-author of the DataLad and Canadian Open Neuroscience "
    "Platform papers and author of a published data management plan template. Bilingual (EN/FR), Montreal-based."
)

COMPETENCIES = [
    ("Research data curation and repositories",
     ": worked directly with DataCite, BIDS, DOI, ORCID, and DMP standards in developing data repositories; "
     "worked on projects depositing data to and sourcing data from FRDR, Zenodo, and OpenNeuro"),
    ("Metadata, documentation, and data quality",
     ": developed metadata standards and quality practices for daily data workflows; investigated data "
     "discrepancies to root cause; built a public knowledge base on open science and FAIR data"),
    ("Researcher consulting and support",
     ": advised principal investigators and research teams on data collection, organization, management, and "
     "sharing; authored a published data management plan template (Zenodo)"),
    ("Sensitive data, privacy, and security",
     ": protected human-subject data from clinical and research contexts was a daily responsibility at the BIC; "
     "attended workshops and conferences on Indigenous data sovereignty"),
    ("Training, community, and national engagement",
     ": designed and delivered data management workshops; built communities of practice; contributed to "
     "national research data policy through Portage, CONP, and the Research Data Alliance; CSCCE-certified "
     "in community engagement"),
    ("Open science and tools",
     ": co-author, DataLad (JOSS), HeuDiConv (DICOM conversion), and the Canadian Open Neuroscience Platform "
     "paper (Scientific Data); Python, R, SQL, and Git"),
    ("Bilingual EN/FR",
     ": professional working proficiency in French and English, Montreal-based"),
]

# each: (dates, title, org, location, [bullets]) - ordered by relevance
EXPERIENCE = [
    ("09/2016 - 10/2021", "Manager of Data Resources", "McConnell Brain Imaging Centre, The Neuro, McGill University", "Montreal, QC", [
        "Worked directly with DataCite, BIDS, DOI, ORCID, and DMP standards in developing data repositories, on projects that involved depositing data to and sourcing data from FRDR, Zenodo, and OpenNeuro",
        "Supported dozens of principal investigators and hundreds of researchers on data collection, management, and sharing, and built the pipelines and research applications that captured their data",
        "Handled protected human-subject data from clinical and research contexts, with privacy and security a daily responsibility of the role",
        "Developed data acquisition policies, metadata standards, and quality practices, and investigated data discrepancies to root cause with engineering and research teams",
        "Designed and launched an OwnCloud self-serve portal that replaced manual data-request handling, and a Moodle-based self-serve education portal to embed data practices in researchers' routines",
        "Contributed to national research data policy through Portage, the Canadian Open Neuroscience Platform (CONP), and Research Data Alliance committee work",
    ]),
    ("08/2020 - 10/2021", "Data Management and Sharing Coordinator (p/t)", "Tanenbaum Open Science Institute, The Neuro, McGill University", "Montreal, QC", [
        "Designed and administered data management training workshops for students and faculty",
        "Built a public online knowledge base on open science and FAIR data concepts",
    ]),
    ("02/2015 - 09/2016", "Software Developer, national CCNA dementia study", "McGill Centre for Integrative Neuroscience", "Montreal, QC", [
        "Programmed the data acquisition pipeline for a pan-Canadian longitudinal study on aging (LORIS platform: PHP, MySQL, JavaScript)",
    ]),
    ("07/2025 - 01/2026", "Moodle Analyst", "Premergency Inc.", "Canada", [
        "Coordinated between government clients and development teams on learning-platform solutions, maintaining regulatory compliance for government deployments; migrated legacy LMS systems to Moodle",
    ]),
    ("01/2024 - 10/2024", "Full Stack Developer", "KnowledgeOne", "Montreal, QC", [
        "Developed learning platform features for active clients (LAMP++, Moodle, LXD); evaluated and adopted AI tools including Claude Code",
    ]),
    ("04/2017 - 08/2021", "Founder and Director", "Sur Place Media", "Montreal, QC", [
        "Founded and directed a neighbourhood arts school; coordinated 20 instructors across 150 workshop groups and led the pivot to virtual delivery during the pandemic",
    ]),
]

ADDITIONAL = [
    "Photonic Knowledge - Database administrator and software developer (2012-2014): Drupal ERP system and its data model",
    "Language Mate; Moss Rehabilitation Hospital; University of Connecticut; Columbia University; Universität Bremen - research assistant and research software developer roles (2003-2010)",
]

EDUCATION = [
    ("1998-2003", "Bachelor of Arts (Honours), Cognitive Science", "McGill University", "Montreal, QC",
     ["Minor: Music Technology. Topics: psycholinguistics, cognitive science, computational modeling of language."]),
    ("2021", "Scientific Community Engagement Fundamentals", "The Center for Scientific Collaboration and Community Engagement (CSCCE)", "", []),
]

PUBLICATIONS = [
    "Halchenko, Y. O., et al. (2024). HeuDiConv: flexible DICOM conversion into structured directory layouts. Journal of Open Source Software, 9(99), 5839.",
    "Poline, J.-B., et al. (2023). Data and tools integration in the Canadian open neuroscience platform. Scientific Data, 10(1), 189.",
    "Halchenko, Y. O., et al. (2021). DataLad: distributed system for joint management of code, data, and their relationship. Journal of Open Source Software, 6(63), 3262.",
    "Strauss, T. (2021). Data Management Plan Template: Neuroimaging in the Neurosciences. Zenodo. doi:10.5281/zenodo.4673558",
]
PUB_FOOT = "9 publications | citation count 630 | h-index 9 | Full list: scholar.google.com/citations"

# ------------------------------------------------------------------
# COVER LETTER
# ------------------------------------------------------------------

COVER_DATE = "October 2, 2026"
COVER_SUBJECT = "Subject: Application: Curation Officer"
COVER_SALUTATION = "Dear Members of the Selection Committee,"
COVER_PARAS = [
    "I am applying for the Curation Officer position on the Alliance's data curation team. The role is to "
    "make new deposits well documented and independently understandable, and to guide researchers in "
    "organizing and describing their data while the research is still active. That is the work I did for "
    "five years as Manager of Data Resources at McGill's McConnell Brain Imaging Centre, and I would like "
    "to do it at national scale.",

    "FRDR is a bilingual repository run by the Alliance, where a curation team reviews each submission to "
    "improve its FAIR compliance before publication. I have been on both sides of that process: at the "
    "centre I worked on projects that deposited data to and sourced data from FRDR, Zenodo, and OpenNeuro, "
    "and worked directly with DataCite, BIDS, DOI, ORCID, and DMP standards in developing data "
    "repositories. In practice, I would bring:",

    "I want to be clear about two edges of my background. I do not hold a library or information science "
    "degree, and I have not worked in a library; I come to curation as a research data manager, which "
    "your posting recognizes as a qualifying route, with a neuroscience domain and publications in the "
    "field. On Indigenous data sovereignty, I have attended workshops and conferences and would treat it "
    "as an area for continued learning, escalating to specialists as the role expects.",

    "I work in French and English, am based in Montreal and available now, and am comfortable in a "
    "distributed team. I would welcome the opportunity to discuss the role. Thank you for your "
    "consideration.",
]
COVER_BULLETS = [
    "Standards in practice: DataCite, BIDS, DOI, ORCID, and DMP standards applied in building repositories, plus a published data management plan template and co-authorship of the DataLad and Canadian Open Neuroscience Platform papers",
    "Researcher support: advice to dozens of principal investigators on organizing, documenting, and sharing data, backed by a public knowledge base and workshops that made good practice easy to follow",
    "Sensitive data: protected human-subject data from clinical and research contexts, where privacy and security were a daily responsibility, which is the judgment the role asks for when escalating governance and privacy questions",
    "Community: communities of practice and national committee work (Portage, CONP, and the Research Data Alliance) that grew curation and data management capacity",
]
COVER_CLOSING = "Kind regards,"

# ==================================================================
# DOCX RENDERING
# ==================================================================
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BLUE = RGBColor(0x21, 0x6F, 0xAA)
BLACK = RGBColor(0x00, 0x00, 0x00)
GRAY = RGBColor(0x55, 0x55, 0x55)


def _margins(doc, top=1.5, bottom=1.5, left=2.0, right=2.0):
    for s in doc.sections:
        s.top_margin, s.bottom_margin = Cm(top), Cm(bottom)
        s.left_margin, s.right_margin = Cm(left), Cm(right)


def _name_header(doc, name, contact):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(name); r.bold = True; r.font.size = Pt(28); r.font.color.rgb = BLUE
    p2 = doc.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(contact); r2.font.size = Pt(10); r2.font.color.rgb = GRAY
    p2.paragraph_format.space_after = Pt(6)


def _heading(doc, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title.upper()); r.bold = True; r.font.size = Pt(11); r.font.color.rgb = BLUE
    pPr = p._p.get_or_add_pPr(); pBdr = OxmlElement('w:pBdr'); bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single'); bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1'); bottom.set(qn('w:color'), '216FAA')
    pBdr.append(bottom); pPr.append(pBdr)


def _entry(doc, dates, title, org, location, bullets):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    r1 = p.add_run(title); r1.bold = True; r1.font.size = Pt(10.5); r1.font.color.rgb = BLACK
    r2 = p.add_run(f"  —  {org}"); r2.font.size = Pt(10.5); r2.font.color.rgb = GRAY
    if location:
        r3 = p.add_run(f", {location}"); r3.font.size = Pt(10.5); r3.font.color.rgb = GRAY
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_before = Pt(0); p2.paragraph_format.space_after = Pt(2)
    p2.paragraph_format.keep_with_next = True
    r4 = p2.add_run(dates); r4.italic = True; r4.font.size = Pt(9.5); r4.font.color.rgb = GRAY
    for b in bullets:
        _bullet(doc, b)


def _bullet(doc, text, indent=0.25):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(1); p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.left_indent = Inches(indent)
    for i, part in enumerate(text.split('**')):
        if not part:
            continue
        r = p.add_run(part); r.font.size = Pt(10); r.bold = (i % 2 == 1)


def _para(doc, text, size=10, italic=False, color=None, space_after=4):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(text); r.font.size = Pt(size); r.italic = italic
    if color:
        r.font.color.rgb = color
    return p


def build_docx_cv(doc=None):
    standalone = doc is None
    if standalone:
        doc = Document(); _margins(doc)
    _name_header(doc, NAME, CV_CONTACT)
    _para(doc, PROFILE, size=10, space_after=6)
    _heading(doc, "Core Competencies")
    for label, rest in COMPETENCIES:
        _bullet(doc, f"**{label}**{rest}")
    _heading(doc, "Professional Experience")
    for e in EXPERIENCE:
        _entry(doc, *e)
    _heading(doc, "Additional Experience")
    for a in ADDITIONAL:
        _bullet(doc, a)
    _heading(doc, "Education")
    for e in EDUCATION:
        _entry(doc, *e)
    _heading(doc, "Languages")
    _para(doc, "English (native), French (bilingual, professional working proficiency).", size=10)
    _heading(doc, "Selected Publications")
    for pub in PUBLICATIONS:
        _bullet(doc, pub)
    _para(doc, PUB_FOOT, size=9, italic=True, color=GRAY)
    _heading(doc, "References")
    _para(doc, "Available upon request.", size=10)
    if standalone:
        doc.save("docx/cv_alliance_curation.docx")
        print("OK cv_alliance_curation.docx")


def build_docx_cover(doc=None):
    standalone = doc is None
    if standalone:
        doc = Document(); _margins(doc, top=2.0, bottom=2.0, left=2.5, right=2.5)
    _name_header(doc, NAME, COVER_CONTACT)
    _para(doc, COVER_DATE, size=10, space_after=12)
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(6)
    r = p.add_run(COVER_SUBJECT); r.bold = True; r.font.size = Pt(10)
    _para(doc, COVER_SALUTATION, size=10, space_after=6)
    for i, para in enumerate(COVER_PARAS):
        pp = doc.add_paragraph(); pp.paragraph_format.space_after = Pt(8)
        rr = pp.add_run(para); rr.font.size = Pt(10)
        if i == 1:
            for b in COVER_BULLETS:
                _bullet(doc, b)
    _para(doc, COVER_CLOSING, size=10, space_after=24)
    _para(doc, NAME, size=10)
    if standalone:
        doc.save("docx/cover_alliance_curation.docx")
        print("OK cover_alliance_curation.docx")


# ==================================================================
# PDF RENDERING
# ==================================================================
from reportlab.lib.pagesizes import letter
from reportlab.lib.units import cm
from reportlab.lib.colors import HexColor
from reportlab.lib.enums import TA_CENTER
from reportlab.platypus import (
    BaseDocTemplate, PageTemplate, Frame, Paragraph, Spacer,
    HRFlowable, ListFlowable, ListItem, KeepTogether,
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

PDF_BLUE = HexColor("#216FAA")
PDF_GRAY = HexColor("#555555")


def _esc(t):
    return t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _bold_md(t):
    t = _esc(t)
    parts = t.split("**")
    out = ""
    for i, p in enumerate(parts):
        out += f"<b>{p}</b>" if i % 2 == 1 else p
    return out


def _styles():
    styles = getSampleStyleSheet()
    return {
        "name": ParagraphStyle("Name", parent=styles["Normal"], fontName="Helvetica-Bold",
                               fontSize=22, textColor=PDF_BLUE, alignment=TA_CENTER, spaceAfter=2, leading=26),
        "contact": ParagraphStyle("Contact", parent=styles["Normal"], fontName="Helvetica",
                                  fontSize=9, textColor=PDF_GRAY, alignment=TA_CENTER, spaceAfter=8, leading=12),
        "heading": ParagraphStyle("Heading", parent=styles["Normal"], fontName="Helvetica-Bold",
                                  fontSize=11, textColor=PDF_BLUE, spaceBefore=8, spaceAfter=2, leading=13),
        "body": ParagraphStyle("Body", parent=styles["Normal"], fontName="Helvetica", fontSize=10,
                               leading=13, spaceAfter=3),
        "bullet": ParagraphStyle("Bullet", parent=styles["Normal"], fontName="Helvetica",
                                 fontSize=9.5, leading=12, spaceAfter=1),
        "entry_title": ParagraphStyle("EntryTitle", parent=styles["Normal"], fontName="Helvetica",
                                      fontSize=10.5, spaceBefore=5, spaceAfter=0, leading=12),
        "entry_dates": ParagraphStyle("EntryDates", parent=styles["Normal"], fontName="Helvetica",
                                      fontSize=9, textColor=PDF_GRAY, spaceAfter=2, leading=11),
        "letter": ParagraphStyle("Letter", parent=styles["Normal"], fontName="Helvetica",
                                 fontSize=10.5, leading=14, spaceAfter=7),
        "small_italic": ParagraphStyle("SmallItalic", parent=styles["Normal"],
                                       fontName="Helvetica-Oblique", fontSize=8.5,
                                       textColor=PDF_GRAY, spaceAfter=2, leading=11),
    }


def _render(path, story, top=1.5, bottom=1.3, left=2.0, right=2.0):
    doc = BaseDocTemplate(path, pagesize=letter,
                          topMargin=top * cm, bottomMargin=bottom * cm,
                          leftMargin=left * cm, rightMargin=right * cm)
    frame = Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id="main")
    doc.addPageTemplates([PageTemplate(id="all", frames=[frame])])
    doc.build(story)
    print(f"OK {os.path.basename(path)}")


def build_pdf_cv(story_only=False):
    S = _styles()

    def rule():
        return HRFlowable(width="100%", thickness=0.6, color=PDF_BLUE, spaceBefore=1, spaceAfter=3)

    def heading(text):
        return [Paragraph(_esc(text.upper()), S["heading"]), rule()]

    def bullets(items):
        return ListFlowable(
            [ListItem(Paragraph(_bold_md(t), S["bullet"]), leftIndent=10, value="•") for t in items],
            bulletType="bullet", bulletColor=PDF_BLUE, leftIndent=12, bulletFontSize=8,
        )

    def entry(dates, title, org, location, bl):
        loc = f", {_esc(location)}" if location else ""
        head = (f'<b>{_esc(title)}</b>'
                f'<font color="#555555">&nbsp;&nbsp;—&nbsp;&nbsp;{_esc(org)}{loc}</font>')
        header = [Paragraph(head, S["entry_title"]), Paragraph(_esc(dates), S["entry_dates"])]
        if not bl:
            return [KeepTogether(header)]
        if len(bl) <= 2:
            return [KeepTogether(header + [bullets(bl)])]
        out = [KeepTogether(header + [bullets(bl[:1])])]
        if len(bl) > 1:
            out.append(bullets(bl[1:]))
        return out

    story = [Paragraph(_esc(NAME), S["name"]),
             Paragraph(_esc(CV_CONTACT), S["contact"]),
             Paragraph(_esc(PROFILE), S["body"])]
    story += heading("Core Competencies")
    story.append(bullets([f"**{l}**{r}" for l, r in COMPETENCIES]))
    story += heading("Professional Experience")
    for e in EXPERIENCE:
        story += entry(*e)
    story += heading("Additional Experience")
    story.append(bullets(ADDITIONAL))
    story += heading("Education")
    for e in EDUCATION:
        story += entry(*e)
    story += heading("Languages")
    story.append(Paragraph("English (native), French (bilingual, professional working proficiency).", S["body"]))
    story += heading("Selected Publications")
    story.append(bullets(PUBLICATIONS))
    story.append(Paragraph(_esc(PUB_FOOT), S["small_italic"]))
    story += heading("References")
    story.append(Paragraph("Available upon request.", S["body"]))

    if story_only:
        return story
    _render("docx/cv_alliance_curation.pdf", story)


def build_pdf_cover(story_only=False):
    S = _styles()
    story = [Paragraph(_esc(NAME), S["name"]),
             Paragraph(_esc(COVER_CONTACT), S["contact"]),
             Spacer(1, 6),
             Paragraph(_esc(COVER_DATE), S["body"]),
             Spacer(1, 8),
             Paragraph(f"<b>{_esc(COVER_SUBJECT)}</b>", S["body"]),
             Spacer(1, 4),
             Paragraph(_esc(COVER_SALUTATION), S["letter"])]
    for i, para in enumerate(COVER_PARAS):
        story.append(Paragraph(_esc(para), S["letter"]))
        if i == 1:
            story.append(ListFlowable(
                [ListItem(Paragraph(_esc(t), S["bullet"]), leftIndent=10, value="•") for t in COVER_BULLETS],
                bulletType="bullet", bulletColor=PDF_BLUE, leftIndent=12, bulletFontSize=8,
            ))
            story.append(Spacer(1, 4))
    story += [Spacer(1, 2),
              KeepTogether([Paragraph(_esc(COVER_CLOSING), S["body"]),
                            Spacer(1, 10),
                            Paragraph(_esc(NAME), S["body"])])]
    if story_only:
        return story
    _render("docx/cover_alliance_curation.pdf", story, top=1.8, bottom=1.1, left=2.5, right=2.5)


# ==================================================================
# COMBINED (cover letter + CV in one document, as the posting requests)
# ==================================================================
def build_docx_combined():
    from docx.enum.section import WD_SECTION
    doc = Document(); _margins(doc, top=2.0, bottom=2.0, left=2.5, right=2.5)
    build_docx_cover(doc)
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    sec.top_margin, sec.bottom_margin = Cm(1.5), Cm(1.5)
    sec.left_margin, sec.right_margin = Cm(2.0), Cm(2.0)
    build_docx_cv(doc)
    doc.save("docx/application_alliance_curation.docx")
    print("OK application_alliance_curation.docx")


def build_pdf_combined():
    from reportlab.platypus import NextPageTemplate, PageBreak
    cover = build_pdf_cover(story_only=True)
    cv = build_pdf_cv(story_only=True)
    path = "docx/application_alliance_curation.pdf"
    doc = BaseDocTemplate(path, pagesize=letter)
    f_cover = Frame(2.5 * cm, 1.1 * cm, letter[0] - 5.0 * cm, letter[1] - 2.9 * cm, id="cover")
    f_cv = Frame(2.0 * cm, 1.3 * cm, letter[0] - 4.0 * cm, letter[1] - 2.8 * cm, id="cv")
    doc.addPageTemplates([PageTemplate(id="cover", frames=[f_cover]),
                          PageTemplate(id="cv", frames=[f_cv])])
    doc.build(cover + [NextPageTemplate("cv"), PageBreak()] + cv)
    print("OK application_alliance_curation.pdf")


if __name__ == "__main__":
    build_docx_cv()
    build_docx_cover()
    build_docx_combined()
    build_pdf_cv()
    build_pdf_cover()
    build_pdf_combined()
    print()
    print("Done. Files in docx/")
