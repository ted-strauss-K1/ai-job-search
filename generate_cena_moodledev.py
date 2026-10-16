"""
Generate application docs for the CENA project (Centre d'excellence québécois
en matière d'environnement numérique d'apprentissage), a Moodle consortium of
Quebec universities and colleges led from the Université de Sherbrooke.
Role: Programmeur développeur Web sénior (Moodle), 35 h/week, Montreal,
hybrid (1 day/week on site, to be confirmed). Sent by a recruiter on LinkedIn.
Outputs (in docx/):
  - cv_cena_moodledev.docx     / .pdf
  - cover_cena_moodledev.docx  / .pdf

Content is defined ONCE below, then rendered to both docx (python-docx) and
pdf (reportlab) so formats never drift.

Language: written in French. The posting is in French, the employer is a
Quebec university consortium, and it requires "maîtrise du français et de
l'anglais", so a French application doubles as proof of the language bar.

Ordering: experience is strictly reverse chronological (by start date), per
Ted's request. Pre-2012 research software roles sit in "Autres expériences"
to support the posting's 8+ years seniority bar without crowding page 1.

Honesty notes:
  - Linux experience is CentOS (a Red Hat derivative), Ubuntu and Debian, not
    RHEL itself; stated as "CentOS (famille Red Hat)" rather than "Red Hat".
  - Claude Code is named as an AI dev tool in competencies only; it is not
    attributed to any specific employer.

Company facts: CENA's full name and its Université de Sherbrooke home were
verified via WebSearch. The cover letter otherwise only restates the
project's goals as described in the posting itself (Moodle mutualization,
community of practice); no third-party claims.
"""
import os

os.makedirs("docx", exist_ok=True)

# ------------------------------------------------------------------
# SHARED CONTENT
# ------------------------------------------------------------------

NAME = "Ted Strauss"
CV_CONTACT = "Montréal (Québec) Canada  |  514-475-2374  |  ted.strauss@gmail.com  |  linkedin.com/in/tedstrauss-63b01821"
COVER_CONTACT = "514-475-2374  |  ted.strauss@gmail.com  |  linkedin.com/in/tedstrauss-63b01821"

PROFILE = (
    "Développeur Web et analyste Moodle cumulant plus de 15 ans d'expérience en développement logiciel "
    "et en infrastructure numérique, dont plus de dix ans en milieu universitaire (Université McGill, "
    "University of Connecticut, Columbia University). Développe des fonctionnalités Moodle en PHP, MySQL "
    "et JavaScript, migre des plateformes d'apprentissage existantes vers Moodle et fait le lien entre "
    "les besoins des utilisateurs et les équipes de développement. A géré les ressources de données d'un "
    "grand centre de recherche de l'Université McGill, où il a conçu un portail de formation Moodle, et "
    "représenté l'établissement au sein de regroupements nationaux (Portage, CONP, RDA). Bilingue "
    "français-anglais, basé à Montréal."
)

COMPETENCIES = [
    ("Moodle",
     " : développement de fonctionnalités, migration de plateformes LMS existantes, configuration "
     "d'environnements, conception de portails de formation autonomes, analyse des besoins des utilisateurs"),
    ("Développement Web",
     " : PHP (programmation orientée objet), JavaScript, AJAX, jQuery, Bootstrap, HTML, CSS, JSON, "
     "standards Web ; CMS Drupal et WordPress"),
    ("Bases de données",
     " : MySQL, PostgreSQL, SQL, modélisation de bases de données, administration (DBA)"),
    ("Systèmes et DevOps",
     " : Linux (CentOS, famille Red Hat ; Ubuntu ; Debian), Apache, Nginx, Docker, LXD, Kubernetes, "
     "Git, LDAP, SSO"),
    ("Méthodes et documentation",
     " : Agile Scrum, documentation et transfert de connaissances (bases de connaissances, ateliers "
     "de formation), coordination entre clients et équipes techniques"),
    ("Outils d'IA",
     " : développement assisté par IA (Claude Code), évaluation et adoption d'outils d'IA en contexte "
     "de formation"),
]

# each: (dates, title, org, location, [bullets]) - strictly reverse chronological by start date
EXPERIENCE = [
    ("07/2025 - 01/2026", "Analyste Moodle", "Premergency Inc.", "Canada", [
        "Migration de plateformes d'apprentissage (LMS) existantes vers Moodle pour des clients gouvernementaux",
        "Coordination entre les clients et l'équipe de développement : analyse des besoins et arrimage avec les fonctionnalités de Moodle",
        "Respect des exigences réglementaires propres aux déploiements LMS gouvernementaux",
        "Conception de solutions de formation sur mesure pour les clients",
    ]),
    ("01/2024 - 10/2024", "Développeur full stack", "KnowledgeOne", "Montréal (QC)", [
        "Développement de fonctionnalités Moodle pour des clients actifs sur la pile LAMP (PHP, MySQL, JavaScript)",
        "Déploiement des environnements sur conteneurs LXD",
        "Évaluation et adoption d'outils d'IA en contexte de formation en entreprise",
    ]),
    ("08/2020 - 10/2021", "Coordonnateur, gestion et partage des données (temps partiel)",
     "Institut Tanenbaum pour la science ouverte, Le Neuro, Université McGill", "Montréal (QC)", [
        "Soutien au lancement d'une nouvelle unité universitaire consacrée à la science ouverte",
        "Conception et animation d'ateliers de formation sur la gestion des données pour les étudiants et le corps professoral",
        "Création d'une base de connaissances en ligne sur les concepts de la science ouverte",
    ]),
    ("09/2016 - 10/2021", "Gestionnaire des ressources de données",
     "Centre d'imagerie cérébrale McConnell, Le Neuro, Université McGill", "Montréal (QC)", [
        "Gestion des flux de données d'un des principaux centres d'imagerie cérébrale au Canada, au service de dizaines de chercheurs principaux et de centaines de chercheurs",
        "Développement et lancement de plusieurs applications Web de recherche centrées sur les données (pile LAMP)",
        "Conception d'un portail de formation autonome basé sur Moodle et lancement d'un portail OwnCloud de gestion autonome des données",
        "Administration de bases de données et de services d'authentification (LDAP, SSO) ; travail en mode Agile",
        "Développement de pipelines de données d'imagerie et de politiques pour l'acquisition quotidienne des données",
        "Représentation de l'établissement au sein de comités nationaux (Portage, CONP, Research Data Alliance)",
    ]),
    ("02/2015 - 09/2016", "Développeur logiciel",
     "Centre de neuroscience intégrative de McGill (MCIN), Université McGill", "Montréal (QC)", [
        "Programmation du pipeline d'acquisition de données de l'étude nationale CCNA sur le vieillissement et le cerveau, sur la plateforme libre LORIS (PHP, MySQL, JavaScript)",
        "Administration des serveurs Linux et de la pile LAMP ; développement en mode Agile",
    ]),
    ("04/2012 - 10/2014", "Administrateur de bases de données et développeur logiciel",
     "Photonic Knowledge", "", [
        "Développement d'un progiciel de gestion intégré (ERP) sous Drupal pour l'imagerie de carottes de forage dans le secteur de l'exploration minière, incluant la modélisation de sa base de données",
        "Administration des serveurs et de la pile LAMP ; développement en mode Agile",
    ]),
]

ADDITIONAL = [
    "Sur Place Media, Montréal - Fondateur et directeur (2017-2021) : école d'arts de quartier, 150 groupes d'ateliers, 20 enseignants ; virage vers la formation en ligne pendant la pandémie",
    "Language Mate - Développeur et concepteur UX (2007-2010)",
    "Moss Rehabilitation Hospital - Développeur de logiciels de recherche (2007-2010) : base de données Web interrogeable sur l'aphasie (publiée en 2010)",
    "University of Connecticut - Assistant de recherche et développeur logiciel (2005-2008)",
    "Columbia University - Assistant de recherche et développeur logiciel (2004-2005)",
    "Universität Bremen - Assistant de recherche et développeur logiciel (2003-2004)",
]

EDUCATION = [
    ("1998 - 2003", "B.A. (spécialisé), sciences cognitives", "Université McGill", "Montréal (QC)",
     ["Mineure en technologie musicale. Psycholinguistique, modélisation computationnelle du langage."]),
    ("2021", "Scientific Community Engagement Fundamentals (certification)", "CSCCE", "", []),
]

LANGUAGES = "Anglais (langue maternelle), français (bilingue, maîtrise professionnelle à l'écrit et à l'oral)."

PUBLICATIONS = [
    "Poline, J.-B., et al. (2023). Data and tools integration in the Canadian open neuroscience platform. Scientific Data, 10(1), 189.",
    "Mirman, D., Strauss, T. J., et al. (2010). A large, searchable, web-based database of aphasic performance on picture naming and other tests of cognitive function. Cognitive Neuropsychology, 27(6), 495-504.",
    "Strauss, T. J., Harris, H. D., & Magnuson, J. S. (2007). jTRACE: A reimplementation and extension of the TRACE model of speech perception and spoken word recognition. Behavior Research Methods, 39(1), 19-30.",
]
PUB_FOOT = "9 publications | 630 citations | indice h : 9 | Liste complète : scholar.google.com/citations"

H_COMP = "Compétences clés"
H_EXP = "Expérience professionnelle"
H_ADD = "Autres expériences"
H_EDU = "Formation"
H_LANG = "Langues"
H_PUB = "Publications choisies"

# ------------------------------------------------------------------
# COVER LETTER
# ------------------------------------------------------------------

COVER_DATE = "Montréal, le 21 septembre 2026"
COVER_SUBJECT = "Objet : Candidature au poste de programmeur développeur Web sénior, projet CENA"
COVER_SALUTATION = "Madame, Monsieur,"
COVER_PARAS = [
    "Je vous soumets ma candidature au poste de programmeur développeur Web sénior pour le projet CENA. "
    "Mutualiser dans un même environnement Moodle les développements technopédagogiques de plusieurs "
    "universités et collèges, tout en bâtissant une communauté de pratique autour de ces développements, "
    "rejoint directement les deux volets de mon parcours : le développement Web sur Moodle et la mise en "
    "place d'infrastructures numériques partagées dans le milieu universitaire.",

    "Chez KnowledgeOne, j'ai développé des fonctionnalités Moodle pour des clients actifs sur la pile LAMP "
    "(PHP, MySQL, JavaScript), avec des environnements déployés sur conteneurs LXD. Chez Premergency, j'ai "
    "ensuite migré des plateformes d'apprentissage existantes vers Moodle pour des clients gouvernementaux, "
    "en faisant le lien entre leurs besoins et l'équipe de développement, soit précisément le premier défi "
    "décrit dans votre affichage. Mon expérience de développement remonte toutefois bien avant Moodle :",

    "J'ai passé plus de six ans à l'Université McGill. Au Centre d'imagerie cérébrale McConnell, j'ai géré "
    "les ressources de données d'un centre au service de centaines de chercheurs, développé plusieurs "
    "applications Web de recherche et conçu un portail de formation autonome basé sur Moodle. J'y ai aussi "
    "représenté l'établissement au sein de Portage, du CONP et de la Research Data Alliance. La gouvernance "
    "partagée entre établissements, la priorisation de développements communs et l'animation d'une "
    "communauté de pratique sont donc des réalités que je connais de l'intérieur.",

    "Bilingue, basé à Montréal, je suis disponible immédiatement et la formule hybride proposée me convient. "
    "Je serais heureux de discuter avec vous de ma contribution au projet CENA.",

    "Je vous prie d'agréer, Madame, Monsieur, l'expression de mes salutations distinguées.",
]
COVER_BULLETS = [
    "Programmation du pipeline d'acquisition de données de l'étude nationale CCNA sur la plateforme libre LORIS (PHP, MySQL, JavaScript), en mode Agile",
    "Développement d'un ERP sous Drupal, incluant la modélisation de sa base de données",
    "Administration de serveurs Linux : CentOS (famille Red Hat), Ubuntu et Debian",
]

# ==================================================================
# DOCX RENDERING
# ==================================================================
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BLUE = RGBColor(0x21, 0x6F, 0xAA)
BLACK = RGBColor(0x00, 0x00, 0x00)
GRAY = RGBColor(0x55, 0x55, 0x55)


def _margins(doc, top=1.5, bottom=1.5, left=2.0, right=2.0):
    for s in doc.sections:
        s.top_margin, s.bottom_margin = Cm(top), Cm(bottom)
        s.left_margin, s.right_margin = Cm(left), Cm(right)


def _name_header(doc, name, contact):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(name); r.bold = True; r.font.size = Pt(28); r.font.color.rgb = BLUE
    p2 = doc.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(contact); r2.font.size = Pt(10); r2.font.color.rgb = GRAY
    p2.paragraph_format.space_after = Pt(6)


def _heading(doc, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title.upper()); r.bold = True; r.font.size = Pt(11); r.font.color.rgb = BLUE
    pPr = p._p.get_or_add_pPr(); pBdr = OxmlElement('w:pBdr'); bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single'); bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1'); bottom.set(qn('w:color'), '216FAA')
    pBdr.append(bottom); pPr.append(pBdr)


def _entry(doc, dates, title, org, location, bullets):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    r1 = p.add_run(title); r1.bold = True; r1.font.size = Pt(10.5); r1.font.color.rgb = BLACK
    r2 = p.add_run(f"  |  {org}"); r2.font.size = Pt(10.5); r2.font.color.rgb = GRAY
    if location:
        r3 = p.add_run(f", {location}"); r3.font.size = Pt(10.5); r3.font.color.rgb = GRAY
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_before = Pt(0); p2.paragraph_format.space_after = Pt(2)
    p2.paragraph_format.keep_with_next = True
    r4 = p2.add_run(dates); r4.italic = True; r4.font.size = Pt(9.5); r4.font.color.rgb = GRAY
    for b in bullets:
        _bullet(doc, b)


def _bullet(doc, text, indent=0.25):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(1); p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.left_indent = Inches(indent)
    for i, part in enumerate(text.split('**')):
        if not part:
            continue
        r = p.add_run(part); r.font.size = Pt(10); r.bold = (i % 2 == 1)


def _para(doc, text, size=10, italic=False, color=None, space_after=4):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(text); r.font.size = Pt(size); r.italic = italic
    if color:
        r.font.color.rgb = color
    return p


def build_docx_cv():
    doc = Document(); _margins(doc)
    _name_header(doc, NAME, CV_CONTACT)
    _para(doc, PROFILE, size=10, space_after=6)
    _heading(doc, H_COMP)
    for label, rest in COMPETENCIES:
        _bullet(doc, f"**{label}**{rest}")
    _heading(doc, H_EXP)
    for e in EXPERIENCE:
        _entry(doc, *e)
    _heading(doc, H_ADD)
    for a in ADDITIONAL:
        _bullet(doc, a)
    _heading(doc, H_EDU)
    for e in EDUCATION:
        _entry(doc, *e)
    _heading(doc, H_LANG)
    _para(doc, LANGUAGES, size=10)
    _heading(doc, H_PUB)
    for pub in PUBLICATIONS:
        _bullet(doc, pub)
    _para(doc, PUB_FOOT, size=9, italic=True, color=GRAY)
    doc.save("docx/cv_cena_moodledev.docx")
    print("OK cv_cena_moodledev.docx")


def build_docx_cover():
    doc = Document(); _margins(doc, top=2.0, bottom=2.0, left=2.5, right=2.5)
    _name_header(doc, NAME, COVER_CONTACT)
    _para(doc, COVER_DATE, size=10, space_after=12)
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(6)
    r = p.add_run(COVER_SUBJECT); r.bold = True; r.font.size = Pt(10)
    _para(doc, COVER_SALUTATION, size=10, space_after=6)
    for i, para in enumerate(COVER_PARAS):
        pp = doc.add_paragraph(); pp.paragraph_format.space_after = Pt(8)
        rr = pp.add_run(para); rr.font.size = Pt(10)
        if i == 1:
            for b in COVER_BULLETS:
                _bullet(doc, b)
    _para(doc, "", size=10, space_after=18)
    _para(doc, NAME, size=10)
    doc.save("docx/cover_cena_moodledev.docx")
    print("OK cover_cena_moodledev.docx")


# ==================================================================
# PDF RENDERING
# ==================================================================
from reportlab.lib.pagesizes import letter
from reportlab.lib.units import cm
from reportlab.lib.colors import HexColor
from reportlab.lib.enums import TA_CENTER
from reportlab.platypus import (
    BaseDocTemplate, PageTemplate, Frame, Paragraph, Spacer,
    HRFlowable, ListFlowable, ListItem, KeepTogether,
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

PDF_BLUE = HexColor("#216FAA")
PDF_GRAY = HexColor("#555555")


def _esc(t):
    return t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _bold_md(t):
    t = _esc(t)
    parts = t.split("**")
    out = ""
    for i, p in enumerate(parts):
        out += f"<b>{p}</b>" if i % 2 == 1 else p
    return out


def _styles():
    styles = getSampleStyleSheet()
    return {
        "name": ParagraphStyle("Name", parent=styles["Normal"], fontName="Helvetica-Bold",
                               fontSize=22, textColor=PDF_BLUE, alignment=TA_CENTER, spaceAfter=2, leading=26),
        "contact": ParagraphStyle("Contact", parent=styles["Normal"], fontName="Helvetica",
                                  fontSize=9, textColor=PDF_GRAY, alignment=TA_CENTER, spaceAfter=8, leading=12),
        "heading": ParagraphStyle("Heading", parent=styles["Normal"], fontName="Helvetica-Bold",
                                  fontSize=11, textColor=PDF_BLUE, spaceBefore=8, spaceAfter=2, leading=13),
        "body": ParagraphStyle("Body", parent=styles["Normal"], fontName="Helvetica", fontSize=10,
                               leading=13, spaceAfter=3),
        "bullet": ParagraphStyle("Bullet", parent=styles["Normal"], fontName="Helvetica",
                                 fontSize=9.5, leading=12, spaceAfter=1),
        "entry_title": ParagraphStyle("EntryTitle", parent=styles["Normal"], fontName="Helvetica",
                                      fontSize=10.5, spaceBefore=5, spaceAfter=0, leading=12),
        "entry_dates": ParagraphStyle("EntryDates", parent=styles["Normal"], fontName="Helvetica",
                                      fontSize=9, textColor=PDF_GRAY, spaceAfter=2, leading=11),
        "letter": ParagraphStyle("Letter", parent=styles["Normal"], fontName="Helvetica",
                                 fontSize=10.5, leading=14, spaceAfter=7),
        "small_italic": ParagraphStyle("SmallItalic", parent=styles["Normal"],
                                       fontName="Helvetica-Oblique", fontSize=8.5,
                                       textColor=PDF_GRAY, spaceAfter=2, leading=11),
    }


def _render(path, story, top=1.5, bottom=1.3, left=2.0, right=2.0):
    doc = BaseDocTemplate(path, pagesize=letter,
                          topMargin=top * cm, bottomMargin=bottom * cm,
                          leftMargin=left * cm, rightMargin=right * cm)
    frame = Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id="main")
    doc.addPageTemplates([PageTemplate(id="all", frames=[frame])])
    doc.build(story)
    print(f"OK {os.path.basename(path)}")


def build_pdf_cv():
    S = _styles()

    def rule():
        return HRFlowable(width="100%", thickness=0.6, color=PDF_BLUE, spaceBefore=1, spaceAfter=3)

    def heading(text):
        return [Paragraph(_esc(text.upper()), S["heading"]), rule()]

    def bullets(items):
        return ListFlowable(
            [ListItem(Paragraph(_bold_md(t), S["bullet"]), leftIndent=10, value="•") for t in items],
            bulletType="bullet", bulletColor=PDF_BLUE, leftIndent=12, bulletFontSize=8,
        )

    def entry(dates, title, org, location, bl):
        loc = f", {_esc(location)}" if location else ""
        head = (f'<b>{_esc(title)}</b>'
                f'<font color="#555555">&nbsp;&nbsp;|&nbsp;&nbsp;{_esc(org)}{loc}</font>')
        header = [Paragraph(head, S["entry_title"]), Paragraph(_esc(dates), S["entry_dates"])]
        if not bl:
            return [KeepTogether(header)]
        if len(bl) <= 2:
            return [KeepTogether(header + [bullets(bl)])]
        out = [KeepTogether(header + [bullets(bl[:1])])]
        if len(bl) > 1:
            out.append(bullets(bl[1:]))
        return out

    story = [Paragraph(_esc(NAME), S["name"]),
             Paragraph(_esc(CV_CONTACT), S["contact"]),
             Paragraph(_esc(PROFILE), S["body"])]
    story += heading(H_COMP)
    story.append(bullets([f"**{l}**{r}" for l, r in COMPETENCIES]))
    story += heading(H_EXP)
    for e in EXPERIENCE:
        story += entry(*e)
    story += heading(H_ADD)
    story.append(bullets(ADDITIONAL))
    story += heading(H_EDU)
    for e in EDUCATION:
        story += entry(*e)
    story += heading(H_LANG)
    story.append(Paragraph(_esc(LANGUAGES), S["body"]))
    story += heading(H_PUB)
    story.append(bullets(PUBLICATIONS))
    story.append(Paragraph(_esc(PUB_FOOT), S["small_italic"]))

    _render("docx/cv_cena_moodledev.pdf", story)


def build_pdf_cover():
    S = _styles()
    story = [Paragraph(_esc(NAME), S["name"]),
             Paragraph(_esc(COVER_CONTACT), S["contact"]),
             Spacer(1, 6),
             Paragraph(_esc(COVER_DATE), S["body"]),
             Spacer(1, 8),
             Paragraph(f"<b>{_esc(COVER_SUBJECT)}</b>", S["body"]),
             Spacer(1, 4),
             Paragraph(_esc(COVER_SALUTATION), S["letter"])]
    for i, para in enumerate(COVER_PARAS):
        story.append(Paragraph(_esc(para), S["letter"]))
        if i == 1:
            story.append(ListFlowable(
                [ListItem(Paragraph(_esc(t), S["bullet"]), leftIndent=10, value="•") for t in COVER_BULLETS],
                bulletType="bullet", bulletColor=PDF_BLUE, leftIndent=12, bulletFontSize=8,
            ))
            story.append(Spacer(1, 4))
    story += [Spacer(1, 16), Paragraph(_esc(NAME), S["body"])]
    _render("docx/cover_cena_moodledev.pdf", story, top=1.8, bottom=1.1, left=2.5, right=2.5)


if __name__ == "__main__":
    build_docx_cv()
    build_docx_cover()
    build_pdf_cv()
    build_pdf_cover()
    print("\nDone. Files in docx/")
