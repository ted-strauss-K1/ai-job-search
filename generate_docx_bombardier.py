"""
Generate DOCX CV and cover letter for Bombardier Learning Business Partner application.
Job ID: 13164 — Organizational Learning, Dorval QC (hybrid)
"""
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

os.makedirs("docx", exist_ok=True)

BLUE = RGBColor(0x21, 0x6F, 0xAA)
BLACK = RGBColor(0x00, 0x00, 0x00)
GRAY = RGBColor(0x55, 0x55, 0x55)


def set_margins(doc, top=1.5, bottom=1.5, left=2.0, right=2.0):
    for section in doc.sections:
        section.top_margin = Cm(top)
        section.bottom_margin = Cm(bottom)
        section.left_margin = Cm(left)
        section.right_margin = Cm(right)


def add_name_header(doc, name, contact):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(name)
    run.bold = True
    run.font.size = Pt(28)
    run.font.color.rgb = BLUE

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run2 = p2.add_run(contact)
    run2.font.size = Pt(10)
    run2.font.color.rgb = GRAY
    p2.paragraph_format.space_after = Pt(6)


def add_section_heading(doc, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(title.upper())
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = BLUE
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), '216FAA')
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_job_entry(doc, dates, title, org, location, bullets):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(0)
    r1 = p.add_run(title)
    r1.bold = True
    r1.font.size = Pt(10.5)
    r1.font.color.rgb = BLACK
    r2 = p.add_run(f"  —  {org}")
    r2.font.size = Pt(10.5)
    r2.font.color.rgb = GRAY
    if location:
        r3 = p.add_run(f", {location}")
        r3.font.size = Pt(10.5)
        r3.font.color.rgb = GRAY

    p2 = doc.add_paragraph()
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(2)
    r4 = p2.add_run(dates)
    r4.italic = True
    r4.font.size = Pt(9.5)
    r4.font.color.rgb = GRAY

    for b in bullets:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_before = Pt(1)
        bp.paragraph_format.space_after = Pt(1)
        bp.paragraph_format.left_indent = Inches(0.25)
        parts = b.split('**')
        for i, part in enumerate(parts):
            if not part:
                continue
            r = bp.add_run(part)
            r.font.size = Pt(10)
            r.bold = (i % 2 == 1)


def add_bullet(doc, text, indent=0.25):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.left_indent = Inches(indent)
    parts = text.split('**')
    for i, part in enumerate(parts):
        if not part:
            continue
        r = p.add_run(part)
        r.font.size = Pt(10)
        r.bold = (i % 2 == 1)


def add_para(doc, text, size=10, italic=False, color=None, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.italic = italic
    if color:
        r.font.color.rgb = color
    return p


# ============================================================
# CV — Bombardier Learning Business Partner
# ============================================================

def cv_bombardier():
    doc = Document()
    set_margins(doc)
    add_name_header(doc, "Ted Strauss",
        "Montreal, Quebec, Canada  |  514-475-2374  |  ted.strauss@gmail.com  |  linkedin.com/in/tedstrauss-63b01821")

    add_para(doc,
        "Learning & development specialist and LMS expert with experience designing training programs, "
        "managing Moodle platforms, and coordinating L&D initiatives across government, academic, and enterprise "
        "environments. Served as liaison between clients, management, and technical teams to analyze training needs "
        "and deploy solutions in compliance-driven contexts. Brings direct experience managing LMS systems, "
        "facilitating bilingual training programs, and driving organizational adoption through change management. "
        "Bilingual (EN/FR), based in Montreal.",
        size=10, space_after=6)

    add_section_heading(doc, "Core Competencies")
    add_bullet(doc, "**LMS platform management**: Moodle administration, development, migration, and deployment across government, higher education, and enterprise environments; platform ownership from configuration through user adoption; LXD, EdX")
    add_bullet(doc, "**Training needs analysis & L&D design**: analyzed training requirements and designed blended learning solutions for researchers, government clients, faculty, and professional teams; workshops, e-learning modules, knowledge bases, and self-serve portals")
    add_bullet(doc, "**Cross-functional liaison & coordination**: served as the bridge between clients/business units and development/technical teams; translated business requirements into training specifications; managed multiple concurrent projects under Agile delivery")
    add_bullet(doc, "**Change management & adoption**: embedded learning programs into organizational workflows through stakeholder engagement rather than compliance enforcement; facilitated adoption across complex multi-stakeholder environments")
    add_bullet(doc, "**Regulatory compliance**: maintained compliance requirements for government LMS deployments; developed data acquisition and handling policies under national research standards; contributed to federal-level policy frameworks through Portage, CONP, and RDA")
    add_bullet(doc, "**Bilingual facilitation**: designed and delivered programs in French and English; managed bilingual LMS environments for government clients; bilingual content development across diverse audiences")

    add_section_heading(doc, "Professional Experience")
    add_job_entry(doc, "07/2025 – 01/2026", "Moodle Analyst", "Premergency Inc.", "Canada", [
        "Served as liaison between government clients and development teams, translating client training requirements into LMS configurations and ensuring regulatory compliance throughout",
        "Analyzed training needs and designed bespoke learning & development solutions tailored to each client's operational context",
        "Managed full migrations of legacy LMS systems to Moodle, including platform configuration, data migration, and user onboarding",
    ])
    add_job_entry(doc, "01/2024 – 10/2024", "Full Stack Developer", "KnowledgeOne", "Montreal, QC", [
        "Developed and maintained LMS features for active clients on the eConcordia Moodle platform (LAMP++, LXD)",
        "Evaluated and adopted AI tools in enterprise learning contexts; assessed usability, adoption pathways, and integration with existing LMS workflows",
    ])
    add_job_entry(doc, "09/2016 – 10/2021", "Manager of Data Resources", "McConnell Brain Imaging Centre, The Neuro, McGill University", "Montreal, QC", [
        "Designed and deployed a Moodle-based self-serve education portal for researcher onboarding and training across dozens of research teams",
        "Delivered bilingual training workshops for researchers, faculty, and non-specialist teams on platform use, data practices, and compliance requirements",
        "Coordinated cross-functional delivery involving IT, research operations, and institutional administration; managed competing priorities across multiple concurrent initiatives",
        "Contributed to national research policy through Portage, CONP, and RDA committee work; developed frameworks adopted at the federal level",
    ])
    add_job_entry(doc, "08/2020 – 10/2021", "Data Management & Sharing Coordinator (p/t)", "Tanenbaum Open Science Institute, The Neuro, McGill University", "Montreal, QC", [
        "Designed and administered training workshops for students, faculty, and research staff on data management and open science practices",
        "Built a public-facing online knowledge base; developed bilingual content for diverse audiences across the research community",
    ])
    add_job_entry(doc, "04/2017 – 08/2021", "Founder and Director", "Sur Place Media", "Montreal, QC", [
        "Founded and operated a neighbourhood arts school; managed 20 instructors across 150 workshop groups",
        "Led a successful pivot to fully virtual delivery during the COVID-19 pandemic, redesigning all programming for online engagement",
    ])
    add_job_entry(doc, "02/2015 – 09/2016", "Software Developer", "McGill Centre for Integrative Neuroscience, McGill University", "Montreal, QC", [
        "Programmed the data acquisition pipeline for the national CCNA longitudinal dementia study (PHP, MySQL, JavaScript); contributed to national research infrastructure",
    ])

    add_section_heading(doc, "Additional Experience")
    add_bullet(doc, "Photonic Knowledge — Database administrator & software developer (2012–2014)")
    add_bullet(doc, "Language Mate — Developer and UX designer (2007–2010)")
    add_bullet(doc, "Moss Rehabilitation Hospital — Research software developer (2007–2010)")
    add_bullet(doc, "University of Connecticut & Columbia University — Research assistant & software developer (2004–2008)")

    add_section_heading(doc, "Education")
    add_job_entry(doc, "1998–2003", "B.A. (Honours), Cognitive Science", "McGill University", "Montreal, QC", [
        "Minor: Music Technology. Topics: Psycholinguistics, cognitive science, computational modeling of language.",
    ])
    add_job_entry(doc, "2021", "Scientific Community Engagement Fundamentals", "CSCCE", "", [])

    add_section_heading(doc, "Languages")
    add_para(doc, "English (native), French (bilingual).", size=10)

    add_section_heading(doc, "References")
    add_para(doc, "Available upon request.", size=10)

    doc.save("docx/cv_bombardier_lbp.docx")
    print("OK cv_bombardier_lbp.docx")


# ============================================================
# Cover Letter — Bombardier Learning Business Partner
# ============================================================

def cover_bombardier():
    doc = Document()
    set_margins(doc, top=2.0, bottom=2.0, left=2.5, right=2.5)
    add_name_header(doc, "Ted Strauss",
        "514-475-2374  |  ted.strauss@gmail.com  |  linkedin.com/in/tedstrauss-63b01821")

    add_para(doc, "June 22, 2026", size=10, space_after=12)
    add_para(doc, "Dear Bombardier Hiring Team,", size=10, space_after=6)

    paragraphs = [
        (
            "I am applying for the Learning Business Partner position on Bombardier's Organizational Learning team "
            "(Job ID 13164). Training design, LMS management, and cross-functional coordination are where my "
            "experience runs deepest, and the role's focus on translating business priorities into learning solutions "
            "matches the work I have been doing across government, academic, and enterprise environments."
        ),
        (
            "At Premergency Inc., I served as the coordinator between government clients and development teams for "
            "Moodle LMS deployments: analyzing client training requirements, translating them into platform "
            "configurations, managing regulatory compliance throughout, and overseeing the full handoff from "
            "migration to user onboarding. That experience maps directly to the Learning Business Partner "
            "description: acting as a liaison between the business and the learning function, monitoring performance "
            "indicators, and designing solutions that satisfy both organizational priorities and compliance "
            "obligations. At KnowledgeOne, I developed and maintained LMS features for the eConcordia platform and "
            "evaluated AI tools for enterprise learning adoption. At McGill's McConnell Brain Imaging Centre, I "
            "designed a Moodle-based education portal and delivered bilingual training workshops for research teams "
            "on platform use and data compliance requirements."
        ),
        (
            "I am bilingual in French and English, based in Montreal, and familiar with the demands of "
            "compliance-driven training environments. Bombardier's commitment to developing industry talent through "
            "its Organizational Learning programs is exactly the kind of strategic L&D function I am motivated to "
            "contribute to."
        ),
        (
            "I would welcome the opportunity to bring this combination of LMS expertise, training design, and "
            "cross-functional coordination to Bombardier. Thank you for your consideration."
        ),
    ]

    for para in paragraphs:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(8)
        r = p.add_run(para)
        r.font.size = Pt(10)

    add_para(doc, "Kind regards,", size=10, space_after=24)
    add_para(doc, "Ted Strauss", size=10)

    doc.save("docx/cover_bombardier_lbp.docx")
    print("OK cover_bombardier_lbp.docx")


cv_bombardier()
cover_bombardier()
print("\nDone. Files in docx/")
