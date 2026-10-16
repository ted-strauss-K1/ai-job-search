"""
Shared renderer for the combined application files (cover letter page + CV pages).
A per-application generator builds a `cfg` dict and calls build_docx(cfg) / build_pdf(cfg).

cfg keys:
  out        output path without extension (e.g. "docx/brp_datagov_advisor_application")
  cover      dict(contact, date, subject, salutation, paras, closing)
  cv         dict(profile, competencies, experience, additional, education, skills, publications)
             or None for a letter-only file
Content is defined once in the generator, then rendered to docx (python-docx) and pdf (reportlab).
"""
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from reportlab.lib.pagesizes import letter
from reportlab.lib.units import cm
from reportlab.lib.colors import HexColor
from reportlab.lib.enums import TA_CENTER
from reportlab.platypus import (
    BaseDocTemplate, PageTemplate, Frame, Paragraph, Spacer, PageBreak,
    HRFlowable, ListFlowable, ListItem, KeepTogether, CondPageBreak,
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

NAME = "Ted Strauss"
CV_CONTACT = "Montreal, Quebec, Canada  |  514-475-2374  |  ted.strauss@gmail.com  |  linkedin.com/in/tedstrauss-63b01821"
PUB_FOOT = "9 publications | citation count 630 | h-index 9 | Full list: scholar.google.com/citations"
LANG_LINE = "English (native), French (bilingual, professional working proficiency)."

BLUE = RGBColor(0x21, 0x6F, 0xAA)
GRAY = RGBColor(0x55, 0x55, 0x55)
PDF_BLUE = HexColor("#216FAA")
PDF_GRAY = HexColor("#555555")


# ---------------------------- DOCX ----------------------------
def _name_header(doc, contact):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(NAME); r.bold = True; r.font.size = Pt(26); r.font.color.rgb = BLUE
    p2 = doc.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(contact); r2.font.size = Pt(9.5); r2.font.color.rgb = GRAY
    p2.paragraph_format.space_after = Pt(6)


def _heading(doc, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(title.upper()); r.bold = True; r.font.size = Pt(11); r.font.color.rgb = BLUE
    pPr = p._p.get_or_add_pPr(); pBdr = OxmlElement('w:pBdr'); bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single'); bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1'); bottom.set(qn('w:color'), '216FAA')
    pBdr.append(bottom); pPr.append(pBdr)


def _bullet(doc, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(1); p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.left_indent = Inches(0.25)
    for i, part in enumerate(text.split('**')):
        if part:
            r = p.add_run(part); r.font.size = Pt(10); r.bold = (i % 2 == 1)


def _entry(doc, dates, title, org, location, bullets):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    r1 = p.add_run(title); r1.bold = True; r1.font.size = Pt(10.5)
    r2 = p.add_run(f" | {org}"); r2.font.size = Pt(10.5); r2.font.color.rgb = GRAY
    if location:
        r3 = p.add_run(f", {location}"); r3.font.size = Pt(10.5); r3.font.color.rgb = GRAY
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_before = Pt(0); p2.paragraph_format.space_after = Pt(2)
    p2.paragraph_format.keep_with_next = bool(bullets)
    r4 = p2.add_run(dates); r4.italic = True; r4.font.size = Pt(9.5); r4.font.color.rgb = GRAY
    for b in bullets:
        _bullet(doc, b)


def _para(doc, text, size=10, italic=False, color=None, space_after=4):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(text); r.font.size = Pt(size); r.italic = italic
    if color:
        r.font.color.rgb = color


def build_docx(cfg):
    c, cv = cfg["cover"], cfg["cv"]
    doc = Document()
    for s in doc.sections:
        s.top_margin, s.bottom_margin = Cm(1.6), Cm(1.5)
        s.left_margin, s.right_margin = Cm(2.2), Cm(2.2)
    _name_header(doc, c["contact"])
    _para(doc, c["date"], size=10.5, space_after=10)
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(6)
    r = p.add_run(c["subject"]); r.bold = True; r.font.size = Pt(10.5)
    _para(doc, c["salutation"], size=10.5, space_after=6)
    for para in c["paras"]:
        pp = doc.add_paragraph(); pp.paragraph_format.space_after = Pt(7)
        pp.add_run(para).font.size = Pt(10.5)
    _para(doc, c["closing"], size=10.5, space_after=14)
    _para(doc, NAME, size=10.5)
    if cv:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        _name_header(doc, CV_CONTACT)
        _para(doc, cv["profile"], size=10, space_after=6)
        _heading(doc, "Core Competencies")
        for label, rest in cv["competencies"]:
            _bullet(doc, f"**{label}**{rest}")
        _heading(doc, "Relevant Experience")
        for e in cv["experience"]:
            _entry(doc, *e)
        _heading(doc, "Additional Experience")
        for a in cv["additional"]:
            _bullet(doc, a)
        _heading(doc, "Education")
        for e in cv["education"]:
            _entry(doc, *e)
        _heading(doc, "Technical Skills")
        _para(doc, cv["skills"], size=10)
        _heading(doc, "Languages")
        _para(doc, LANG_LINE, size=10)
        _heading(doc, "Selected Publications")
        for pub in cv["publications"]:
            _bullet(doc, pub)
        _para(doc, PUB_FOOT, size=9, italic=True, color=GRAY)
    doc.save(cfg["out"] + ".docx")
    print("OK", cfg["out"] + ".docx")


# ---------------------------- PDF -----------------------------
def _esc(t):
    return t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _bold_md(t):
    parts = _esc(t).split("**")
    return "".join(f"<b>{p}</b>" if i % 2 == 1 else p for i, p in enumerate(parts))


def _styles():
    s = getSampleStyleSheet()
    return {
        "name": ParagraphStyle("Name", parent=s["Normal"], fontName="Helvetica-Bold",
                               fontSize=22, textColor=PDF_BLUE, alignment=TA_CENTER, spaceAfter=2, leading=26),
        "contact": ParagraphStyle("Contact", parent=s["Normal"], fontName="Helvetica",
                                  fontSize=9, textColor=PDF_GRAY, alignment=TA_CENTER, spaceAfter=8, leading=12),
        "heading": ParagraphStyle("Heading", parent=s["Normal"], fontName="Helvetica-Bold",
                                  fontSize=11, textColor=PDF_BLUE, spaceBefore=8, spaceAfter=2, leading=13),
        "body": ParagraphStyle("Body", parent=s["Normal"], fontName="Helvetica", fontSize=10,
                               leading=13, spaceAfter=3),
        "bullet": ParagraphStyle("Bullet", parent=s["Normal"], fontName="Helvetica",
                                 fontSize=9.5, leading=12, spaceAfter=1),
        "entry_title": ParagraphStyle("EntryTitle", parent=s["Normal"], fontName="Helvetica",
                                      fontSize=10.5, spaceBefore=5, spaceAfter=0, leading=12),
        "entry_dates": ParagraphStyle("EntryDates", parent=s["Normal"], fontName="Helvetica",
                                      fontSize=9, textColor=PDF_GRAY, spaceAfter=2, leading=11),
        "letter": ParagraphStyle("Letter", parent=s["Normal"], fontName="Helvetica",
                                 fontSize=10, leading=13.2, spaceAfter=6),
        "small_italic": ParagraphStyle("SmallItalic", parent=s["Normal"],
                                       fontName="Helvetica-Oblique", fontSize=8.5,
                                       textColor=PDF_GRAY, spaceAfter=2, leading=11),
    }


def _bullets(S, items, md=True):
    f = _bold_md if md else _esc
    return ListFlowable(
        [ListItem(Paragraph(f(t), S["bullet"]), leftIndent=10, value="•") for t in items],
        bulletType="bullet", bulletColor=PDF_BLUE, leftIndent=12, bulletFontSize=8,
    )


def build_pdf(cfg):
    c, cv = cfg["cover"], cfg["cv"]
    S = _styles()

    def heading(text):
        return [CondPageBreak(2.5 * cm), Paragraph(_esc(text.upper()), S["heading"]),
                HRFlowable(width="100%", thickness=0.6, color=PDF_BLUE, spaceBefore=1, spaceAfter=3)]

    def entry(dates, title, org, location, bl):
        loc = f", {_esc(location)}" if location else ""
        head = (f'<b>{_esc(title)}</b>'
                f'<font color="#555555">&nbsp;|&nbsp;{_esc(org)}{loc}</font>')
        header = [Paragraph(head, S["entry_title"]), Paragraph(_esc(dates), S["entry_dates"])]
        if not bl:
            return [KeepTogether(header)]
        if len(bl) <= 2:
            return [KeepTogether(header + [_bullets(S, bl)])]
        return [KeepTogether(header + [_bullets(S, bl[:1])]), _bullets(S, bl[1:])]

    story = [Paragraph(_esc(NAME), S["name"]), Paragraph(_esc(c["contact"]), S["contact"]),
             Spacer(1, 4), Paragraph(_esc(c["date"]), S["body"]), Spacer(1, 6),
             Paragraph(f"<b>{_esc(c['subject'])}</b>", S["body"]), Spacer(1, 3),
             Paragraph(_esc(c["salutation"]), S["letter"])]
    for para in c["paras"]:
        story.append(Paragraph(_esc(para), S["letter"]))
    story.append(KeepTogether([Paragraph(_esc(c["closing"]), S["body"]),
                               Spacer(1, 8), Paragraph(_esc(NAME), S["body"])]))
    if cv:
        story.append(PageBreak())
        story += [Paragraph(_esc(NAME), S["name"]), Paragraph(_esc(CV_CONTACT), S["contact"]),
                  Paragraph(_esc(cv["profile"]), S["body"])]
        story += heading("Core Competencies")
        story.append(_bullets(S, [f"**{l}**{r}" for l, r in cv["competencies"]]))
        story += heading("Relevant Experience")
        for e in cv["experience"]:
            story += entry(*e)
        story += heading("Additional Experience")
        story.append(_bullets(S, cv["additional"]))
        story += heading("Education")
        for e in cv["education"]:
            story += entry(*e)
        story += heading("Technical Skills")
        story.append(Paragraph(_esc(cv["skills"]), S["body"]))
        story += heading("Languages")
        story.append(Paragraph(_esc(LANG_LINE), S["body"]))
        story += heading("Selected Publications")
        story.append(_bullets(S, cv["publications"], md=False))
        story.append(Paragraph(_esc(PUB_FOOT), S["small_italic"]))
    doc = BaseDocTemplate(cfg["out"] + ".pdf", pagesize=letter,
                          topMargin=1.5 * cm, bottomMargin=1.3 * cm,
                          leftMargin=2.2 * cm, rightMargin=2.2 * cm)
    frame = Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id="main")
    doc.addPageTemplates([PageTemplate(id="all", frames=[frame])])
    doc.build(story)
    print("OK", cfg["out"] + ".pdf")


def build_all(cfg):
    build_docx(cfg)
    build_pdf(cfg)
