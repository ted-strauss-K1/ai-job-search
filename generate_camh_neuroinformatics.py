"""
Generate application docs for CAMH — Research Methods Specialist (Research Data
Integration), Krembil Centre for Neuroinformatics. Toronto (250 College St).
12-month contract. $42.67 - $57.55 hourly.
Outputs (in docx/):
  - cv_camh_neuroinformatics.docx     / .pdf
  - cover_camh_neuroinformatics.docx  / .pdf

Content is defined ONCE below, then rendered to both docx (python-docx) and
pdf (reportlab) so the two formats never drift.

Targeting note: operational management of a neuroinformatics platform and an
institutional research data repository (BrainHealth Databank) with FAIR mentorship
in an open science environment — the closest posting yet to Ted's actual Neuro
role, so the CV leads hard with the Brain Imaging Centre and the open science
record. REDCap and XNAT/LabKey acknowledged as gaps in the letter, never claimed.
"""
import os

os.makedirs("docx", exist_ok=True)

# ------------------------------------------------------------------
# SHARED CONTENT
# ------------------------------------------------------------------

NAME = "Ted Strauss"
CV_CONTACT = "Montreal, Quebec, Canada  |  514-475-2374  |  ted.strauss@gmail.com  |  linkedin.com/in/tedstrauss-63b01821"
COVER_CONTACT = "514-475-2374  |  ted.strauss@gmail.com  |  linkedin.com/in/tedstrauss-63b01821"

PROFILE = (
    "Research data specialist with five years managing the data resources of McGill's McConnell "
    "Brain Imaging Centre: stewardship of sensitive multimodal neuroimaging research data across "
    "its full lifecycle, from acquisition and ingestion through curation, sharing, and archiving, "
    "under ethics and privacy requirements. Published open science contributor (Canadian Open "
    "Neuroscience Platform, DataLad) with hands-on FAIR practice and a track record of mentoring "
    "researchers into better data stewardship. Designed and delivered research data management "
    "training at McGill's Tanenbaum Open Science Institute. Python, Git, LAMP/SQL. Bilingual "
    "(EN/FR), prepared to relocate to Toronto."
)

COMPETENCIES = [
    ("Research data platform operations",
     ": managed daily data operations for a leading brain imaging centre, onboarding studies to "
     "acquisition and storage workflows; designed and launched an OwnCloud self-serve portal that "
     "replaced manual data-request handling for access to research data"),
    ("Repository lifecycle and data sharing",
     ": built pipelines and data integration workflows spanning ingestion through curation and "
     "archiving; coordinated researcher data deposits and access under documented policy; "
     "contributor to the Canadian Open Neuroscience Platform, a national neuroinformatics platform "
     "integrating distributed neuroscience data repositories"),
    ("FAIR principles and open science",
     ": hands-on FAIR practice through CONP; co-author on DataLad, a distributed system for joint "
     "management of code and data with provenance tracking; authored a published neuroimaging Data "
     "Management Plan template; built a public knowledge base on open science concepts"),
    ("Compliance and sensitive data",
     ": five years of stewardship of sensitive human neuroimaging data under research ethics, "
     "privacy, and consent requirements; maintained regulatory compliance for government "
     "learning-platform deployments"),
    ("Mentorship and training",
     ": designed and administered research data management workshops for students and faculty at "
     "the Tanenbaum Open Science Institute; built a Moodle-based self-serve education portal for "
     "researcher onboarding; certified in Scientific Community Engagement Fundamentals (CSCCE)"),
    ("Technical toolkit",
     ": Python (pandas), R, SQL via a decade of LAMP-stack development, Git version control, "
     "Docker and Kubernetes, AWS and Azure, SSO/LDAP; evaluated and adopted AI tooling including "
     "Claude Code"),
]

EXPERIENCE = [
    ("09/2016 – 10/2021", "Manager of Data Resources", "McConnell Brain Imaging Centre, The Neuro, McGill University", "Montreal, QC", [
        "Managed data resources for one of Canada's leading brain imaging centres, supporting dozens of principal investigators and hundreds of researchers working with sensitive multimodal neuroimaging data",
        "Onboarded studies to the centre's data workflows and maintained data integrity and quality from instrument output (MRI, MEG/EEG) through storage and archiving",
        "Designed and launched an OwnCloud portal that moved researchers to self-serve data management, replacing manual data-request handling",
        "Led metadata standards, documentation practices, and data sharing policy; represented the centre on national research data bodies (Portage, CONP, RDA)",
        "Designed a Moodle-based education portal for researcher onboarding, and delivered hands-on training to build a culture of data stewardship",
    ]),
    ("08/2020 – 10/2021", "Data Management & Sharing Coordinator (p/t)", "Tanenbaum Open Science Institute, The Neuro, McGill University", "Montreal, QC", [
        "Designed and administered research data management and open science training workshops for students and faculty",
        "Built a public-facing online knowledge base explaining open science concepts to non-specialist audiences",
    ]),
    ("2015–2016", "Software Developer, national CCNA dementia study", "McGill Centre for Integrative Neuroscience", "Montreal, QC", [
        "Built research software for the Canadian Consortium on Neurodegeneration in Aging, a pan-Canadian multi-site health research study",
    ]),
    ("07/2025 – 01/2026", "Moodle Analyst", "Premergency Inc.", "Canada", [
        "Coordinated between government clients and development teams to deliver learning-platform solutions, maintaining regulatory compliance for government deployments",
    ]),
    ("01/2024 – 10/2024", "Full Stack Developer", "KnowledgeOne", "Montreal, QC", [
        "Developed learning platform features for active clients (LAMP++, Moodle, LXD); evaluated and adopted AI tools including Claude Code in enterprise learning contexts",
    ]),
    ("04/2017 – 08/2021", "Founder and Director", "Sur Place Media", "Montreal, QC", [
        "Founded and directed a neighbourhood arts school; hired and coordinated 20 instructors across 150 workshop groups, and led the pivot to virtual delivery during the pandemic",
    ]),
]

ADDITIONAL = [
    "Photonic Knowledge — Database administrator and software developer (2012–2014)",
    "Kitikmeot Heritage Society — Researcher, writer, and producer (2010–2012)",
]

EDUCATION = [
    ("1998–2003", "Bachelor of Arts (Honours), Cognitive Science", "McGill University", "Montreal, QC",
     ["Minor: Music Technology. Topics: psycholinguistics, cognitive science, computational modeling of language."]),
    ("2021", "Scientific Community Engagement Fundamentals", "The Center for Scientific Collaboration and Community Engagement (CSCCE)", "", []),
]

PUBLICATIONS = [
    "Poline, J.-B., et al. (2023). Data and tools integration in the Canadian open neuroscience platform. Scientific Data, 10(1), 189.",
    "Halchenko, Y. O., et al. (2021). DataLad: distributed system for joint management of code, data, and their relationship. Journal of Open Source Software, 6(63).",
    "Strauss, T. (2021). Data Management Plan Template: Neuroimaging in the Neurosciences. Zenodo. doi:10.5281/zenodo.4673558",
]
PUB_FOOT = "9 publications | citation count 630 | h-index 9 | Full list: scholar.google.com/citations"

COVER_DATE = "August 31, 2026"
COVER_SUBJECT = "Subject: Application: Research Methods Specialist (Research Data Integration), Krembil Centre for Neuroinformatics"
COVER_SALUTATION = "Dear Hiring Committee,"
COVER_PARAS = [
    "I am applying for the Research Methods Specialist (Research Data Integration) position at the "
    "Krembil Centre for Neuroinformatics. For five years I managed the data resources of McGill's "
    "McConnell Brain Imaging Centre, and the work this posting describes, onboarding studies, "
    "keeping multimodal brain data trustworthy through ingestion and preparation, coordinating a "
    "repository lifecycle, and mentoring researchers in FAIR practice, is the closest description "
    "of that job I have seen since I held it.",

    "CAMH's formal adoption of Open Science Principles is a large part of why I am applying. My "
    "open science record is practical rather than aspirational: I contributed to the Canadian Open "
    "Neuroscience Platform, a national effort to make sensitive neuroscience data findable and "
    "reusable under FAIR principles, co-authored its data and tools integration paper in "
    "Scientific Data (2023), and co-authored DataLad, a distributed system for managing code and "
    "data with provenance tracking. At McGill's Tanenbaum Open Science Institute I designed and "
    "ran the research data management workshops and built the public knowledge base that turned "
    "open science from policy into practice for students and faculty.",

    "The operational core of the role matches what I did daily at the Brain Imaging Centre: "
    "onboarding studies into acquisition and storage workflows for MRI and other multimodal "
    "neuroimaging data, maintaining integrity and documentation from instrument output through "
    "archiving, and handling access to sensitive data under ethics and consent requirements. When "
    "manual data-request handling became a bottleneck, I designed and launched an OwnCloud "
    "self-serve portal with documented permissions, the same pattern as coordinating a databank "
    "portal's setup, user training, and access requests. I authored a published Data Management "
    "Plan template for neuroimaging, and I work in Python and Git, with SQL from a decade of "
    "LAMP-stack development.",

    "One honest note: I have not administered REDCap, XNAT, or LabKey in production. I have lived "
    "next to that tool class for years, managing the data those systems capture and staging it for "
    "curation, and electronic data capture configuration is the kind of tooling I pick up quickly, "
    "as I have done across Moodle, OwnCloud, Drupal, and cloud platforms. The 12-month contract "
    "term works for me, and I am prepared to relocate to Toronto for an on-site role at 250 "
    "College Street. I am also bilingual in French and English.",

    "I would welcome the opportunity to discuss how I can support the Neuroinformatics Platform "
    "and the BrainHealth Databank. Thank you for your consideration.",
]
COVER_CLOSING = "Kind regards,"

# ==================================================================
# DOCX RENDERING
# ==================================================================
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BLUE = RGBColor(0x21, 0x6F, 0xAA)
BLACK = RGBColor(0x00, 0x00, 0x00)
GRAY = RGBColor(0x55, 0x55, 0x55)


def _margins(doc, top=1.5, bottom=1.5, left=2.0, right=2.0):
    for s in doc.sections:
        s.top_margin, s.bottom_margin = Cm(top), Cm(bottom)
        s.left_margin, s.right_margin = Cm(left), Cm(right)


def _name_header(doc, name, contact):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(name); r.bold = True; r.font.size = Pt(28); r.font.color.rgb = BLUE
    p2 = doc.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(contact); r2.font.size = Pt(10); r2.font.color.rgb = GRAY
    p2.paragraph_format.space_after = Pt(6)


def _heading(doc, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title.upper()); r.bold = True; r.font.size = Pt(11); r.font.color.rgb = BLUE
    pPr = p._p.get_or_add_pPr(); pBdr = OxmlElement('w:pBdr'); bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single'); bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1'); bottom.set(qn('w:color'), '216FAA')
    pBdr.append(bottom); pPr.append(pBdr)


def _entry(doc, dates, title, org, location, bullets):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    r1 = p.add_run(title); r1.bold = True; r1.font.size = Pt(10.5); r1.font.color.rgb = BLACK
    r2 = p.add_run(f"  —  {org}"); r2.font.size = Pt(10.5); r2.font.color.rgb = GRAY
    if location:
        r3 = p.add_run(f", {location}"); r3.font.size = Pt(10.5); r3.font.color.rgb = GRAY
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_before = Pt(0); p2.paragraph_format.space_after = Pt(2)
    p2.paragraph_format.keep_with_next = True
    r4 = p2.add_run(dates); r4.italic = True; r4.font.size = Pt(9.5); r4.font.color.rgb = GRAY
    for b in bullets:
        _bullet(doc, b)


def _bullet(doc, text, indent=0.25):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(1); p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.left_indent = Inches(indent)
    for i, part in enumerate(text.split('**')):
        if not part:
            continue
        r = p.add_run(part); r.font.size = Pt(10); r.bold = (i % 2 == 1)


def _para(doc, text, size=10, italic=False, color=None, space_after=4):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(text); r.font.size = Pt(size); r.italic = italic
    if color:
        r.font.color.rgb = color
    return p


def build_docx_cv():
    doc = Document(); _margins(doc)
    _name_header(doc, NAME, CV_CONTACT)
    _para(doc, PROFILE, size=10, space_after=6)
    _heading(doc, "Core Competencies")
    for label, rest in COMPETENCIES:
        _bullet(doc, f"**{label}**{rest}")
    _heading(doc, "Relevant Experience")
    for e in EXPERIENCE:
        _entry(doc, *e)
    _heading(doc, "Additional Experience")
    for a in ADDITIONAL:
        _bullet(doc, a)
    _heading(doc, "Education")
    for e in EDUCATION:
        _entry(doc, *e)
    _heading(doc, "Languages")
    _para(doc, "English (native), French (bilingual, professional working proficiency).", size=10)
    _heading(doc, "Selected Publications")
    for pub in PUBLICATIONS:
        _bullet(doc, pub)
    _para(doc, PUB_FOOT, size=9, italic=True, color=GRAY)
    doc.save("docx/cv_camh_neuroinformatics.docx")
    print("OK cv_camh_neuroinformatics.docx")


def build_docx_cover():
    doc = Document(); _margins(doc, top=2.0, bottom=2.0, left=2.5, right=2.5)
    _name_header(doc, NAME, COVER_CONTACT)
    _para(doc, COVER_DATE, size=10, space_after=12)
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(6)
    r = p.add_run(COVER_SUBJECT); r.bold = True; r.font.size = Pt(10)
    _para(doc, COVER_SALUTATION, size=10, space_after=6)
    for para in COVER_PARAS:
        pp = doc.add_paragraph(); pp.paragraph_format.space_after = Pt(8)
        rr = pp.add_run(para); rr.font.size = Pt(10)
    _para(doc, COVER_CLOSING, size=10, space_after=24)
    _para(doc, NAME, size=10)
    doc.save("docx/cover_camh_neuroinformatics.docx")
    print("OK cover_camh_neuroinformatics.docx")


# ==================================================================
# PDF RENDERING
# ==================================================================
from reportlab.lib.pagesizes import letter
from reportlab.lib.units import cm
from reportlab.lib.colors import HexColor
from reportlab.lib.enums import TA_CENTER
from reportlab.platypus import (
    BaseDocTemplate, PageTemplate, Frame, Paragraph, Spacer,
    HRFlowable, ListFlowable, ListItem, KeepTogether,
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

PDF_BLUE = HexColor("#216FAA")
PDF_GRAY = HexColor("#555555")


def _esc(t):
    return t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _bold_md(t):
    t = _esc(t)
    parts = t.split("**")
    out = ""
    for i, p in enumerate(parts):
        out += f"<b>{p}</b>" if i % 2 == 1 else p
    return out


def _styles():
    styles = getSampleStyleSheet()
    return {
        "name": ParagraphStyle("Name", parent=styles["Normal"], fontName="Helvetica-Bold",
                               fontSize=22, textColor=PDF_BLUE, alignment=TA_CENTER, spaceAfter=2, leading=26),
        "contact": ParagraphStyle("Contact", parent=styles["Normal"], fontName="Helvetica",
                                  fontSize=9, textColor=PDF_GRAY, alignment=TA_CENTER, spaceAfter=8, leading=12),
        "heading": ParagraphStyle("Heading", parent=styles["Normal"], fontName="Helvetica-Bold",
                                  fontSize=11, textColor=PDF_BLUE, spaceBefore=8, spaceAfter=2, leading=13),
        "body": ParagraphStyle("Body", parent=styles["Normal"], fontName="Helvetica", fontSize=10,
                               leading=13, spaceAfter=3),
        "bullet": ParagraphStyle("Bullet", parent=styles["Normal"], fontName="Helvetica",
                                 fontSize=9.5, leading=12, spaceAfter=1),
        "entry_title": ParagraphStyle("EntryTitle", parent=styles["Normal"], fontName="Helvetica",
                                      fontSize=10.5, spaceBefore=5, spaceAfter=0, leading=12),
        "entry_dates": ParagraphStyle("EntryDates", parent=styles["Normal"], fontName="Helvetica",
                                      fontSize=9, textColor=PDF_GRAY, spaceAfter=2, leading=11),
        "letter": ParagraphStyle("Letter", parent=styles["Normal"], fontName="Helvetica",
                                 fontSize=10.5, leading=15, spaceAfter=8),
        "small_italic": ParagraphStyle("SmallItalic", parent=styles["Normal"],
                                       fontName="Helvetica-Oblique", fontSize=8.5,
                                       textColor=PDF_GRAY, spaceAfter=2, leading=11),
    }


def _render(path, story, top=1.5, bottom=1.3, left=2.0, right=2.0):
    doc = BaseDocTemplate(path, pagesize=letter,
                          topMargin=top * cm, bottomMargin=bottom * cm,
                          leftMargin=left * cm, rightMargin=right * cm)
    frame = Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id="main")
    doc.addPageTemplates([PageTemplate(id="all", frames=[frame])])
    doc.build(story)
    print(f"OK {os.path.basename(path)}")


def build_pdf_cv():
    S = _styles()

    def rule():
        return HRFlowable(width="100%", thickness=0.6, color=PDF_BLUE, spaceBefore=1, spaceAfter=3)

    def heading(text):
        return [Paragraph(_esc(text.upper()), S["heading"]), rule()]

    def bullets(items):
        return ListFlowable(
            [ListItem(Paragraph(_bold_md(t), S["bullet"]), leftIndent=10, value="•") for t in items],
            bulletType="bullet", bulletColor=PDF_BLUE, leftIndent=12, bulletFontSize=8,
        )

    def entry(dates, title, org, location, bl):
        loc = f", {_esc(location)}" if location else ""
        head = (f'<b>{_esc(title)}</b>'
                f'<font color="#555555">&nbsp;&nbsp;—&nbsp;&nbsp;{_esc(org)}{loc}</font>')
        header = [Paragraph(head, S["entry_title"]), Paragraph(_esc(dates), S["entry_dates"])]
        if not bl:
            return [KeepTogether(header)]
        out = [KeepTogether(header + [bullets(bl[:1])])]
        if len(bl) > 1:
            out.append(bullets(bl[1:]))
        return out

    story = [Paragraph(_esc(NAME), S["name"]),
             Paragraph(_esc(CV_CONTACT), S["contact"]),
             Paragraph(_esc(PROFILE), S["body"])]
    story += heading("Core Competencies")
    story.append(bullets([f"**{l}**{r}" for l, r in COMPETENCIES]))
    story += heading("Relevant Experience")
    for e in EXPERIENCE:
        story += entry(*e)
    story += heading("Additional Experience")
    story.append(bullets(ADDITIONAL))
    story += heading("Education")
    for e in EDUCATION:
        story += entry(*e)
    story += heading("Languages")
    story.append(Paragraph("English (native), French (bilingual, professional working proficiency).", S["body"]))
    story += heading("Selected Publications")
    story.append(bullets(PUBLICATIONS))
    story.append(Paragraph(_esc(PUB_FOOT), S["small_italic"]))

    _render("docx/cv_camh_neuroinformatics.pdf", story)


def build_pdf_cover():
    S = _styles()
    story = [Paragraph(_esc(NAME), S["name"]),
             Paragraph(_esc(COVER_CONTACT), S["contact"]),
             Spacer(1, 6),
             Paragraph(_esc(COVER_DATE), S["body"]),
             Spacer(1, 8),
             Paragraph(f"<b>{_esc(COVER_SUBJECT)}</b>", S["body"]),
             Spacer(1, 4),
             Paragraph(_esc(COVER_SALUTATION), S["letter"])]
    for para in COVER_PARAS:
        story.append(Paragraph(_esc(para), S["letter"]))
    story += [Spacer(1, 6),
              KeepTogether([Paragraph(_esc(COVER_CLOSING), S["body"]),
                            Spacer(1, 12),
                            Paragraph(_esc(NAME), S["body"])])]
    _render("docx/cover_camh_neuroinformatics.pdf", story, top=1.6, bottom=1.2, left=2.5, right=2.5)


if __name__ == "__main__":
    build_docx_cv()
    build_docx_cover()
    build_pdf_cv()
    build_pdf_cover()
    print("\nDone. Files in docx/")
