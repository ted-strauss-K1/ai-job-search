"""Generate DOCX for McGill Digital Literacy and Documentation Specialist."""
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

os.makedirs("docx", exist_ok=True)

BLUE = RGBColor(0x21, 0x6F, 0xAA)
BLACK = RGBColor(0x00, 0x00, 0x00)
GRAY = RGBColor(0x55, 0x55, 0x55)

def set_margins(doc, top=1.5, bottom=1.5, left=2.0, right=2.0):
    for section in doc.sections:
        section.top_margin = Cm(top)
        section.bottom_margin = Cm(bottom)
        section.left_margin = Cm(left)
        section.right_margin = Cm(right)

def add_name_header(doc, name, contact):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(name)
    run.bold = True
    run.font.size = Pt(28)
    run.font.color.rgb = BLUE
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run2 = p2.add_run(contact)
    run2.font.size = Pt(10)
    run2.font.color.rgb = GRAY
    p2.paragraph_format.space_after = Pt(6)

def add_section_heading(doc, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(title.upper())
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = BLUE
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), '216FAA')
    pBdr.append(bottom)
    pPr.append(pBdr)

def add_job_entry(doc, dates, title, org, location, bullets):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(0)
    r1 = p.add_run(title)
    r1.bold = True
    r1.font.size = Pt(10.5)
    r2 = p.add_run(f"  —  {org}")
    r2.font.size = Pt(10.5)
    r2.font.color.rgb = GRAY
    if location:
        r3 = p.add_run(f", {location}")
        r3.font.size = Pt(10.5)
        r3.font.color.rgb = GRAY
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(2)
    r4 = p2.add_run(dates)
    r4.italic = True
    r4.font.size = Pt(9.5)
    r4.font.color.rgb = GRAY
    for b in bullets:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_before = Pt(1)
        bp.paragraph_format.space_after = Pt(1)
        bp.paragraph_format.left_indent = Inches(0.25)
        parts = b.split('**')
        for i, part in enumerate(parts):
            if not part:
                continue
            r = bp.add_run(part)
            r.font.size = Pt(10)
            r.bold = (i % 2 == 1)

def add_bullet(doc, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.left_indent = Inches(0.25)
    parts = text.split('**')
    for i, part in enumerate(parts):
        if not part:
            continue
        r = p.add_run(part)
        r.font.size = Pt(10)
        r.bold = (i % 2 == 1)

def add_para(doc, text, size=10, italic=False, color=None, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.italic = italic
    if color:
        r.font.color.rgb = color
    return p

# ============================================================
# CV — McGill Digital Literacy and Documentation Specialist
# ============================================================

def cv_mcgill():
    doc = Document()
    set_margins(doc)
    add_name_header(doc, "Ted Strauss",
        "Montreal, Quebec, Canada  |  514-475-2374  |  ted.strauss@gmail.com  |  linkedin.com/in/tedstrauss-63b01821")

    add_para(doc, "Digital literacy practitioner and documentation specialist with experience designing and delivering technology training programs for academic, government, and institutional audiences. At the McConnell Brain Imaging Centre and Tanenbaum Open Science Institute at McGill University, built self-serve learning platforms, ran digital literacy workshops, and developed public knowledge bases translating complex technical workflows for non-specialist audiences. Brings a record of making IT systems accessible through documentation and training, grounded in direct experience with McGill's research community and bilingual institutional environment. Bilingual (EN/FR).", size=10, space_after=6)

    add_section_heading(doc, "Core Competencies")
    add_bullet(doc, "**Digital literacy programming**: designed and delivered IT and data literacy workshops for faculty, students, researchers, and government staff; built self-serve learning platforms to reduce time-to-productivity for new users")
    add_bullet(doc, "**Documentation and knowledge base development**: built and maintained public-facing knowledge bases on technical workflows and open science practices for non-specialist audiences; wrote platform guides and onboarding documentation for diverse user groups")
    add_bullet(doc, "**LMS and training platform management**: Moodle administration, content development, and deployment across academic and government contexts; designed e-learning modules, blended learning programs, and self-directed learning paths")
    add_bullet(doc, "**Bilingual facilitation (EN/FR)**: designed and delivered training programs in English and French; produced bilingual documentation; at ease in McGill's bilingual institutional environment")
    add_bullet(doc, "**IT adoption and community engagement**: embedded technology literacy into workflows through ongoing community engagement; change management approaches to increase platform uptake and reduce support friction")

    add_section_heading(doc, "Professional Experience")
    add_job_entry(doc, "07/2025 – 01/2026", "Moodle Analyst", "Premergency Inc.", "Canada", [
        "Designed bespoke learning & development solutions for government clients, including course content, user documentation, and training delivery aligned to regulatory requirements",
        "Managed full LMS migrations to Moodle, producing onboarding documentation and support materials for end users across multiple client organizations",
        "Coordinated between government clients and technical teams to translate requirements into training programs and platform configurations",
    ])
    add_job_entry(doc, "01/2024 – 10/2024", "Full Stack Developer", "KnowledgeOne", "Montreal, QC", [
        "Developed and maintained LMS features for active clients on the eConcordia Moodle platform (LAMP++, LXD)",
        "Evaluated AI tools for enterprise learning, assessing usability and adoption pathways for non-specialist users",
    ])
    add_job_entry(doc, "09/2016 – 10/2021", "Manager of Data Resources", "McConnell Brain Imaging Centre, The Neuro, McGill University", "Montreal, QC", [
        "Designed and deployed a Moodle-based self-serve education portal for researcher onboarding, reducing time-to-productivity for new staff across a multi-PI neuroimaging environment",
        "Developed platform guides and data management documentation for researchers with varying levels of technical proficiency",
        "Delivered training workshops on digital tools, data workflows, and platform adoption for students, faculty, and research staff",
        "Contributed to national research data policy through Portage, CONP, and RDA committee work; co-authored open science frameworks adopted across Canadian institutions",
    ])
    add_job_entry(doc, "08/2020 – 10/2021", "Data Management & Sharing Coordinator (p/t)", "Tanenbaum Open Science Institute, The Neuro, McGill University", "Montreal, QC", [
        "Built a public-facing knowledge base on open science, data management, and research workflows, written for general audiences and researchers at varying career stages in English and French",
        "Designed and delivered training workshops for students, faculty, and research teams on digital tools and data practices",
    ])
    add_job_entry(doc, "04/2017 – 08/2021", "Founder and Director", "Sur Place Media", "Montreal, QC", [
        "Founded and operated a neighbourhood arts school; hosted 150 workshop groups with 20 instructors; designed programming for a bilingual (EN/FR) learner community; led full pivot to virtual delivery during COVID-19",
    ])
    add_job_entry(doc, "02/2015 – 09/2016", "Software Developer", "McGill Centre for Integrative Neuroscience, McGill University", "Montreal, QC", [
        "Programmed data acquisition pipeline for the national CCNA longitudinal study using the LORIS platform; contributed to research platform documentation and user onboarding",
    ])

    add_section_heading(doc, "Additional Experience")
    add_bullet(doc, "Language Mate — Developer and UX designer (2007–2010)")
    add_bullet(doc, "Moss Rehabilitation Hospital — Research software developer (2007–2010)")
    add_bullet(doc, "University of Connecticut & Columbia University — Research assistant & software developer (2004–2008)")

    add_section_heading(doc, "Education")
    add_job_entry(doc, "1998–2003", "B.A. (Honours), Cognitive Science", "McGill University", "Montreal, QC", [
        "Minor: Music Technology. Topics: Psycholinguistics, cognitive science, computational modeling of language.",
    ])
    add_job_entry(doc, "2021", "Scientific Community Engagement Fundamentals", "CSCCE", "", [])

    add_section_heading(doc, "Languages")
    add_para(doc, "English (native), French (bilingual).", size=10)

    add_section_heading(doc, "Selected Publications")
    add_bullet(doc, "Poline, J.-B., et al. (2023). Data and tools integration in the Canadian open neuroscience platform. Scientific Data, 10(1), 189.")
    add_bullet(doc, "Strauss, T. (2021). Data Management Plan Template: Neuroimaging in the Neurosciences. Zenodo. doi:10.5281/zenodo.4673558")
    add_para(doc, "9 publications | citation count: 630 | h-index: 9 | scholar.google.com/citations", size=9, italic=True, color=GRAY)

    add_section_heading(doc, "References")
    add_para(doc, "Available upon request.", size=10)

    doc.save("docx/cv_mcgill_digitalliteracy.docx")
    print("cv_mcgill_digitalliteracy.docx")

# ============================================================
# Cover Letter — McGill Digital Literacy and Documentation Specialist
# ============================================================

def cover_mcgill():
    doc = Document()
    set_margins(doc, top=2.0, bottom=2.0, left=2.5, right=2.5)
    add_name_header(doc, "Ted Strauss",
        "514-475-2374  |  ted.strauss@gmail.com  |  linkedin.com/in/tedstrauss-63b01821")

    add_para(doc, "May 25, 2026", size=10, space_after=12)
    add_para(doc, "Dear McGill IT Services Hiring Team,", size=10, space_after=8)

    for para in [
        "I am applying for the Digital Literacy and Documentation Specialist position because it combines two things I have spent my career doing: making complex technical systems accessible to non-specialist users through documentation and training, and building the learning communities that make adoption stick. The role sits exactly within McGill IT's Communications, Training and Community Engagement team, which is the model I have worked toward across every position I have held.",
        "I know this environment from the inside. From 2016 to 2021, I worked at the McConnell Brain Imaging Centre and Tanenbaum Open Science Institute at McGill, where I designed and deployed a Moodle-based self-serve education portal for researcher onboarding, ran digital literacy workshops for faculty and students in both official languages, and built a public-facing knowledge base on data management practices for a general audience. The challenge in all of that work was not the content itself, it was meeting users where they were: calibrating documentation and training for audiences with very different comfort levels around technology, and writing clearly enough that the documentation actually got used. At Premergency Inc., I brought that same approach to government LMS clients, producing onboarding documentation, user guides, and training programs aligned to regulatory compliance requirements.",
        "I am bilingual in French and English, based in Montreal, and have worked extensively in McGill's institutional context. I understand the rhythms of the academic calendar, the diversity of the McGill user community, and the specific challenge of supporting technology adoption across a bilingual, research-intensive university environment.",
        "I would welcome the opportunity to contribute to McGill IT's digital literacy and community engagement programs. Thank you for your consideration.",
    ]:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(8)
        r = p.add_run(para)
        r.font.size = Pt(10)

    add_para(doc, "Kind regards,", size=10, space_after=24)
    add_para(doc, "Ted Strauss", size=10)

    doc.save("docx/cover_mcgill_digitalliteracy.docx")
    print("cover_mcgill_digitalliteracy.docx")

cv_mcgill()
cover_mcgill()
print("Done.")
