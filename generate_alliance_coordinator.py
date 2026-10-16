"""
Generate application docs for the Digital Research Alliance of Canada —
Coordinator, Events and Engagement Support (Contract), Requisition ID 1086.
Remote. Term position until March 31, 2027. Salary $70,550 - $95,450 CAD.
Outputs (in docx/):
  - cv_alliance_coordinator.docx     / .pdf
  - cover_alliance_coordinator.docx  / .pdf

Content is defined ONCE below, then rendered to both docx (python-docx) and
pdf (reportlab) so the two formats never drift.

Targeting note: this posting is an events/engagement coordinator role, NOT a data
governance role. The CV is deliberately re-anchored — Sur Place Media and the TOSI
workshop programme lead, and the governance record appears as supporting evidence for
the posting's REQUIRED "records management, scheduling and documentation" and
"attention to detail in handling data, records" qualifications.
"""
import os

os.makedirs("docx", exist_ok=True)

# ------------------------------------------------------------------
# SHARED CONTENT
# ------------------------------------------------------------------

NAME = "Ted Strauss"
CV_CONTACT = "Montreal, Quebec, Canada  |  514-475-2374  |  ted.strauss@gmail.com  |  linkedin.com/in/tedstrauss-63b01821"
COVER_CONTACT = "514-475-2374  |  ted.strauss@gmail.com  |  linkedin.com/in/tedstrauss-63b01821"

PROFILE = (
    "Coordinator and program organizer with over a decade of experience across Canada's academic, "
    "research, and learning-technology sectors, combining hands-on event delivery with the records "
    "and documentation discipline that keeps engagement programs running. Founded and ran a "
    "neighbourhood arts school that hosted 150 workshop groups with 20 instructors, and led its full "
    "pivot to virtual delivery during the pandemic. Designed and administered data management "
    "workshops for students and faculty at McGill's Tanenbaum Open Science Institute, and represented "
    "the McConnell Brain Imaging Centre on the Portage Network, whose operations were integrated into "
    "the Alliance in 2021, and on CONP and RDA. Bilingual (EN/FR), Montreal-based, and already fluent "
    "in the research data management and research software communities this role serves."
)

# (bold label, remaining text) — label is emphasised
COMPETENCIES = [
    ("Event coordination and delivery",
     ": 150 workshop groups planned and delivered with 20 instructors, in person and online, covering "
     "scheduling, instructor coordination, registration, materials, and day-of support; led a full "
     "pivot to virtual delivery under time pressure during the COVID-19 pandemic"),
    ("Stakeholder engagement and outreach",
     ": coordinated across faculty, students, and institutional partners for the launch of a new "
     "academic unit at McGill, a primary member of the Alliance; represented an institution on "
     "national multi-stakeholder committees where alignment depended on sustained relationships "
     "rather than authority"),
    ("Records management and data integrity",
     ": maintained institutional records, metadata, and documentation standards for a major research "
     "centre, including data policies for daily acquisition workflows; separately kept the participant "
     "and scheduling records for a 150-group workshop program"),
    ("Communications and content",
     ": built a public-facing knowledge base explaining research data concepts to non-specialists; "
     "drafted and formatted training and outreach material; produced video, podcast, and multimedia "
     "content (Adobe suite)"),
    ("Research-ecosystem fluency",
     ": six years inside Canada's research data community (McGill, 2015-2021), including committee "
     "work with the Portage Network, whose operations were folded into the Alliance in 2021, and with "
     "CONP and RDA; certified in Scientific Community Engagement Fundamentals (CSCCE); contributor to "
     "the Canadian Open Neuroscience Platform"),
    ("Administrative and digital tools",
     ": Microsoft Office 365 and Teams; scheduling, invoice and expense tracking, and purchasing as a "
     "small-business operator; content and participant-communication platforms including WordPress, "
     "Drupal, OwnCloud, and Moodle (course enrolment, announcements, and bulk participant messaging); "
     "evaluated and adopted AI tools including Claude Code"),
]

# each: (dates, title, org, location, [bullets])
EXPERIENCE = [
    ("04/2017 – 08/2021", "Founder and Director", "Sur Place Media", "Montreal, QC", [
        "Founded and operated a neighbourhood arts school offering workshops in creative writing, media arts, and visual arts; hosted 150 workshop groups with 20 instructors",
        "Owned the full event cycle: scheduling and timelines, instructor coordination, registration and participant communication, materials preparation, purchasing, and on-site delivery",
        "Led the organization's pivot to fully virtual delivery during the COVID-19 pandemic, rebuilding programming and participant communication for online delivery under time pressure",
        "Ran the administrative side end to end: records, invoicing, expense tracking, purchasing, and scheduling",
        "Produced video, podcast, and multimedia content to support programming and outreach",
    ]),
    ("08/2020 – 10/2021", "Data Management & Sharing Coordinator (p/t)", "Tanenbaum Open Science Institute, The Neuro, McGill University", "Montreal, QC", [
        "Designed and administered training workshops for students and faculty, coordinating scheduling, materials, and participant follow-up",
        "Supported the launch of a new academic unit focused on open science, coordinating across faculty, students, and institutional stakeholders",
        "Built a public-facing online knowledge base explaining open science concepts to non-specialist audiences",
    ]),
    ("07/2025 – 01/2026", "Moodle Analyst", "Premergency Inc.", "Canada", [
        "Coordinated between government clients and development teams, tracking requirements, timelines, and deliverables across both sides and keeping documentation aligned with compliance obligations",
    ]),
    ("09/2016 – 10/2021", "Manager of Data Resources", "McConnell Brain Imaging Centre, The Neuro, McGill University", "Montreal, QC", [
        "Managed data workflows, records, and documentation standards for one of Canada's leading brain imaging centres, supporting dozens of principal investigators and hundreds of researchers",
        "Represented the centre on national research data bodies (Portage, CONP, RDA), sustaining multi-institutional relationships and contributing to national research data policy",
        "Designed a Moodle-based self-serve education portal for researcher onboarding, and delivered hands-on training to build data literacy",
        "Launched an OwnCloud portal that moved researchers to self-serve data management, replacing manual data-request handling and removing an administrative bottleneck",
    ]),
    ("01/2024 – 10/2024", "Full Stack Developer", "KnowledgeOne", "Montreal, QC", [
        "Developed learning platform features for active clients (LAMP++, Moodle, LXD); evaluated and adopted AI tools in enterprise learning contexts",
    ]),
]

ADDITIONAL = [
    "Kitikmeot Heritage Society — Researcher, writer, and producer (2010–2012): co-authored The Art of Inuinnait Drum Dancing (2012), published by the Kitikmeot Heritage Society",
    "McGill Centre for Integrative Neuroscience — Software developer, national CCNA dementia study (2015–2016)",
    "Photonic Knowledge — Database administrator and software developer (2012–2014)",
]

EDUCATION = [
    ("1998–2003", "B.A. (Honours), Cognitive Science", "McGill University", "Montreal, QC",
     ["Minor: Music Technology. Topics: psycholinguistics, cognitive science, computational modeling of language, music technology."]),
    ("2021", "Scientific Community Engagement Fundamentals", "The Center for Scientific Collaboration and Community Engagement (CSCCE)", "", []),
]

PUBLICATIONS = [
    "Poline, J.-B., et al. (2023). Data and tools integration in the Canadian open neuroscience platform. Scientific Data, 10(1), 189.",
    "Strauss, T. (2021). Data Management Plan Template: Neuroimaging in the Neurosciences. Zenodo. doi:10.5281/zenodo.4673558",
]
PUB_FOOT = "9 publications | citation count 630 | h-index 9 | Full list: scholar.google.com/citations"

COVER_DATE = "August 7, 2026"
COVER_SUBJECT = "Subject: Application: Coordinator, Events and Engagement Support (Contract), Requisition ID 1086"
COVER_SALUTATION = "Dear Hiring Team,"
COVER_PARAS = [
    "I am applying for the Coordinator, Events and Engagement Support position. I have spent six years "
    "inside the research data management community this role serves, and four years running the events "
    "end of an organization myself.",

    "The Alliance's engagement work runs on town halls, an AGM, member focus groups, bilingual "
    "webinars, and cross-country consultations, and I know those audiences from the inside. I "
    "represented the McConnell Brain Imaging Centre on the Portage Network, whose operations were "
    "integrated into the Alliance in 2021, and on CONP and RDA. I designed and administered data "
    "management workshops for students and faculty at the Tanenbaum Open Science Institute, and built "
    "a public knowledge base explaining open science to non-specialists. A coordinator who already "
    "understands what a session on RDM or research software is about can anticipate what speakers "
    "need and catch a problem in a draft agenda before it reaches participants.",

    "The event coordination itself I have owned end to end. At Sur Place Media, the arts school I "
    "founded and directed, I delivered 150 workshop groups with 20 instructors: scheduling and "
    "timelines, instructor coordination, registration and participant communication, materials, venue "
    "and supplier arrangements, and day-of support. When the pandemic closed in-person "
    "delivery, I led the pivot to fully virtual delivery and kept the programming running. I also "
    "ran the administration behind it: records, invoicing, expense tracking, purchasing, and "
    "scheduling. On records and documentation, I spent five "
    "years maintaining records, metadata, and documentation standards for a major research centre; "
    "applied to registrations, attendance, and contact data, that discipline is the difference between "
    "a CRM people trust and one they quietly work around. I have not administered a dedicated CRM, "
    "though I have managed contact and content systems including WordPress, Drupal, and OwnCloud.",

    "On seniority: my recent titles have been more senior than this one, and I am applying "
    "deliberately. The parts of my career I have most wanted to do more of, building communities of "
    "practice, running workshops, and making technical work accessible to the people it serves, are "
    "the substance of this role rather than a side activity in it, and the Alliance is the "
    "organization I would most like to do that for. The posted range and the term through March 2027 "
    "both fit what I am looking for. I am bilingual in French and English, which matters for an "
    "engagement team that runs its webinars in both languages, and I am based in Montreal, one of "
    "your consultation cities.",

    "I would welcome the opportunity to discuss how I can support the Engagement team's events. Thank "
    "you for your consideration.",
]
COVER_CLOSING = "Kind regards,"

# ==================================================================
# DOCX RENDERING
# ==================================================================
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BLUE = RGBColor(0x21, 0x6F, 0xAA)
BLACK = RGBColor(0x00, 0x00, 0x00)
GRAY = RGBColor(0x55, 0x55, 0x55)


def _margins(doc, top=1.5, bottom=1.5, left=2.0, right=2.0):
    for s in doc.sections:
        s.top_margin, s.bottom_margin = Cm(top), Cm(bottom)
        s.left_margin, s.right_margin = Cm(left), Cm(right)


def _name_header(doc, name, contact):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(name); r.bold = True; r.font.size = Pt(28); r.font.color.rgb = BLUE
    p2 = doc.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(contact); r2.font.size = Pt(10); r2.font.color.rgb = GRAY
    p2.paragraph_format.space_after = Pt(6)


def _heading(doc, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title.upper()); r.bold = True; r.font.size = Pt(11); r.font.color.rgb = BLUE
    pPr = p._p.get_or_add_pPr(); pBdr = OxmlElement('w:pBdr'); bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single'); bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1'); bottom.set(qn('w:color'), '216FAA')
    pBdr.append(bottom); pPr.append(pBdr)


def _entry(doc, dates, title, org, location, bullets):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    r1 = p.add_run(title); r1.bold = True; r1.font.size = Pt(10.5); r1.font.color.rgb = BLACK
    r2 = p.add_run(f"  —  {org}"); r2.font.size = Pt(10.5); r2.font.color.rgb = GRAY
    if location:
        r3 = p.add_run(f", {location}"); r3.font.size = Pt(10.5); r3.font.color.rgb = GRAY
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_before = Pt(0); p2.paragraph_format.space_after = Pt(2)
    p2.paragraph_format.keep_with_next = True
    r4 = p2.add_run(dates); r4.italic = True; r4.font.size = Pt(9.5); r4.font.color.rgb = GRAY
    for b in bullets:
        _bullet(doc, b)


def _bullet(doc, text, indent=0.25):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(1); p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.left_indent = Inches(indent)
    for i, part in enumerate(text.split('**')):
        if not part:
            continue
        r = p.add_run(part); r.font.size = Pt(10); r.bold = (i % 2 == 1)


def _para(doc, text, size=10, italic=False, color=None, space_after=4):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(text); r.font.size = Pt(size); r.italic = italic
    if color:
        r.font.color.rgb = color
    return p


def build_docx_cv():
    doc = Document(); _margins(doc)
    _name_header(doc, NAME, CV_CONTACT)
    _para(doc, PROFILE, size=10, space_after=6)
    _heading(doc, "Core Competencies")
    for label, rest in COMPETENCIES:
        _bullet(doc, f"**{label}**{rest}")
    _heading(doc, "Relevant Experience")
    for e in EXPERIENCE:
        _entry(doc, *e)
    _heading(doc, "Additional Experience")
    for a in ADDITIONAL:
        _bullet(doc, a)
    _heading(doc, "Education")
    for e in EDUCATION:
        _entry(doc, *e)
    _heading(doc, "Languages")
    _para(doc, "English (native), French (bilingual, professional working proficiency).", size=10)
    _heading(doc, "Selected Publications")
    for pub in PUBLICATIONS:
        _bullet(doc, pub)
    _para(doc, PUB_FOOT, size=9, italic=True, color=GRAY)
    doc.save("docx/cv_alliance_coordinator.docx")
    print("OK cv_alliance_coordinator.docx")


def build_docx_cover():
    doc = Document(); _margins(doc, top=2.0, bottom=2.0, left=2.5, right=2.5)
    _name_header(doc, NAME, COVER_CONTACT)
    _para(doc, COVER_DATE, size=10, space_after=12)
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(6)
    r = p.add_run(COVER_SUBJECT); r.bold = True; r.font.size = Pt(10)
    _para(doc, COVER_SALUTATION, size=10, space_after=6)
    for para in COVER_PARAS:
        pp = doc.add_paragraph(); pp.paragraph_format.space_after = Pt(8)
        rr = pp.add_run(para); rr.font.size = Pt(10)
    _para(doc, COVER_CLOSING, size=10, space_after=24)
    _para(doc, NAME, size=10)
    doc.save("docx/cover_alliance_coordinator.docx")
    print("OK cover_alliance_coordinator.docx")


# ==================================================================
# PDF RENDERING
# ==================================================================
from reportlab.lib.pagesizes import letter
from reportlab.lib.units import cm
from reportlab.lib.colors import HexColor
from reportlab.lib.enums import TA_CENTER
from reportlab.platypus import (
    BaseDocTemplate, PageTemplate, Frame, Paragraph, Spacer,
    HRFlowable, ListFlowable, ListItem, KeepTogether,
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

PDF_BLUE = HexColor("#216FAA")
PDF_GRAY = HexColor("#555555")


def _esc(t):
    return t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _bold_md(t):
    t = _esc(t)
    parts = t.split("**")
    out = ""
    for i, p in enumerate(parts):
        out += f"<b>{p}</b>" if i % 2 == 1 else p
    return out


def _styles():
    styles = getSampleStyleSheet()
    return {
        "name": ParagraphStyle("Name", parent=styles["Normal"], fontName="Helvetica-Bold",
                               fontSize=22, textColor=PDF_BLUE, alignment=TA_CENTER, spaceAfter=2, leading=26),
        "contact": ParagraphStyle("Contact", parent=styles["Normal"], fontName="Helvetica",
                                  fontSize=9, textColor=PDF_GRAY, alignment=TA_CENTER, spaceAfter=8, leading=12),
        "heading": ParagraphStyle("Heading", parent=styles["Normal"], fontName="Helvetica-Bold",
                                  fontSize=11, textColor=PDF_BLUE, spaceBefore=8, spaceAfter=2, leading=13),
        "body": ParagraphStyle("Body", parent=styles["Normal"], fontName="Helvetica", fontSize=10,
                               leading=13, spaceAfter=3),
        "bullet": ParagraphStyle("Bullet", parent=styles["Normal"], fontName="Helvetica",
                                 fontSize=9.5, leading=12, spaceAfter=1),
        "entry_title": ParagraphStyle("EntryTitle", parent=styles["Normal"], fontName="Helvetica",
                                      fontSize=10.5, spaceBefore=5, spaceAfter=0, leading=12),
        "entry_dates": ParagraphStyle("EntryDates", parent=styles["Normal"], fontName="Helvetica",
                                      fontSize=9, textColor=PDF_GRAY, spaceAfter=2, leading=11),
        "letter": ParagraphStyle("Letter", parent=styles["Normal"], fontName="Helvetica",
                                 fontSize=10.5, leading=15, spaceAfter=8),
        "small_italic": ParagraphStyle("SmallItalic", parent=styles["Normal"],
                                       fontName="Helvetica-Oblique", fontSize=8.5,
                                       textColor=PDF_GRAY, spaceAfter=2, leading=11),
    }


def _render(path, story, top=1.5, bottom=1.3, left=2.0, right=2.0):
    doc = BaseDocTemplate(path, pagesize=letter,
                          topMargin=top * cm, bottomMargin=bottom * cm,
                          leftMargin=left * cm, rightMargin=right * cm)
    frame = Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id="main")
    doc.addPageTemplates([PageTemplate(id="all", frames=[frame])])
    doc.build(story)
    print(f"OK {os.path.basename(path)}")


def build_pdf_cv():
    S = _styles()

    def rule():
        return HRFlowable(width="100%", thickness=0.6, color=PDF_BLUE, spaceBefore=1, spaceAfter=3)

    def heading(text):
        return [Paragraph(_esc(text.upper()), S["heading"]), rule()]

    def bullets(items):
        return ListFlowable(
            [ListItem(Paragraph(_bold_md(t), S["bullet"]), leftIndent=10, value="•") for t in items],
            bulletType="bullet", bulletColor=PDF_BLUE, leftIndent=12, bulletFontSize=8,
        )

    def entry(dates, title, org, location, bl):
        loc = f", {_esc(location)}" if location else ""
        head = (f'<b>{_esc(title)}</b>'
                f'<font color="#555555">&nbsp;&nbsp;—&nbsp;&nbsp;{_esc(org)}{loc}</font>')
        header = [Paragraph(head, S["entry_title"]), Paragraph(_esc(dates), S["entry_dates"])]
        if not bl:
            return [KeepTogether(header)]
        # Title, dates and first bullet stay together; the rest flows, so a long entry
        # does not jump wholesale to the next page and strand a half-empty one behind it.
        out = [KeepTogether(header + [bullets(bl[:1])])]
        if len(bl) > 1:
            out.append(bullets(bl[1:]))
        return out

    story = [Paragraph(_esc(NAME), S["name"]),
             Paragraph(_esc(CV_CONTACT), S["contact"]),
             Paragraph(_esc(PROFILE), S["body"])]
    story += heading("Core Competencies")
    story.append(bullets([f"**{l}**{r}" for l, r in COMPETENCIES]))
    story += heading("Relevant Experience")
    for e in EXPERIENCE:
        story += entry(*e)
    story += heading("Additional Experience")
    story.append(bullets(ADDITIONAL))
    story += heading("Education")
    for e in EDUCATION:
        story += entry(*e)
    story += heading("Languages")
    story.append(Paragraph("English (native), French (bilingual, professional working proficiency).", S["body"]))
    story += heading("Selected Publications")
    story.append(bullets(PUBLICATIONS))
    story.append(Paragraph(_esc(PUB_FOOT), S["small_italic"]))

    _render("docx/cv_alliance_coordinator.pdf", story)


def build_pdf_cover():
    S = _styles()
    story = [Paragraph(_esc(NAME), S["name"]),
             Paragraph(_esc(COVER_CONTACT), S["contact"]),
             Spacer(1, 6),
             Paragraph(_esc(COVER_DATE), S["body"]),
             Spacer(1, 8),
             Paragraph(f"<b>{_esc(COVER_SUBJECT)}</b>", S["body"]),
             Spacer(1, 4),
             Paragraph(_esc(COVER_SALUTATION), S["letter"])]
    for para in COVER_PARAS:
        story.append(Paragraph(_esc(para), S["letter"]))
    story += [Spacer(1, 6),
              KeepTogether([Paragraph(_esc(COVER_CLOSING), S["body"]),
                            Spacer(1, 12),
                            Paragraph(_esc(NAME), S["body"])])]
    _render("docx/cover_alliance_coordinator.pdf", story, top=2.0, bottom=1.3, left=2.5, right=2.5)


if __name__ == "__main__":
    build_docx_cv()
    build_docx_cover()
    build_pdf_cv()
    build_pdf_cover()
    print("\nDone. Files in docx/")
