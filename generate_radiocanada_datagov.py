"""
Generate application docs for CBC/Radio-Canada - Director, Enterprise Data Governance (T&I).
Montreal (also Toronto/Ottawa), hybrid, permanent, full-time. Seen on LinkedIn 2026-10-02
(jobs 4471114587 and 4469484537). Posting text carries "Unposting Date 2026-09-27" (already
passed) but was reposted 2026-10-01/02; employer's own posting could NOT be fetched
(careers site unreachable), so open status is unverified. No salary or req ID stated.
Requires 5+ years data management/governance, enterprise governance program experience,
catalog tooling, business glossary, national matrixed org, EN/FR bilingualism.
Outputs (in docx/): cv_radiocanada_datagov, cover_radiocanada_datagov,
application_radiocanada_datagov (.docx and .pdf each). English documents.

Content is defined ONCE below, then rendered to both docx (python-docx) and
pdf (reportlab) so the two formats never drift.

GAPS STATED HONESTLY: governance experience is in research, not at a national broadcaster or
enterprise scale; no formal direct reports; hands-on tools on record are Collibra (used at
the centre), SQL, Python, data modeling, ETL; NOT claimed: Alation, Informatica, Purview,
GCP, Jira, business glossary ownership, lineage tooling, steering-committee chairing.

Company fact used (mandate: programming in English and French under the Broadcasting Act)
verified at canada.ca (Canadian Heritage backgrounder, Feb 2025), not from the posting.
"""
import os

os.makedirs("docx", exist_ok=True)

NAME = "Ted Strauss"
CV_CONTACT = "Montreal, Quebec, Canada  |  514-475-2374  |  ted.strauss@gmail.com  |  linkedin.com/in/tedstrauss-63b01821"
COVER_CONTACT = "514-475-2374  |  ted.strauss@gmail.com  |  linkedin.com/in/tedstrauss-63b01821"

PROFILE = (
    "Data governance specialist who builds the policies, metadata standards, and training that make "
    "governance work across organizations that do not report to one boss. Led data governance and "
    "stewardship for five years at McGill's McConnell Brain Imaging Centre for dozens of principal "
    "investigators and hundreds of researchers, and contributed to national data governance practice through "
    "Portage, the Canadian Open Neuroscience Platform (CONP), and the Research Data Alliance. Treats data "
    "literacy and change management as the core of a governance program: communities of practice, "
    "workshops, and self-serve tools that people actually adopt. Hands-on with SQL, data modeling, ETL, "
    "Python, and Collibra. Bilingual (EN/FR), Montreal-based."
)

COMPETENCIES = [
    ("Data governance frameworks, policies, and standards",
     ": defined metadata standards, stewardship workflows, and data policies for a leading brain imaging "
     "centre; contributed to national research data governance through Portage, CONP, and the Research Data Alliance"),
    ("Data catalog, metadata, and governance tooling",
     ": worked with Collibra and applied metadata standards (FAIR principles) to research data; "
     "built data pipelines and a self-serve OwnCloud portal for managed data access"),
    ("Data literacy, training, and change management",
     ": designed and delivered data management workshops, built a public knowledge base and a "
     "Moodle-based education portal, and grew communities of practice; completed CSCCE community engagement training"),
    ("Stakeholder influence across a national, matrixed setting",
     ": coordinated researchers, IT, and policy bodies across institutions on national committees, "
     "translating technical and regulatory concepts for non-specialists"),
    ("Technical foundation",
     ": SQL, NoSQL, ETL, data modeling, Python (pandas, plotly), R, Tableau; AWS and Azure; "
     "LDAP/SSO, Docker, Agile/Scrum"),
    ("AI readiness and adoption",
     ": evaluated and adopted AI tools including Claude Code in enterprise learning contexts"),
    ("Bilingual EN/FR",
     ": English (native), French (professional working proficiency), Montreal-based"),
]

EXPERIENCE = [
    ("09/2016 - 10/2021", "Manager of Data Resources", "McConnell Brain Imaging Centre, The Neuro, McGill University", "Montreal, QC", [
        "Led data governance, metadata standards, and stewardship workflows for a leading brain imaging centre, supporting dozens of principal investigators and hundreds of researchers",
        "Developed brain imaging data pipelines and data policies for the daily data acquisition workflow; tools included Collibra, LDAP, SSO, and R",
        "Grew adoption by treating governance as a literacy problem rather than an enforcement one: a Moodle-based self-serve education portal, hands-on training, and communities of practice",
        "Launched an OwnCloud portal for self-serve data management, making compliant data handling the easy path",
        "Contributed to national research policy through Portage, CONP, and Research Data Alliance committee work, building consensus across institutions",
    ]),
    ("08/2020 - 10/2021", "Data Management and Sharing Coordinator (p/t)", "Tanenbaum Open Science Institute, The Neuro, McGill University", "Montreal, QC", [
        "Designed and administered training workshops on data management practices for students and faculty",
        "Built a public online knowledge base explaining open science concepts to non-specialists",
    ]),
    ("07/2025 - 01/2026", "Moodle Analyst", "Premergency Inc.", "Canada", [
        "Coordinated between client and development teams on LMS solutions for government clients, maintaining regulatory compliance requirements; migrated legacy LMS systems to Moodle",
    ]),
    ("01/2024 - 10/2024", "Full Stack Developer", "KnowledgeOne", "Montreal, QC", [
        "Developed learning platform features for active clients (LAMP++, Moodle, LXD); evaluated and adopted AI tools including Claude Code in an enterprise learning context",
    ]),
]

ADDITIONAL = [
    "McGill Centre for Integrative Neuroscience - Software Developer (2015-2016): data acquisition pipeline for the national CCNA longitudinal dementia study (PHP, MySQL, JavaScript)",
    "Sur Place Media - Founder and Director (2017-2021): founded and ran a neighbourhood arts school; 150 workshop groups with 20 teachers, pandemic pivot to virtual delivery",
    "Photonic Knowledge - Database Administrator and Software Developer (2012-2014): Drupal ERP system for core sample imaging",
]

EDUCATION = [
    ("1998-2003", "Bachelor of Arts (Honours), Cognitive Science", "McGill University", "Montreal, QC",
     ["Minor: Music Technology. Topics: psycholinguistics, cognitive science, computational modeling of language."]),
    ("2021", "Scientific Community Engagement Fundamentals", "The Center for Scientific Collaboration and Community Engagement (CSCCE)", "", []),
]

PUBLICATIONS = [
    "Poline, J.-B., et al. (2023). Data and tools integration in the Canadian open neuroscience platform. Scientific Data, 10(1), 189.",
    "Halchenko, Y. O., et al. (2021). DataLad: distributed system for joint management of code, data, and their relationship. Journal of Open Source Software, 6(63), 3262.",
    "Strauss, T. (2021). Data Management Plan Template: Neuroimaging in the Neurosciences. Zenodo. doi:10.5281/zenodo.4673558",
]
PUB_FOOT = "9 publications | citation count 630 | h-index 9 | Full list: scholar.google.com/citations"

COVER_DATE = "October 5, 2026"
COVER_SUBJECT = "Subject: Application: Director, Enterprise Data Governance (T&I)"
COVER_SALUTATION = "Dear Hiring Team,"
COVER_PARAS = [
    "I am applying for the Director, Enterprise Data Governance position in Technology & Infrastructure. "
    "The role is to stand up an enterprise governance program: policies and standards, a data catalog and "
    "metadata practice, a data quality program, the steering committee and data council, and the "
    "data literacy and change management that make people use them. That combination of governance and "
    "adoption is the work I did for five years as Manager of Data Resources at McGill's McConnell Brain "
    "Imaging Centre.",

    "CBC/Radio-Canada's mandate under the Broadcasting Act is to serve audiences in both English and French, "
    "so a governance program there has to work across a national organization in both official languages. "
    "Your posting describes a network of data owners, stewards, and domain leads "
    "who do not all report to the director. I have worked that way throughout my career: at the centre I "
    "supported dozens of principal investigators and hundreds of researchers, and I:",

    "My governance track record was built in research rather than broadcasting, and the parts that "
    "transfer are the ones this role depends on: I have coordinated multi-stakeholder delivery and led "
    "through influence, though I would be building a formal governance team for the first time. On "
    "tooling, I have worked with Collibra, SQL, data modeling, ETL, and Python; Alation, Purview, and "
    "Google Cloud data tooling would be new to me, and I would approach an enterprise catalog rollout by "
    "pairing my standards and adoption experience with strong engineering partners. What I bring is the "
    "discipline underneath the tools: clear ownership, documented standards, and the ability to explain a "
    "control to an engineer and to an executive so that it is followed. For AI readiness, I have "
    "evaluated and adopted AI tools, including Claude Code, in enterprise learning contexts.",

    "I am bilingual in French and English, based in Montreal, and available now, and the hybrid "
    "arrangement suits me. I would welcome the opportunity to discuss how I would help build "
    "CBC/Radio-Canada's data governance program. Thank you for your consideration.",
]
COVER_BULLETS = [
    "defined metadata standards, stewardship workflows, and data policies for the daily data acquisition workflow",
    "ran workshops, a public knowledge base, a Moodle-based education portal, and communities of practice so that governance was adopted, not just published",
    "launched an OwnCloud self-serve portal for data management and built the brain imaging data pipelines",
    "built consensus across institutions through Portage, CONP, and Research Data Alliance committee work",
]
COVER_CLOSING = "Kind regards,"

# ==================================================================
# DOCX RENDERING
# ==================================================================
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BLUE = RGBColor(0x21, 0x6F, 0xAA)
BLACK = RGBColor(0x00, 0x00, 0x00)
GRAY = RGBColor(0x55, 0x55, 0x55)


def _margins(doc, top=1.5, bottom=1.5, left=2.0, right=2.0):
    for s in doc.sections:
        s.top_margin, s.bottom_margin = Cm(top), Cm(bottom)
        s.left_margin, s.right_margin = Cm(left), Cm(right)


def _name_header(doc, name, contact):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(name); r.bold = True; r.font.size = Pt(28); r.font.color.rgb = BLUE
    p2 = doc.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(contact); r2.font.size = Pt(10); r2.font.color.rgb = GRAY
    p2.paragraph_format.space_after = Pt(6)


def _heading(doc, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title.upper()); r.bold = True; r.font.size = Pt(11); r.font.color.rgb = BLUE
    pPr = p._p.get_or_add_pPr(); pBdr = OxmlElement('w:pBdr'); bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single'); bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1'); bottom.set(qn('w:color'), '216FAA')
    pBdr.append(bottom); pPr.append(pBdr)


def _entry(doc, dates, title, org, location, bullets):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    r1 = p.add_run(title); r1.bold = True; r1.font.size = Pt(10.5); r1.font.color.rgb = BLACK
    r2 = p.add_run(f"  |  {org}"); r2.font.size = Pt(10.5); r2.font.color.rgb = GRAY
    if location:
        r3 = p.add_run(f", {location}"); r3.font.size = Pt(10.5); r3.font.color.rgb = GRAY
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_before = Pt(0); p2.paragraph_format.space_after = Pt(2)
    p2.paragraph_format.keep_with_next = True
    r4 = p2.add_run(dates); r4.italic = True; r4.font.size = Pt(9.5); r4.font.color.rgb = GRAY
    for b in bullets:
        _bullet(doc, b)


def _bullet(doc, text, indent=0.25):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(1); p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.left_indent = Inches(indent)
    for i, part in enumerate(text.split('**')):
        if not part:
            continue
        r = p.add_run(part); r.font.size = Pt(10); r.bold = (i % 2 == 1)


def _para(doc, text, size=10, italic=False, color=None, space_after=4):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(text); r.font.size = Pt(size); r.italic = italic
    if color:
        r.font.color.rgb = color
    return p


def build_docx_cv(doc=None):
    standalone = doc is None
    if standalone:
        doc = Document(); _margins(doc)
    _name_header(doc, NAME, CV_CONTACT)
    _para(doc, PROFILE, size=10, space_after=6)
    _heading(doc, "Core Competencies")
    for label, rest in COMPETENCIES:
        _bullet(doc, f"**{label}**{rest}")
    _heading(doc, "Professional Experience")
    for e in EXPERIENCE:
        _entry(doc, *e)
    _heading(doc, "Additional Experience")
    for a in ADDITIONAL:
        _bullet(doc, a)
    _heading(doc, "Education")
    for e in EDUCATION:
        _entry(doc, *e)
    _heading(doc, "Languages")
    _para(doc, "English (native), French (bilingual, professional working proficiency).", size=10)
    _heading(doc, "Selected Publications")
    for pub in PUBLICATIONS:
        _bullet(doc, pub)
    _para(doc, PUB_FOOT, size=9, italic=True, color=GRAY)
    _heading(doc, "References")
    _para(doc, "Available upon request.", size=10)
    if standalone:
        doc.save("docx/cv_radiocanada_datagov.docx")
        print("OK cv_radiocanada_datagov.docx")


def build_docx_cover(doc=None):
    standalone = doc is None
    if standalone:
        doc = Document(); _margins(doc, top=2.0, bottom=2.0, left=2.5, right=2.5)
    _name_header(doc, NAME, COVER_CONTACT)
    _para(doc, COVER_DATE, size=10, space_after=12)
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(6)
    r = p.add_run(COVER_SUBJECT); r.bold = True; r.font.size = Pt(10)
    _para(doc, COVER_SALUTATION, size=10, space_after=6)
    for i, para in enumerate(COVER_PARAS):
        pp = doc.add_paragraph(); pp.paragraph_format.space_after = Pt(8)
        rr = pp.add_run(para); rr.font.size = Pt(10)
        if i == 1:
            for b in COVER_BULLETS:
                _bullet(doc, b)
    _para(doc, COVER_CLOSING, size=10, space_after=24)
    _para(doc, NAME, size=10)
    if standalone:
        doc.save("docx/cover_radiocanada_datagov.docx")
        print("OK cover_radiocanada_datagov.docx")


# ==================================================================
# PDF RENDERING
# ==================================================================
from reportlab.lib.pagesizes import letter
from reportlab.lib.units import cm
from reportlab.lib.colors import HexColor
from reportlab.lib.enums import TA_CENTER
from reportlab.platypus import (
    BaseDocTemplate, PageTemplate, Frame, Paragraph, Spacer,
    HRFlowable, ListFlowable, ListItem, KeepTogether,
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

PDF_BLUE = HexColor("#216FAA")
PDF_GRAY = HexColor("#555555")


def _esc(t):
    return t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _bold_md(t):
    t = _esc(t)
    parts = t.split("**")
    out = ""
    for i, p in enumerate(parts):
        out += f"<b>{p}</b>" if i % 2 == 1 else p
    return out


def _styles():
    styles = getSampleStyleSheet()
    return {
        "name": ParagraphStyle("Name", parent=styles["Normal"], fontName="Helvetica-Bold",
                               fontSize=22, textColor=PDF_BLUE, alignment=TA_CENTER, spaceAfter=2, leading=26),
        "contact": ParagraphStyle("Contact", parent=styles["Normal"], fontName="Helvetica",
                                  fontSize=9, textColor=PDF_GRAY, alignment=TA_CENTER, spaceAfter=8, leading=12),
        "heading": ParagraphStyle("Heading", parent=styles["Normal"], fontName="Helvetica-Bold",
                                  fontSize=11, textColor=PDF_BLUE, spaceBefore=8, spaceAfter=2, leading=13),
        "body": ParagraphStyle("Body", parent=styles["Normal"], fontName="Helvetica", fontSize=10,
                               leading=13, spaceAfter=3),
        "bullet": ParagraphStyle("Bullet", parent=styles["Normal"], fontName="Helvetica",
                                 fontSize=9.5, leading=12, spaceAfter=1),
        "entry_title": ParagraphStyle("EntryTitle", parent=styles["Normal"], fontName="Helvetica",
                                      fontSize=10.5, spaceBefore=5, spaceAfter=0, leading=12),
        "entry_dates": ParagraphStyle("EntryDates", parent=styles["Normal"], fontName="Helvetica",
                                      fontSize=9, textColor=PDF_GRAY, spaceAfter=2, leading=11),
        "letter": ParagraphStyle("Letter", parent=styles["Normal"], fontName="Helvetica",
                                 fontSize=10.5, leading=14, spaceAfter=7),
        "small_italic": ParagraphStyle("SmallItalic", parent=styles["Normal"],
                                       fontName="Helvetica-Oblique", fontSize=8.5,
                                       textColor=PDF_GRAY, spaceAfter=2, leading=11),
    }


def _render(path, story, top=1.5, bottom=1.3, left=2.0, right=2.0):
    doc = BaseDocTemplate(path, pagesize=letter,
                          topMargin=top * cm, bottomMargin=bottom * cm,
                          leftMargin=left * cm, rightMargin=right * cm)
    frame = Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id="main")
    doc.addPageTemplates([PageTemplate(id="all", frames=[frame])])
    doc.build(story)
    print(f"OK {os.path.basename(path)}")


def build_pdf_cv(story_only=False):
    S = _styles()

    def rule():
        return HRFlowable(width="100%", thickness=0.6, color=PDF_BLUE, spaceBefore=1, spaceAfter=3)

    def heading(text):
        return [Paragraph(_esc(text.upper()), S["heading"]), rule()]

    def bullets(items):
        return ListFlowable(
            [ListItem(Paragraph(_bold_md(t), S["bullet"]), leftIndent=10, value="•") for t in items],
            bulletType="bullet", bulletColor=PDF_BLUE, leftIndent=12, bulletFontSize=8,
        )

    def entry(dates, title, org, location, bl):
        loc = f", {_esc(location)}" if location else ""
        head = (f'<b>{_esc(title)}</b>'
                f'<font color="#555555">&nbsp;&nbsp;|&nbsp;&nbsp;{_esc(org)}{loc}</font>')
        header = [Paragraph(head, S["entry_title"]), Paragraph(_esc(dates), S["entry_dates"])]
        if not bl:
            return [KeepTogether(header)]
        if len(bl) <= 2:
            return [KeepTogether(header + [bullets(bl)])]
        out = [KeepTogether(header + [bullets(bl[:1])])]
        if len(bl) > 1:
            out.append(bullets(bl[1:]))
        return out

    story = [Paragraph(_esc(NAME), S["name"]),
             Paragraph(_esc(CV_CONTACT), S["contact"]),
             Paragraph(_esc(PROFILE), S["body"])]
    story += heading("Core Competencies")
    story.append(bullets([f"**{l}**{r}" for l, r in COMPETENCIES]))
    story += heading("Professional Experience")
    for e in EXPERIENCE:
        story += entry(*e)
    story += heading("Additional Experience")
    story.append(bullets(ADDITIONAL))
    story += heading("Education")
    for e in EDUCATION:
        story += entry(*e)
    story += heading("Languages")
    story.append(Paragraph("English (native), French (bilingual, professional working proficiency).", S["body"]))
    story += heading("Selected Publications")
    story.append(bullets(PUBLICATIONS))
    story.append(Paragraph(_esc(PUB_FOOT), S["small_italic"]))
    story += heading("References")
    story.append(Paragraph("Available upon request.", S["body"]))

    if story_only:
        return story
    _render("docx/cv_radiocanada_datagov.pdf", story)


def build_pdf_cover(story_only=False):
    S = _styles()
    story = [Paragraph(_esc(NAME), S["name"]),
             Paragraph(_esc(COVER_CONTACT), S["contact"]),
             Spacer(1, 6),
             Paragraph(_esc(COVER_DATE), S["body"]),
             Spacer(1, 8),
             Paragraph(f"<b>{_esc(COVER_SUBJECT)}</b>", S["body"]),
             Spacer(1, 4),
             Paragraph(_esc(COVER_SALUTATION), S["letter"])]
    for i, para in enumerate(COVER_PARAS):
        story.append(Paragraph(_esc(para), S["letter"]))
        if i == 1:
            story.append(ListFlowable(
                [ListItem(Paragraph(_esc(t), S["bullet"]), leftIndent=10, value="•") for t in COVER_BULLETS],
                bulletType="bullet", bulletColor=PDF_BLUE, leftIndent=12, bulletFontSize=8,
            ))
            story.append(Spacer(1, 4))
    story += [Spacer(1, 2),
              KeepTogether([Paragraph(_esc(COVER_CLOSING), S["body"]),
                            Spacer(1, 10),
                            Paragraph(_esc(NAME), S["body"])])]
    if story_only:
        return story
    _render("docx/cover_radiocanada_datagov.pdf", story, top=1.8, bottom=1.1, left=2.5, right=2.5)


# ==================================================================
# COMBINED (cover letter + CV in one document, for a single upload)
# ==================================================================
def build_docx_combined():
    from docx.enum.section import WD_SECTION
    doc = Document(); _margins(doc, top=2.0, bottom=2.0, left=2.5, right=2.5)
    build_docx_cover(doc)
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    sec.top_margin, sec.bottom_margin = Cm(1.5), Cm(1.5)
    sec.left_margin, sec.right_margin = Cm(2.0), Cm(2.0)
    build_docx_cv(doc)
    doc.save("docx/application_radiocanada_datagov.docx")
    print("OK application_radiocanada_datagov.docx")


def build_pdf_combined():
    from reportlab.platypus import NextPageTemplate, PageBreak
    cover = build_pdf_cover(story_only=True)
    cv = build_pdf_cv(story_only=True)
    path = "docx/application_radiocanada_datagov.pdf"
    doc = BaseDocTemplate(path, pagesize=letter)
    f_cover = Frame(2.5 * cm, 1.1 * cm, letter[0] - 5.0 * cm, letter[1] - 2.9 * cm, id="cover")
    f_cv = Frame(2.0 * cm, 1.3 * cm, letter[0] - 4.0 * cm, letter[1] - 2.8 * cm, id="cv")
    doc.addPageTemplates([PageTemplate(id="cover", frames=[f_cover]),
                          PageTemplate(id="cv", frames=[f_cv])])
    doc.build(cover + [NextPageTemplate("cv"), PageBreak()] + cv)
    print("OK application_radiocanada_datagov.pdf")


if __name__ == "__main__":
    build_docx_cv()
    build_docx_cover()
    build_docx_combined()
    build_pdf_cv()
    build_pdf_cover()
    build_pdf_combined()
    print()
    print("Done. Files in docx/")
