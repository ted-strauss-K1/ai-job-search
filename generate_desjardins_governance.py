"""
Generate application docs for Desjardins — Senior Advisor, Data Governance.
Montreal, hybrid. Full-time, regular. Req R2610157.
Outputs (in docx/):
  - cv_desjardins_governance.docx        / .pdf   (English, per CV-language policy)
  - cover_desjardins_governance_en.docx  / .pdf
  - cover_desjardins_governance_fr.docx  / .pdf

Content is defined ONCE below, then rendered to both docx (python-docx) and
pdf (reportlab) so formats never drift. Two cover letter language versions
were requested by the candidate; the CV stays English per the profile's
fixed CV-language policy.

Targeting note: leads with governance leadership (McConnell Brain Imaging
Centre / Neuro, McGill) rather than the LMS/L&D track, since this posting is
a pure data-governance advisory role. Genuine gaps stated plainly: no
financial-sector or cooperative regulatory experience (AMF/OSFI), and the
Neuro governance tenure (5 years) sits just under the posting's 6-year bar —
both acknowledged in the cover letter, the second bridged using the
posting's own "alternative combinations may be considered" language.

Company facts used in the cover letter were verified independently against
desjardins.com (mission/values page): "largest cooperative financial group
in Canada", democratic one-member-one-vote governance, 10+ million members,
189+ caisses, money framed as a tool for member/community autonomy. NOT
used: anything from the DGIG (insurance subsidiary) board governance PDF
found during research — it describes P&C insurance board committees, not
enterprise data governance, and is not relevant to this role.
"""
import os

os.makedirs("docx", exist_ok=True)

# ------------------------------------------------------------------
# SHARED CONTENT
# ------------------------------------------------------------------

NAME = "Ted Strauss"
CV_CONTACT = "Montreal, Quebec, Canada  |  514-475-2374  |  ted.strauss@gmail.com  |  linkedin.com/in/tedstrauss-63b01821"
COVER_CONTACT = "514-475-2374  |  ted.strauss@gmail.com  |  linkedin.com/in/tedstrauss-63b01821"

PROFILE = (
    "Data governance specialist with direct experience building the standards, controls, and advisory "
    "relationships that let organizations manage data responsibly. Led data governance for one of "
    "Canada's leading research data ecosystems at the McConnell Brain Imaging Centre (McGill University): "
    "defined metadata and stewardship policies, built compliance mechanisms into daily workflows, and "
    "represented the institution on national data governance bodies (Portage, CONP, RDA). Brings a rare "
    "combination of technical grounding (LAMP++, devops, data platforms) and an advisory, "
    "relationship-first approach to governance maturity. Bilingual (EN/FR), Montreal-based."
)

# (bold label, remaining text)
COMPETENCIES = [
    ("Data governance frameworks and standards",
     ": defined metadata standards, stewardship policies, and lifecycle rules (classification, "
     "retention, access) for a national-scale research data environment; built governance mechanisms "
     "that let dozens of research groups manage data responsibly without a central bottleneck"),
    ("Advisory-first governance approach",
     ": grew adoption by treating data governance as a literacy problem rather than an enforcement one, "
     "building communities of practice and hands-on training instead of issuing rules; the same posture "
     "this role's mandate calls for when guiding data producers and users"),
    ("Compliance and controls management",
     ": maintained regulatory compliance requirements across government LMS deployments; developed and "
     "monitored data quality and stewardship controls tied to institutional and national policy "
     "commitments"),
    ("Cross-functional stakeholder collaboration",
     ": coordinated across research, clinical, IT, and national policy stakeholders (Portage, CONP, RDA) "
     "who did not report to me, translating governance requirements into practices teams could actually "
     "follow"),
    ("AI adoption and governance",
     ": evaluated and adopted AI tooling in enterprise and regulated contexts, including Claude Code, "
     "weighing data handling, confidentiality, and governance implications alongside the productivity "
     "case"),
    ("Bilingual communication in a Quebec cooperative context",
     ": professional working proficiency in French and English; Montreal-based with a decade of "
     "experience navigating Quebec's research, public-sector, and cooperative institutions"),
]

# each: (dates, title, org, location, [bullets]) — ordered by relevance to governance, not chronology
EXPERIENCE = [
    ("09/2016 – 10/2021", "Manager of Data Resources", "McConnell Brain Imaging Centre, The Neuro, McGill University", "Montreal, QC", [
        "Led data governance for one of Canada's leading brain imaging research centres, defining metadata standards and lifecycle policies covering classification, retention, and access for dozens of principal investigators and hundreds of researchers",
        "Built compliance and data quality controls into daily acquisition and stewardship workflows rather than issuing them as separate policy documents",
        "Grew governance adoption as a literacy problem, not a compliance one, building communities of practice through hands-on training",
        "Represented the centre on national research data governance bodies (Portage, CONP, RDA), contributing to shared Canadian data governance standards",
        "Launched a self-serve OwnCloud portal and a Moodle-based education portal, making governance-compliant data management the path of least resistance for researchers",
    ]),
    ("08/2020 – 10/2021", "Data Management & Sharing Coordinator (p/t)", "Tanenbaum Open Science Institute, The Neuro, McGill University", "Montreal, QC", [
        "Designed and administered training workshops on data governance and management practices for faculty and graduate students",
        "Built a public knowledge base translating governance and open science requirements into practical guidance for non-specialists",
    ]),
    ("07/2025 – 01/2026", "Moodle Analyst", "Premergency Inc.", "Canada", [
        "Maintained regulatory compliance requirements across government LMS deployments, coordinating between client and development teams",
        "Migrated legacy systems and designed solutions against documented capability gaps rather than a fixed catalogue",
    ]),
    ("01/2024 – 10/2024", "Full Stack Developer", "KnowledgeOne", "Montreal, QC", [
        "Evaluated and adopted AI tools in enterprise contexts, assessing data handling and governance implications alongside the productivity case",
    ]),
]

ADDITIONAL = [
    "McGill Centre for Integrative Neuroscience — Software Developer (2015–2016): built the data acquisition pipeline for the national CCNA longitudinal dementia study (PHP, MySQL, JavaScript)",
    "Sur Place Media — Founder and Director (2017–2021): founded and ran a neighbourhood arts school, owning program operations for 150 workshop groups and 20 instructors",
    "Photonic Knowledge — Database Administrator and Software Developer (2012–2014): built a Drupal ERP system and its underlying data model for the mining exploration sector",
]

EDUCATION = [
    ("1998–2003", "B.A. (Honours), Cognitive Science", "McGill University", "Montreal, QC",
     ["Minor: Music Technology. Topics: psycholinguistics, cognitive science, computational modeling of language, music technology."]),
    ("2021", "Scientific Community Engagement Fundamentals", "CSCCE", "", []),
]

PUBLICATIONS = [
    "Poline, J.-B., et al. (2023). Data and tools integration in the Canadian open neuroscience platform. Scientific Data, 10(1), 189.",
    "Strauss, T. (2021). Data Management Plan Template: Neuroimaging in the Neurosciences. Zenodo. doi:10.5281/zenodo.4673558",
]
PUB_FOOT = "9 publications | citation count 630 | h-index 9 | Full list: scholar.google.com/citations"

# ------------------------------------------------------------------
# COVER LETTER — ENGLISH
# ------------------------------------------------------------------

COVER_DATE_EN = "August 14, 2026"
COVER_SUBJECT_EN = "Subject: Application: Senior Advisor, Data Governance (R2610157)"
COVER_SALUTATION_EN = "Dear Hiring Team,"
COVER_PARAS_EN = [
    "I am applying for the Senior Advisor, Data Governance position (R2610157) in Montreal. The posting "
    "asks for someone who builds data governance day to day: standards, controls, and advisory "
    "relationships with the people who actually produce and use data, not a policy written and left on "
    "a shelf. That is the work I led at the McConnell Brain Imaging Centre, one of Canada's leading "
    "brain imaging research centres, for five years.",

    "Desjardins describes its mission as putting money to work for the autonomy and development of "
    "members and communities, governed one member, one vote across more than 189 caisses serving over "
    "ten million members. A cooperative that answers to that many member-owners needs data governance "
    "that earns trust rather than one that simply enforces rules from above. That advisory posture, "
    "guiding producers and users toward better data practices instead of policing them, is how I have "
    "approached governance work throughout my career, and it is what your posting describes when it "
    "asks for someone to guide data producers and users in data valorization.",

    "At the McConnell Brain Imaging Centre, I managed data governance for dozens of principal "
    "investigators and hundreds of researchers:",

    "Desjardins' scale and cooperative structure are new to me, but the underlying discipline is not, "
    "and I want to be direct about the two places where my background does not line up cleanly with "
    "your posting. My governance leadership was built inside a research institution, so the regulatory "
    "context specific to a financial cooperative (AMF, OSFI) is new ground. And my five years as "
    "Manager of Data Resources sit just under your six-year bar on their own; taken together with the "
    "governance-adjacent compliance work I did afterward, and your posting's note that alternative "
    "combinations of qualifications and experience are considered, I believe the case holds.",

    "I am bilingual in French and English at professional working proficiency and based in Montreal. I "
    "would welcome the opportunity to discuss how I would contribute to Desjardins' data governance "
    "maturity. Thank you for your consideration.",
]
COVER_BULLETS_EN = [
    "Defined metadata standards and lifecycle policies (classification, retention, access) for a research data environment supporting dozens of PIs and hundreds of researchers",
    "Built data quality and compliance controls into daily workflows, then measured and communicated governance maturity to institutional leadership",
    "Represented the centre on national data governance bodies, Portage, CONP, and the Research Data Alliance, contributing to shared Canadian standards",
    "Evaluated and adopted AI tooling, including Claude Code, weighing data handling and governance implications alongside the productivity case",
]
COVER_CLOSING_EN = "Kind regards,"

# ------------------------------------------------------------------
# COVER LETTER — FRENCH
# ------------------------------------------------------------------

COVER_DATE_FR = "14 août 2026"
COVER_SUBJECT_FR = "Objet : Candidature au poste de conseiller principal, gouvernance des données (R2610157)"
COVER_SALUTATION_FR = "Madame, Monsieur,"
COVER_PARAS_FR = [
    "Je pose ma candidature au poste de conseiller principal, gouvernance des données (R2610157) à "
    "Montréal. Votre affichage recherche une personne capable de bâtir la gouvernance des données au "
    "quotidien : normes, contrôles et relations de conseil avec les personnes qui produisent et "
    "utilisent réellement les données, plutôt qu'une politique rédigée puis laissée de côté. C'est le "
    "travail que j'ai dirigé pendant cinq ans au Centre d'imagerie cérébrale McConnell, l'un des "
    "principaux centres de recherche en imagerie cérébrale au Canada.",

    "Desjardins définit sa mission comme la mise en valeur de l'argent au service de l'autonomie et du "
    "développement des membres et des communautés, selon une gouvernance démocratique, un membre, un "
    "vote, répartie dans plus de 189 caisses au service de plus de dix millions de membres. Une "
    "coopérative qui rend des comptes à autant de membres-propriétaires a besoin d'une gouvernance des "
    "données qui gagne la confiance plutôt que d'imposer des règles d'en haut. Cette approche de conseil, "
    "guider plutôt que surveiller, est celle que j'ai toujours privilégiée dans mon travail de "
    "gouvernance, et c'est exactement ce que votre affichage demande.",

    "Au Centre d'imagerie cérébrale McConnell, j'ai géré la gouvernance des données pour des dizaines de "
    "chercheurs principaux et des centaines de chercheurs :",

    "Je tiens à être direct sur deux points où mon parcours ne correspond pas parfaitement à votre "
    "affichage. Mon leadership en gouvernance s'est construit dans un contexte de recherche, donc le "
    "cadre réglementaire propre à une coopérative financière (AMF, BSIF) est un terrain nouveau pour "
    "moi. Et mes cinq ans comme responsable des ressources de données se situent tout juste sous votre "
    "exigence de six ans; combinés à mon expérience connexe en conformité par la suite, et à la "
    "flexibilité que prévoit votre affichage sur les combinaisons équivalentes, je crois que le dossier "
    "tient la route.",

    "Je suis bilingue en français et en anglais, avec une compétence professionnelle complète dans les "
    "deux langues, et basé à Montréal. Je serais heureux de discuter de la façon dont je pourrais "
    "contribuer à la maturité de la gouvernance des données de Desjardins. Merci de considérer ma "
    "candidature.",
]
COVER_BULLETS_FR = [
    "Défini les normes de métadonnées et les politiques de cycle de vie (classification, conservation, accès) pour un environnement de données de recherche soutenant des dizaines de chercheurs principaux et des centaines de chercheurs",
    "Intégré les contrôles de qualité et de conformité des données directement dans les flux de travail quotidiens, puis mesuré et communiqué la maturité de la gouvernance à la direction",
    "Représenté le centre auprès d'organismes nationaux de gouvernance des données, Portage, le CONP et la Research Data Alliance, contribuant à des normes canadiennes partagées",
    "Évalué et adopté des outils d'intelligence artificielle, dont Claude Code, en pesant les enjeux de traitement des données et de gouvernance au même titre que les gains de productivité",
]
COVER_CLOSING_FR = "Veuillez agréer, Madame, Monsieur, mes salutations distinguées,"

# ==================================================================
# DOCX RENDERING
# ==================================================================
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BLUE = RGBColor(0x21, 0x6F, 0xAA)
BLACK = RGBColor(0x00, 0x00, 0x00)
GRAY = RGBColor(0x55, 0x55, 0x55)


def _margins(doc, top=1.5, bottom=1.5, left=2.0, right=2.0):
    for s in doc.sections:
        s.top_margin, s.bottom_margin = Cm(top), Cm(bottom)
        s.left_margin, s.right_margin = Cm(left), Cm(right)


def _name_header(doc, name, contact):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(name); r.bold = True; r.font.size = Pt(28); r.font.color.rgb = BLUE
    p2 = doc.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(contact); r2.font.size = Pt(10); r2.font.color.rgb = GRAY
    p2.paragraph_format.space_after = Pt(6)


def _heading(doc, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title.upper()); r.bold = True; r.font.size = Pt(11); r.font.color.rgb = BLUE
    pPr = p._p.get_or_add_pPr(); pBdr = OxmlElement('w:pBdr'); bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single'); bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1'); bottom.set(qn('w:color'), '216FAA')
    pBdr.append(bottom); pPr.append(pBdr)


def _entry(doc, dates, title, org, location, bullets):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    r1 = p.add_run(title); r1.bold = True; r1.font.size = Pt(10.5); r1.font.color.rgb = BLACK
    r2 = p.add_run(f"  —  {org}"); r2.font.size = Pt(10.5); r2.font.color.rgb = GRAY
    if location:
        r3 = p.add_run(f", {location}"); r3.font.size = Pt(10.5); r3.font.color.rgb = GRAY
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_before = Pt(0); p2.paragraph_format.space_after = Pt(2)
    p2.paragraph_format.keep_with_next = True
    r4 = p2.add_run(dates); r4.italic = True; r4.font.size = Pt(9.5); r4.font.color.rgb = GRAY
    for b in bullets:
        _bullet(doc, b)


def _bullet(doc, text, indent=0.25):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(1); p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.left_indent = Inches(indent)
    for i, part in enumerate(text.split('**')):
        if not part:
            continue
        r = p.add_run(part); r.font.size = Pt(10); r.bold = (i % 2 == 1)


def _para(doc, text, size=10, italic=False, color=None, space_after=4):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(text); r.font.size = Pt(size); r.italic = italic
    if color:
        r.font.color.rgb = color
    return p


def build_docx_cv():
    doc = Document(); _margins(doc)
    _name_header(doc, NAME, CV_CONTACT)
    _para(doc, PROFILE, size=10, space_after=6)
    _heading(doc, "Core Competencies")
    for label, rest in COMPETENCIES:
        _bullet(doc, f"**{label}**{rest}")
    _heading(doc, "Relevant Experience")
    for e in EXPERIENCE:
        _entry(doc, *e)
    _heading(doc, "Additional Experience")
    for a in ADDITIONAL:
        _bullet(doc, a)
    _heading(doc, "Education")
    for e in EDUCATION:
        _entry(doc, *e)
    _heading(doc, "Languages")
    _para(doc, "English (native), French (bilingual, professional working proficiency).", size=10)
    _heading(doc, "Selected Publications")
    for pub in PUBLICATIONS:
        _bullet(doc, pub)
    _para(doc, PUB_FOOT, size=9, italic=True, color=GRAY)
    doc.save("docx/cv_desjardins_governance.docx")
    print("OK cv_desjardins_governance.docx")


def build_docx_cover(lang):
    date, subject, salutation, paras, bullets, closing = _cover_fields(lang)
    doc = Document(); _margins(doc, top=2.0, bottom=2.0, left=2.5, right=2.5)
    _name_header(doc, NAME, COVER_CONTACT)
    _para(doc, date, size=10, space_after=12)
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(6)
    r = p.add_run(subject); r.bold = True; r.font.size = Pt(10)
    _para(doc, salutation, size=10, space_after=6)
    for i, para in enumerate(paras):
        pp = doc.add_paragraph(); pp.paragraph_format.space_after = Pt(8)
        rr = pp.add_run(para); rr.font.size = Pt(10)
        if i == 2:  # paragraph introducing the bullet list
            for b in bullets:
                _bullet(doc, b)
    _para(doc, closing, size=10, space_after=24)
    _para(doc, NAME, size=10)
    doc.save(f"docx/cover_desjardins_governance_{lang}.docx")
    print(f"OK cover_desjardins_governance_{lang}.docx")


def _cover_fields(lang):
    if lang == "en":
        return COVER_DATE_EN, COVER_SUBJECT_EN, COVER_SALUTATION_EN, COVER_PARAS_EN, COVER_BULLETS_EN, COVER_CLOSING_EN
    return COVER_DATE_FR, COVER_SUBJECT_FR, COVER_SALUTATION_FR, COVER_PARAS_FR, COVER_BULLETS_FR, COVER_CLOSING_FR


# ==================================================================
# PDF RENDERING
# ==================================================================
from reportlab.lib.pagesizes import letter
from reportlab.lib.units import cm
from reportlab.lib.colors import HexColor
from reportlab.lib.enums import TA_CENTER
from reportlab.platypus import (
    BaseDocTemplate, PageTemplate, Frame, Paragraph, Spacer,
    HRFlowable, ListFlowable, ListItem, KeepTogether,
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

PDF_BLUE = HexColor("#216FAA")
PDF_GRAY = HexColor("#555555")


def _esc(t):
    return t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _bold_md(t):
    t = _esc(t)
    parts = t.split("**")
    out = ""
    for i, p in enumerate(parts):
        out += f"<b>{p}</b>" if i % 2 == 1 else p
    return out


def _styles():
    styles = getSampleStyleSheet()
    return {
        "name": ParagraphStyle("Name", parent=styles["Normal"], fontName="Helvetica-Bold",
                               fontSize=22, textColor=PDF_BLUE, alignment=TA_CENTER, spaceAfter=2, leading=26),
        "contact": ParagraphStyle("Contact", parent=styles["Normal"], fontName="Helvetica",
                                  fontSize=9, textColor=PDF_GRAY, alignment=TA_CENTER, spaceAfter=8, leading=12),
        "heading": ParagraphStyle("Heading", parent=styles["Normal"], fontName="Helvetica-Bold",
                                  fontSize=11, textColor=PDF_BLUE, spaceBefore=8, spaceAfter=2, leading=13),
        "body": ParagraphStyle("Body", parent=styles["Normal"], fontName="Helvetica", fontSize=10,
                               leading=13, spaceAfter=3),
        "bullet": ParagraphStyle("Bullet", parent=styles["Normal"], fontName="Helvetica",
                                 fontSize=9.5, leading=12, spaceAfter=1),
        "entry_title": ParagraphStyle("EntryTitle", parent=styles["Normal"], fontName="Helvetica",
                                      fontSize=10.5, spaceBefore=5, spaceAfter=0, leading=12),
        "entry_dates": ParagraphStyle("EntryDates", parent=styles["Normal"], fontName="Helvetica",
                                      fontSize=9, textColor=PDF_GRAY, spaceAfter=2, leading=11),
        "letter": ParagraphStyle("Letter", parent=styles["Normal"], fontName="Helvetica",
                                 fontSize=10.5, leading=14, spaceAfter=7),
        "small_italic": ParagraphStyle("SmallItalic", parent=styles["Normal"],
                                       fontName="Helvetica-Oblique", fontSize=8.5,
                                       textColor=PDF_GRAY, spaceAfter=2, leading=11),
    }


def _render(path, story, top=1.5, bottom=1.3, left=2.0, right=2.0):
    doc = BaseDocTemplate(path, pagesize=letter,
                          topMargin=top * cm, bottomMargin=bottom * cm,
                          leftMargin=left * cm, rightMargin=right * cm)
    frame = Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id="main")
    doc.addPageTemplates([PageTemplate(id="all", frames=[frame])])
    doc.build(story)
    print(f"OK {os.path.basename(path)}")


def build_pdf_cv():
    S = _styles()

    def rule():
        return HRFlowable(width="100%", thickness=0.6, color=PDF_BLUE, spaceBefore=1, spaceAfter=3)

    def heading(text):
        return [Paragraph(_esc(text.upper()), S["heading"]), rule()]

    def bullets(items):
        return ListFlowable(
            [ListItem(Paragraph(_bold_md(t), S["bullet"]), leftIndent=10, value="•") for t in items],
            bulletType="bullet", bulletColor=PDF_BLUE, leftIndent=12, bulletFontSize=8,
        )

    def entry(dates, title, org, location, bl):
        loc = f", {_esc(location)}" if location else ""
        head = (f'<b>{_esc(title)}</b>'
                f'<font color="#555555">&nbsp;&nbsp;—&nbsp;&nbsp;{_esc(org)}{loc}</font>')
        header = [Paragraph(head, S["entry_title"]), Paragraph(_esc(dates), S["entry_dates"])]
        if not bl:
            return [KeepTogether(header)]
        if len(bl) <= 2:
            return [KeepTogether(header + [bullets(bl)])]
        out = [KeepTogether(header + [bullets(bl[:1])])]
        if len(bl) > 1:
            out.append(bullets(bl[1:]))
        return out

    story = [Paragraph(_esc(NAME), S["name"]),
             Paragraph(_esc(CV_CONTACT), S["contact"]),
             Paragraph(_esc(PROFILE), S["body"])]
    story += heading("Core Competencies")
    story.append(bullets([f"**{l}**{r}" for l, r in COMPETENCIES]))
    story += heading("Relevant Experience")
    for e in EXPERIENCE:
        story += entry(*e)
    story += heading("Additional Experience")
    story.append(bullets(ADDITIONAL))
    story += heading("Education")
    for e in EDUCATION:
        story += entry(*e)
    story += heading("Languages")
    story.append(Paragraph("English (native), French (bilingual, professional working proficiency).", S["body"]))
    story += heading("Selected Publications")
    story.append(bullets(PUBLICATIONS))
    story.append(Paragraph(_esc(PUB_FOOT), S["small_italic"]))

    _render("docx/cv_desjardins_governance.pdf", story)


def build_pdf_cover(lang):
    date, subject, salutation, paras, bullet_items, closing = _cover_fields(lang)
    S = _styles()
    story = [Paragraph(_esc(NAME), S["name"]),
             Paragraph(_esc(COVER_CONTACT), S["contact"]),
             Spacer(1, 6),
             Paragraph(_esc(date), S["body"]),
             Spacer(1, 8),
             Paragraph(f"<b>{_esc(subject)}</b>", S["body"]),
             Spacer(1, 4),
             Paragraph(_esc(salutation), S["letter"])]
    for i, para in enumerate(paras):
        story.append(Paragraph(_esc(para), S["letter"]))
        if i == 2:
            story.append(ListFlowable(
                [ListItem(Paragraph(_esc(t), S["bullet"]), leftIndent=10, value="•") for t in bullet_items],
                bulletType="bullet", bulletColor=PDF_BLUE, leftIndent=12, bulletFontSize=8,
            ))
            story.append(Spacer(1, 4))
    story += [Spacer(1, 2),
              KeepTogether([Paragraph(_esc(closing), S["body"]),
                            Spacer(1, 10),
                            Paragraph(_esc(NAME), S["body"])])]
    _render(f"docx/cover_desjardins_governance_{lang}.pdf", story, top=1.8, bottom=1.1, left=2.5, right=2.5)


if __name__ == "__main__":
    build_docx_cv()
    build_docx_cover("en")
    build_docx_cover("fr")
    build_pdf_cv()
    build_pdf_cover("en")
    build_pdf_cover("fr")
    print("\nDone. Files in docx/")
