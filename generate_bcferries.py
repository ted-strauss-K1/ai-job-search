"""
Generate application docs for BC Ferries — Lead Data Governance Specialist (Job 4518).
Location: Victoria / Vancouver Lower Mainland. Salary $132,600-$165,700.
Outputs (in docx/):
  - cv_bcferries_datagov.docx
  - cover_bcferries_datagov.docx
  - bcferries_datagov_combined.pdf   (cover letter + CV in one PDF, as the posting requires)

Content is defined ONCE below, then rendered to both docx (python-docx) and
pdf (reportlab) so the two formats never drift.
"""
import os

os.makedirs("docx", exist_ok=True)

# ------------------------------------------------------------------
# SHARED CONTENT
# ------------------------------------------------------------------

NAME = "Ted Strauss"
CV_CONTACT = "Montreal, Quebec, Canada  |  514-475-2374  |  ted.strauss@gmail.com  |  linkedin.com/in/tedstrauss-63b01821"
COVER_CONTACT = "514-475-2374  |  ted.strauss@gmail.com  |  linkedin.com/in/tedstrauss-63b01821"

PROFILE = (
    "Data governance specialist with over a decade of experience designing and operationalizing "
    "data stewardship programs across Canada's research and public-sector ecosystem. Led the "
    "governance program at McGill University's McConnell Brain Imaging Centre, defining metadata "
    "standards, data quality measures, retention and access policies, and stewardship structures "
    "for dozens of research groups. Shaped national data governance policy through Portage, CONP, "
    "and RDA committee work. Known for making governance practical: building a data-driven culture "
    "through stakeholder engagement and change management rather than compliance enforcement. "
    "Bilingual (EN/FR); 9 peer-reviewed publications; citation count 630; h-index 9."
)

# (bold label, remaining text) — label is emphasised
COMPETENCIES = [
    ("Enterprise governance frameworks",
     ": designed and maintained policies and standards for data quality, metadata, classification, "
     "retention, privacy, and stewardship across a multi-group research environment; authored data "
     "acquisition and handling policies for high-throughput workflows"),
    ("Governance operating model",
     ": established stewardship structures, role assignments, and issue-resolution workflows; deployed "
     "Collibra for institutional data cataloguing and compliance tracking; ran governance as an ongoing "
     "practice with continuous improvement rather than a one-time project"),
    ("Stakeholder engagement & change leadership",
     ": built communities of practice and embedded data literacy into daily workflows; cultivated "
     "trusted relationships across teams to drive adoption; positioned governance as an enabler, not an "
     "operational burden on the business"),
    ("Enterprise & national influence",
     ": represented institutional interests on national data governance bodies (Portage, CONP, RDA); "
     "co-authored data sharing and stewardship frameworks adopted across Canadian institutions; aligned "
     "governance activity with strategic priorities across many stakeholders"),
    ("Regulatory & public-sector compliance",
     ": managed regulatory compliance for government LMS deployments; developed data handling policies "
     "meeting national research and privacy standards"),
    ("Data infrastructure & tooling",
     ": Collibra, OwnCloud, LAMP++ (Linux, Apache/Nginx, MySQL/PostgreSQL, PHP, JavaScript), SSO, LDAP, "
     "Active Directory, Docker, cloud (AWS, Azure); R, Python (pandas, plotly), Tableau; evaluated and "
     "adopted AI tools including Claude Code for enterprise workflows"),
]

# each: (dates, title, org, location, [bullets])
EXPERIENCE = [
    ("07/2025 – 01/2026", "Moodle Analyst", "Premergency Inc.", "Canada", [
        "Coordinated between government clients and development teams to deliver regulated LMS solutions, maintaining compliance obligations throughout; managed stakeholder expectations across client and technical teams",
        "Managed legacy LMS migrations to Moodle; designed bespoke learning & development solutions for public-sector clients",
    ]),
    ("01/2024 – 10/2024", "Full Stack Developer", "KnowledgeOne", "Montreal, QC", [
        "Developed LMS features for active clients on the eConcordia Moodle platform (LAMP++, LXD)",
        "Evaluated and adopted AI tools in enterprise learning contexts, assessing governance implications alongside usability and performance",
    ]),
    ("09/2016 – 10/2021", "Manager of Data Resources", "McConnell Brain Imaging Centre, The Neuro, McGill University", "Montreal, QC", [
        "Led the data governance program for one of Canada's leading brain imaging centres, defining metadata standards, data ownership and stewardship structures, and quality measures for dozens of principal investigators and hundreds of researchers",
        "Designed and enforced data policies covering acquisition, classification, retention, and access across daily neuroimaging pipelines",
        "Deployed Collibra for institutional data cataloguing and governance tracking; launched an OwnCloud self-serve portal; administered databases across SQL and NoSQL systems",
        "Drove adoption through a Moodle-based education portal and hands-on training, building data literacy and a culture of stewardship rather than top-down compliance",
        "Represented the centre on national data governance bodies (Portage, CONP, RDA); co-authored frameworks adopted across Canadian institutions",
    ]),
    ("08/2020 – 10/2021", "Data Management & Sharing Coordinator (p/t)", "Tanenbaum Open Science Institute, The Neuro, McGill University", "Montreal, QC", [
        "Supported the launch of a new academic unit focused on open science; coordinated across faculty, students, and institutional stakeholders",
        "Designed and delivered training on data management and governance practices; built a public-facing knowledge base on FAIR data principles",
    ]),
    ("04/2017 – 08/2021", "Founder and Director", "Sur Place Media", "Montreal, QC", [
        "Founded and operated a neighbourhood arts school; hosted 150 workshop groups with 20 instructors; led a full pivot to virtual delivery during the COVID-19 pandemic",
    ]),
    ("02/2015 – 09/2016", "Software Developer", "McGill Centre for Integrative Neuroscience, McGill University", "Montreal, QC", [
        "Programmed the data acquisition pipeline for the national CCNA longitudinal dementia study using the LORIS platform (PHP, MySQL, JavaScript); contributed to national research data infrastructure",
    ]),
]

ADDITIONAL = [
    "Photonic Knowledge — Database administrator & software developer (2012–2014)",
    "Language Mate — Developer and UX designer (2007–2010)",
    "Moss Rehabilitation Hospital — Research software developer (2007–2010)",
    "University of Connecticut & Columbia University — Research assistant & software developer (2004–2008)",
]

EDUCATION = [
    ("1998–2003", "B.A. (Honours), Cognitive Science", "McGill University", "Montreal, QC",
     ["Minor: Music Technology. Topics: Psycholinguistics, cognitive science, computational modeling of language."]),
    ("2021", "Scientific Community Engagement Fundamentals", "CSCCE", "", []),
]

PUBLICATIONS = [
    "Poline, J.-B., et al. (2023). Data and tools integration in the Canadian open neuroscience platform. Scientific Data, 10(1), 189.",
    "Halchenko, Y. O., et al. (2021). DataLad: distributed system for joint management of code, data, and their relationship. JOSS, 6(63), 3262.",
    "Strauss, T. (2021). Data Management Plan Template: Neuroimaging in the Neurosciences. Zenodo. doi:10.5281/zenodo.4673558",
]
PUB_FOOT = "Citation count: 630 | h-index: 9 | Full list: scholar.google.com/citations"

COVER_DATE = "July 15, 2026"
COVER_SUBJECT = "Subject: Application: Lead Data Governance Specialist"
COVER_SALUTATION = "Dear BC Ferries Hiring Team,"
COVER_PARAS = [
    "I am applying for the Lead Data Governance Specialist role. Your posting names the challenge I have "
    "spent my career on: bringing data governance and the business together, and putting practices in "
    "place that are effective without becoming an operational burden. That is exactly how I have "
    "approached governance across Canada's research and public-sector ecosystem.",

    "At McGill University's McConnell Brain Imaging Centre, I led the data governance program for one of "
    "Canada's leading neuroimaging environments. I defined metadata standards, data quality measures, and "
    "retention and access policies; established stewardship structures and role assignments across dozens "
    "of research groups; and deployed Collibra for institutional cataloguing and compliance tracking. The "
    "recurring obstacle was never the technology, it was adoption. I built governance as a culture of data "
    "literacy, cultivating trust with teams and embedding practices into daily workflows so they held up "
    "in practice rather than only on paper. That is the change-agent work your posting describes.",

    "I have also operated above the single-institution level. Representing our centre on national data "
    "governance bodies, Portage, CONP, and RDA, I helped author data sharing and stewardship frameworks "
    "adopted across Canadian institutions, aligning governance activity with strategic priorities across "
    "many stakeholders. Establishing a Data Governance Centre of Excellence, partnering with data "
    "stewards, domain owners, and a governance council to grow organizational capability, is a natural "
    "extension of that work.",

    "More recently at Premergency Inc., I coordinated regulated LMS deployments for government clients, "
    "maintaining compliance obligations throughout, so I understand governance in a public-sector, "
    "accountability-driven context. I would bring to BC Ferries a governance approach that pairs rigorous "
    "standards for quality, metadata, classification, retention, and privacy with the stakeholder "
    "engagement that makes those standards hold.",

    "I would welcome the opportunity to discuss how I can help operationalize and mature enterprise data "
    "governance at BC Ferries. Thank you for your consideration.",
]
COVER_CLOSING = "Kind regards,"

# ==================================================================
# DOCX RENDERING
# ==================================================================
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BLUE = RGBColor(0x21, 0x6F, 0xAA)
BLACK = RGBColor(0x00, 0x00, 0x00)
GRAY = RGBColor(0x55, 0x55, 0x55)


def _margins(doc, top=1.5, bottom=1.5, left=2.0, right=2.0):
    for s in doc.sections:
        s.top_margin, s.bottom_margin = Cm(top), Cm(bottom)
        s.left_margin, s.right_margin = Cm(left), Cm(right)


def _name_header(doc, name, contact):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(name); r.bold = True; r.font.size = Pt(28); r.font.color.rgb = BLUE
    p2 = doc.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(contact); r2.font.size = Pt(10); r2.font.color.rgb = GRAY
    p2.paragraph_format.space_after = Pt(6)


def _heading(doc, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title.upper()); r.bold = True; r.font.size = Pt(11); r.font.color.rgb = BLUE
    pPr = p._p.get_or_add_pPr(); pBdr = OxmlElement('w:pBdr'); bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single'); bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1'); bottom.set(qn('w:color'), '216FAA')
    pBdr.append(bottom); pPr.append(pBdr)


def _entry(doc, dates, title, org, location, bullets):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(0)
    r1 = p.add_run(title); r1.bold = True; r1.font.size = Pt(10.5); r1.font.color.rgb = BLACK
    r2 = p.add_run(f"  —  {org}"); r2.font.size = Pt(10.5); r2.font.color.rgb = GRAY
    if location:
        r3 = p.add_run(f", {location}"); r3.font.size = Pt(10.5); r3.font.color.rgb = GRAY
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_before = Pt(0); p2.paragraph_format.space_after = Pt(2)
    r4 = p2.add_run(dates); r4.italic = True; r4.font.size = Pt(9.5); r4.font.color.rgb = GRAY
    for b in bullets:
        _bullet(doc, b)


def _bullet(doc, text, indent=0.25):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(1); p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.left_indent = Inches(indent)
    for i, part in enumerate(text.split('**')):
        if not part:
            continue
        r = p.add_run(part); r.font.size = Pt(10); r.bold = (i % 2 == 1)


def _para(doc, text, size=10, italic=False, color=None, space_after=4):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(text); r.font.size = Pt(size); r.italic = italic
    if color:
        r.font.color.rgb = color
    return p


def build_docx_cv():
    doc = Document(); _margins(doc)
    _name_header(doc, NAME, CV_CONTACT)
    _para(doc, PROFILE, size=10, space_after=6)
    _heading(doc, "Core Competencies")
    for label, rest in COMPETENCIES:
        _bullet(doc, f"**{label}**{rest}")
    _heading(doc, "Professional Experience")
    for e in EXPERIENCE:
        _entry(doc, *e)
    _heading(doc, "Additional Experience")
    for a in ADDITIONAL:
        _bullet(doc, a)
    _heading(doc, "Education")
    for e in EDUCATION:
        _entry(doc, *e)
    _heading(doc, "Languages")
    _para(doc, "English (native), French (bilingual).", size=10)
    _heading(doc, "Selected Publications")
    for pub in PUBLICATIONS:
        _bullet(doc, pub)
    _para(doc, PUB_FOOT, size=9, italic=True, color=GRAY)
    _heading(doc, "References")
    _para(doc, "Available upon request.", size=10)
    doc.save("docx/cv_bcferries_datagov.docx")
    print("OK cv_bcferries_datagov.docx")


def build_docx_cover():
    doc = Document(); _margins(doc, top=2.0, bottom=2.0, left=2.5, right=2.5)
    _name_header(doc, NAME, COVER_CONTACT)
    _para(doc, COVER_DATE, size=10, space_after=12)
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(6)
    r = p.add_run(COVER_SUBJECT); r.bold = True; r.font.size = Pt(10)
    _para(doc, COVER_SALUTATION, size=10, space_after=6)
    for para in COVER_PARAS:
        pp = doc.add_paragraph(); pp.paragraph_format.space_after = Pt(8)
        rr = pp.add_run(para); rr.font.size = Pt(10)
    _para(doc, COVER_CLOSING, size=10, space_after=24)
    _para(doc, NAME, size=10)
    doc.save("docx/cover_bcferries_datagov.docx")
    print("OK cover_bcferries_datagov.docx")


# ==================================================================
# PDF RENDERING (combined: cover letter, then CV)
# ==================================================================
from reportlab.lib.pagesizes import letter
from reportlab.lib.units import cm
from reportlab.lib.colors import HexColor
from reportlab.lib.enums import TA_CENTER
from reportlab.platypus import (
    BaseDocTemplate, PageTemplate, Frame, Paragraph, Spacer, PageBreak,
    HRFlowable, ListFlowable, ListItem,
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

PDF_BLUE = HexColor("#216FAA")
PDF_GRAY = HexColor("#555555")


def _esc(t):
    return t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _bold_md(t):
    # convert **x** to <b>x</b> after escaping
    t = _esc(t)
    parts = t.split("**")
    out = ""
    for i, p in enumerate(parts):
        out += f"<b>{p}</b>" if i % 2 == 1 else p
    return out


def build_pdf():
    styles = getSampleStyleSheet()
    name_style = ParagraphStyle("Name", parent=styles["Normal"], fontName="Helvetica-Bold",
                                fontSize=22, textColor=PDF_BLUE, alignment=TA_CENTER, spaceAfter=2, leading=26)
    contact_style = ParagraphStyle("Contact", parent=styles["Normal"], fontName="Helvetica",
                                   fontSize=9, textColor=PDF_GRAY, alignment=TA_CENTER, spaceAfter=8, leading=12)
    heading_style = ParagraphStyle("Heading", parent=styles["Normal"], fontName="Helvetica-Bold",
                                   fontSize=11, textColor=PDF_BLUE, spaceBefore=8, spaceAfter=2, leading=13)
    body = ParagraphStyle("Body", parent=styles["Normal"], fontName="Helvetica", fontSize=10,
                          leading=13, spaceAfter=3)
    bullet_style = ParagraphStyle("Bullet", parent=body, fontSize=9.5, leading=12, spaceAfter=1)
    entry_title = ParagraphStyle("EntryTitle", parent=body, fontSize=10.5, spaceBefore=5, spaceAfter=0, leading=12)
    entry_dates = ParagraphStyle("EntryDates", parent=body, fontSize=9, textColor=PDF_GRAY,
                                 spaceAfter=2, leading=11)
    letter_body = ParagraphStyle("Letter", parent=body, fontSize=10.5, leading=15, spaceAfter=8)
    small_italic = ParagraphStyle("SmallItalic", parent=body, fontSize=8.5, textColor=PDF_GRAY,
                                  fontName="Helvetica-Oblique", spaceAfter=2)

    def rule():
        return HRFlowable(width="100%", thickness=0.6, color=PDF_BLUE, spaceBefore=1, spaceAfter=3)

    def heading(text):
        return [Paragraph(_esc(text.upper()), heading_style), rule()]

    def bullets(items, style=bullet_style):
        return ListFlowable(
            [ListItem(Paragraph(_bold_md(t), style), leftIndent=10, value="•") for t in items],
            bulletType="bullet", bulletColor=PDF_BLUE, leftIndent=12, bulletFontSize=8,
        )

    def entry(dates, title, org, location, bl):
        loc = f", {_esc(location)}" if location else ""
        head = (f'<b>{_esc(title)}</b>'
                f'<font color="#555555">&nbsp;&nbsp;—&nbsp;&nbsp;{_esc(org)}{loc}</font>')
        flow = [Paragraph(head, entry_title), Paragraph(_esc(dates), entry_dates)]
        if bl:
            flow.append(bullets(bl))
        return flow

    story = []

    # ---- COVER LETTER ----
    story.append(Paragraph(_esc(NAME), name_style))
    story.append(Paragraph(_esc(COVER_CONTACT), contact_style))
    story.append(Spacer(1, 6))
    story.append(Paragraph(_esc(COVER_DATE), body))
    story.append(Spacer(1, 8))
    story.append(Paragraph(f"<b>{_esc(COVER_SUBJECT)}</b>", body))
    story.append(Spacer(1, 4))
    story.append(Paragraph(_esc(COVER_SALUTATION), letter_body))
    for para in COVER_PARAS:
        story.append(Paragraph(_esc(para), letter_body))
    story.append(Spacer(1, 10))
    story.append(Paragraph(_esc(COVER_CLOSING), body))
    story.append(Spacer(1, 20))
    story.append(Paragraph(_esc(NAME), body))

    story.append(PageBreak())

    # ---- CV ----
    story.append(Paragraph(_esc(NAME), name_style))
    story.append(Paragraph(_esc(CV_CONTACT), contact_style))
    story.append(Paragraph(_esc(PROFILE), body))
    story += heading("Core Competencies")
    story.append(bullets([f"**{l}**{r}" for l, r in COMPETENCIES]))
    story += heading("Professional Experience")
    for e in EXPERIENCE:
        story += entry(*e)
    story += heading("Additional Experience")
    story.append(bullets(ADDITIONAL))
    story += heading("Education")
    for e in EDUCATION:
        story += entry(*e)
    story += heading("Languages")
    story.append(Paragraph("English (native), French (bilingual).", body))
    story += heading("Selected Publications")
    story.append(bullets(PUBLICATIONS))
    story.append(Paragraph(_esc(PUB_FOOT), small_italic))
    story += heading("References")
    story.append(Paragraph("Available upon request.", body))

    doc = BaseDocTemplate("docx/bcferries_datagov_combined.pdf", pagesize=letter,
                          topMargin=1.5 * cm, bottomMargin=1.3 * cm,
                          leftMargin=2.0 * cm, rightMargin=2.0 * cm)
    frame = Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id="main")
    doc.addPageTemplates([PageTemplate(id="all", frames=[frame])])
    doc.build(story)
    print("OK bcferries_datagov_combined.pdf")


if __name__ == "__main__":
    build_docx_cv()
    build_docx_cover()
    build_pdf()
    print("\nDone. Files in docx/")
