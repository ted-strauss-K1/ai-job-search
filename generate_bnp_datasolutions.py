"""
Generate application docs for BNP Paribas — Data Solutions Manager, Ref 111114839114008.
CIB Americas, Montreal. Permanent. Hybrid, 3 days/week on site (one must be Mon or Fri).
Outputs (in docx/):
  - cv_bnp_datasolutions.docx     / .pdf
  - cover_bnp_datasolutions.docx  / .pdf

Content is defined ONCE below, then rendered to both docx (python-docx) and
pdf (reportlab) so the two formats never drift.

Targeting note: strong match on data governance / architecture / ETL, executive
communication, and building an analytics community (their Data & Digital Week and
Hackathon). Genuine gaps, handled honestly and never hidden: no Master's degree, no
SAS/Alteryx/Power BI, and no ML/clustering/PCA modelling work. The posting explicitly
invites applicants who do not meet the full list.
"""
import os

os.makedirs("docx", exist_ok=True)

# ------------------------------------------------------------------
# SHARED CONTENT
# ------------------------------------------------------------------

NAME = "Ted Strauss"
CV_CONTACT = "Montreal, Quebec, Canada  |  514-475-2374  |  ted.strauss@gmail.com  |  linkedin.com/in/tedstrauss-63b01821"
COVER_CONTACT = "514-475-2374  |  ted.strauss@gmail.com  |  linkedin.com/in/tedstrauss-63b01821"

PROFILE = (
    "Data governance and analytics leader with over a decade of experience building the frameworks, "
    "pipelines, and working relationships that make institutional data trustworthy enough to decide "
    "on. Led data governance for McGill University's McConnell Brain Imaging Centre, defining "
    "metadata standards, data definitions, quality expectations, and retention and access policies "
    "across dozens of research groups, and managing the acquisition pipelines behind them. Built and "
    "ran an organisation of 20 instructors, and contributed to national data standards through "
    "Portage, CONP, and RDA. Known for turning technical work into decisions non-specialists can act "
    "on. Bilingual (EN/FR), Montreal-based."
)

# (bold label, remaining text) — label is emphasised
COMPETENCIES = [
    ("Data governance, architecture and ETL",
     ": defined and maintained governance frameworks covering metadata standards, data definitions, "
     "classification, retention, access, and quality; designed and ran ETL and acquisition pipelines "
     "for high-volume daily workflows; used Collibra as the institutional data catalogue"),
    ("Analytics and visual reporting",
     ": SQL and NoSQL, Python (pandas, plotly), R, and Tableau; data modelling, wrangling, and "
     "cleaning; built reporting and dashboards that let non-technical stakeholders act on findings "
     "without a translator in the room"),
    ("Leading and developing people",
     ": managed a data resources function serving dozens of principal investigators and hundreds of "
     "researchers; recruited, scheduled, and supported a team of 20 instructors as founder and "
     "director; mentored colleagues into data practice rather than issuing rules"),
    ("Executive and cross-functional communication",
     ": briefed institutional leadership and national committees; aligned technical, operational, and "
     "policy stakeholders who did not report to one another; wrote for audiences ranging from "
     "engineers to funders to the general public"),
    ("Building analytics communities and events",
     ": designed and delivered training programmes, workshops, and a Moodle-based self-serve learning "
     "portal that drove adoption; ran 150 workshop groups across in-person and virtual formats; "
     "certified in Scientific Community Engagement Fundamentals (CSCCE)"),
    ("Tool adoption and continuous learning",
     ": evaluated and adopted new tooling in regulated and enterprise settings, including AI tools "
     "such as Claude Code, assessing data handling and governance implications alongside "
     "productivity; comfortable with rapid prototyping and iterative delivery in Agile teams"),
]

# each: (dates, title, org, location, [bullets])
EXPERIENCE = [
    ("09/2016 – 10/2021", "Manager of Data Resources", "McConnell Brain Imaging Centre, The Neuro, McGill University", "Montreal, QC", [
        "Led data governance for one of Canada's leading brain imaging centres, defining metadata standards, data definitions, ownership structures, and quality expectations across dozens of principal investigators and hundreds of researchers",
        "Designed and maintained the governance, architecture, and ETL processes behind daily acquisition workflows, and wrote the policies covering classification, retention, and access",
        "Used Collibra as the institutional data catalogue; maintained the glossary, metadata fields, and naming conventions; launched an OwnCloud portal that replaced manual data-request handling with self-service",
        "Drove adoption through a Moodle-based education portal and hands-on training, building data literacy and communities of practice rather than top-down compliance",
        "Represented the centre on national research data bodies (Portage, CONP, RDA); authored a published data management plan template and co-authored the Canadian Open Neuroscience Platform integration paper (Scientific Data, 2023)",
    ]),
    ("07/2025 – 01/2026", "Moodle Analyst", "Premergency Inc.", "Canada", [
        "Coordinated between government clients and development teams to deliver regulated learning platforms, maintaining compliance obligations and aligning expectations across business and technical stakeholders",
    ]),
    ("04/2017 – 08/2021", "Founder and Director", "Sur Place Media", "Montreal, QC", [
        "Founded and ran an arts school, recruiting and supporting a team of 20 instructors and delivering 150 workshop groups; led a full pivot to virtual delivery during the pandemic",
        "Owned the operating side end to end: budgeting, invoicing, expense tracking, purchasing, scheduling, and records",
    ]),
    ("01/2024 – 10/2024", "Full Stack Developer", "KnowledgeOne", "Montreal, QC", [
        "Developed learning platform features for active clients (LAMP++, Moodle, LXD); evaluated and adopted AI tools in enterprise learning contexts",
    ]),
    ("08/2020 – 10/2021", "Data Management & Sharing Coordinator (p/t)", "Tanenbaum Open Science Institute, The Neuro, McGill University", "Montreal, QC", [
        "Designed and administered data management training for students and faculty; built a public knowledge base explaining data concepts to non-specialists",
    ]),
    ("02/2015 – 09/2016", "Software Developer", "McGill Centre for Integrative Neuroscience, McGill University", "Montreal, QC", [
        "Programmed the data acquisition pipeline for the national CCNA longitudinal dementia study on the LORIS platform (PHP, MySQL, JavaScript)",
    ]),
]

ADDITIONAL = [
    "Photonic Knowledge — Database administrator and software developer (2012–2014): built a Drupal ERP system and its underlying data model for the mining exploration sector",
    "Moss Rehabilitation Hospital — Research software developer (2007–2010)",
    "University of Connecticut & Columbia University — Research assistant and software developer (2004–2008)",
]

EDUCATION = [
    ("1998–2003", "B.A. (Honours), Cognitive Science", "McGill University", "Montreal, QC",
     ["Minor: Music Technology. Topics: psycholinguistics, cognitive science, computational modeling of language, music technology."]),
    ("2021", "Scientific Community Engagement Fundamentals", "CSCCE", "", []),
]

PUBLICATIONS = [
    "Poline, J.-B., et al. (2023). Data and tools integration in the Canadian open neuroscience platform. Scientific Data, 10(1), 189.",
    "Strauss, T. (2021). Data Management Plan Template: Neuroimaging in the Neurosciences. Zenodo. doi:10.5281/zenodo.4673558",
]
PUB_FOOT = "9 publications | citation count 630 | h-index 9 | Full list: scholar.google.com/citations"

COVER_DATE = "August 7, 2026"
COVER_SUBJECT = "Subject: Application: Data Solutions Manager (Reference 111114839114008)"
COVER_SALUTATION = "Dear Hiring Team,"
COVER_PARAS = [
    "I am applying for the Data Solutions Manager position with CIB Americas in Montreal. Your posting "
    "puts data governance, architecture, and ETL frameworks at the centre of the role, then asks for "
    "someone who can turn what comes out of them into something executives will act on, and who can "
    "build an analytics community around the work. That combination is what I have spent my career "
    "doing.",

    "At McGill University's McConnell Brain Imaging Centre I led data governance for one of Canada's "
    "leading research environments. I defined the metadata standards, data definitions, ownership "
    "structures, and quality expectations across dozens of research groups, designed the acquisition "
    "and ETL processes behind daily workflows, and wrote the policies covering classification, "
    "retention, and access. Compliance and data quality were not abstractions there. A definition "
    "that did not hold, or a pipeline that dropped a field, showed up as a result nobody could "
    "reproduce.",

    "The part of your posting I find most interesting is the Data & Digital Week and Hackathon, "
    "because building that kind of community is the work I most want to do more of. I ran adoption at "
    "McGill as a literacy problem rather than a compliance one, building a self-serve learning portal "
    "and running hands-on training, and I separately founded and directed an organisation where I "
    "recruited and supported 20 instructors across 150 workshop groups. I am certified in Scientific "
    "Community Engagement Fundamentals. Championing new tools and rapid prototyping is a natural "
    "extension of that: I have evaluated and adopted AI tooling, including Claude Code, in enterprise "
    "settings, weighing data handling and governance alongside the productivity case.",

    "Two things I want to be straight about, since your posting asks for them. I hold a B.A. "
    "(Honours), not a Master's. And while I work daily in SQL, Python, R, and Tableau, my analytics "
    "practice is governance, pipeline, and reporting work rather than predictive modelling; I have "
    "not built production machine learning, clustering, or PCA models, and I have not used SAS, "
    "Alteryx, or Power BI. I would rather say so now than discover it in an interview. What I bring "
    "instead is the governance, architecture, and ETL half of the posting at depth, the stakeholder "
    "and executive communication that turns analysis into decisions, and a record of building the "
    "kind of analytics culture the role is asked to grow.",

    "I am bilingual in French and English at professional working proficiency, based in Montreal, and "
    "able to meet the on-site requirement. I would welcome the opportunity to discuss how I can "
    "strengthen the data governance foundations under CIB Americas analytics. Thank you for your "
    "consideration.",
]
COVER_CLOSING = "Kind regards,"

# ==================================================================
# DOCX RENDERING
# ==================================================================
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BLUE = RGBColor(0x21, 0x6F, 0xAA)
BLACK = RGBColor(0x00, 0x00, 0x00)
GRAY = RGBColor(0x55, 0x55, 0x55)


def _margins(doc, top=1.5, bottom=1.5, left=2.0, right=2.0):
    for s in doc.sections:
        s.top_margin, s.bottom_margin = Cm(top), Cm(bottom)
        s.left_margin, s.right_margin = Cm(left), Cm(right)


def _name_header(doc, name, contact):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(name); r.bold = True; r.font.size = Pt(28); r.font.color.rgb = BLUE
    p2 = doc.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(contact); r2.font.size = Pt(10); r2.font.color.rgb = GRAY
    p2.paragraph_format.space_after = Pt(6)


def _heading(doc, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title.upper()); r.bold = True; r.font.size = Pt(11); r.font.color.rgb = BLUE
    pPr = p._p.get_or_add_pPr(); pBdr = OxmlElement('w:pBdr'); bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single'); bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1'); bottom.set(qn('w:color'), '216FAA')
    pBdr.append(bottom); pPr.append(pBdr)


def _entry(doc, dates, title, org, location, bullets):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    r1 = p.add_run(title); r1.bold = True; r1.font.size = Pt(10.5); r1.font.color.rgb = BLACK
    r2 = p.add_run(f"  —  {org}"); r2.font.size = Pt(10.5); r2.font.color.rgb = GRAY
    if location:
        r3 = p.add_run(f", {location}"); r3.font.size = Pt(10.5); r3.font.color.rgb = GRAY
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_before = Pt(0); p2.paragraph_format.space_after = Pt(2)
    p2.paragraph_format.keep_with_next = True
    r4 = p2.add_run(dates); r4.italic = True; r4.font.size = Pt(9.5); r4.font.color.rgb = GRAY
    for b in bullets:
        _bullet(doc, b)


def _bullet(doc, text, indent=0.25):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(1); p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.left_indent = Inches(indent)
    for i, part in enumerate(text.split('**')):
        if not part:
            continue
        r = p.add_run(part); r.font.size = Pt(10); r.bold = (i % 2 == 1)


def _para(doc, text, size=10, italic=False, color=None, space_after=4):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(text); r.font.size = Pt(size); r.italic = italic
    if color:
        r.font.color.rgb = color
    return p


def build_docx_cv():
    doc = Document(); _margins(doc)
    _name_header(doc, NAME, CV_CONTACT)
    _para(doc, PROFILE, size=10, space_after=6)
    _heading(doc, "Core Competencies")
    for label, rest in COMPETENCIES:
        _bullet(doc, f"**{label}**{rest}")
    _heading(doc, "Relevant Experience")
    for e in EXPERIENCE:
        _entry(doc, *e)
    _heading(doc, "Additional Experience")
    for a in ADDITIONAL:
        _bullet(doc, a)
    _heading(doc, "Education")
    for e in EDUCATION:
        _entry(doc, *e)
    _heading(doc, "Languages")
    _para(doc, "English (native), French (bilingual, professional working proficiency).", size=10)
    _heading(doc, "Selected Publications")
    for pub in PUBLICATIONS:
        _bullet(doc, pub)
    _para(doc, PUB_FOOT, size=9, italic=True, color=GRAY)
    doc.save("docx/cv_bnp_datasolutions.docx")
    print("OK cv_bnp_datasolutions.docx")


def build_docx_cover():
    doc = Document(); _margins(doc, top=2.0, bottom=2.0, left=2.5, right=2.5)
    _name_header(doc, NAME, COVER_CONTACT)
    _para(doc, COVER_DATE, size=10, space_after=12)
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(6)
    r = p.add_run(COVER_SUBJECT); r.bold = True; r.font.size = Pt(10)
    _para(doc, COVER_SALUTATION, size=10, space_after=6)
    for para in COVER_PARAS:
        pp = doc.add_paragraph(); pp.paragraph_format.space_after = Pt(8)
        rr = pp.add_run(para); rr.font.size = Pt(10)
    _para(doc, COVER_CLOSING, size=10, space_after=24)
    _para(doc, NAME, size=10)
    doc.save("docx/cover_bnp_datasolutions.docx")
    print("OK cover_bnp_datasolutions.docx")


# ==================================================================
# PDF RENDERING
# ==================================================================
from reportlab.lib.pagesizes import letter
from reportlab.lib.units import cm
from reportlab.lib.colors import HexColor
from reportlab.lib.enums import TA_CENTER
from reportlab.platypus import (
    BaseDocTemplate, PageTemplate, Frame, Paragraph, Spacer,
    HRFlowable, ListFlowable, ListItem, KeepTogether,
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

PDF_BLUE = HexColor("#216FAA")
PDF_GRAY = HexColor("#555555")


def _esc(t):
    return t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _bold_md(t):
    t = _esc(t)
    parts = t.split("**")
    out = ""
    for i, p in enumerate(parts):
        out += f"<b>{p}</b>" if i % 2 == 1 else p
    return out


def _styles():
    styles = getSampleStyleSheet()
    return {
        "name": ParagraphStyle("Name", parent=styles["Normal"], fontName="Helvetica-Bold",
                               fontSize=22, textColor=PDF_BLUE, alignment=TA_CENTER, spaceAfter=2, leading=26),
        "contact": ParagraphStyle("Contact", parent=styles["Normal"], fontName="Helvetica",
                                  fontSize=9, textColor=PDF_GRAY, alignment=TA_CENTER, spaceAfter=8, leading=12),
        "heading": ParagraphStyle("Heading", parent=styles["Normal"], fontName="Helvetica-Bold",
                                  fontSize=11, textColor=PDF_BLUE, spaceBefore=8, spaceAfter=2, leading=13),
        "body": ParagraphStyle("Body", parent=styles["Normal"], fontName="Helvetica", fontSize=10,
                               leading=13, spaceAfter=3),
        "bullet": ParagraphStyle("Bullet", parent=styles["Normal"], fontName="Helvetica",
                                 fontSize=9.5, leading=12, spaceAfter=1),
        "entry_title": ParagraphStyle("EntryTitle", parent=styles["Normal"], fontName="Helvetica",
                                      fontSize=10.5, spaceBefore=5, spaceAfter=0, leading=12),
        "entry_dates": ParagraphStyle("EntryDates", parent=styles["Normal"], fontName="Helvetica",
                                      fontSize=9, textColor=PDF_GRAY, spaceAfter=2, leading=11),
        "letter": ParagraphStyle("Letter", parent=styles["Normal"], fontName="Helvetica",
                                 fontSize=10.5, leading=15, spaceAfter=8),
        "small_italic": ParagraphStyle("SmallItalic", parent=styles["Normal"],
                                       fontName="Helvetica-Oblique", fontSize=8.5,
                                       textColor=PDF_GRAY, spaceAfter=2, leading=11),
    }


def _render(path, story, top=1.5, bottom=1.3, left=2.0, right=2.0):
    doc = BaseDocTemplate(path, pagesize=letter,
                          topMargin=top * cm, bottomMargin=bottom * cm,
                          leftMargin=left * cm, rightMargin=right * cm)
    frame = Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id="main")
    doc.addPageTemplates([PageTemplate(id="all", frames=[frame])])
    doc.build(story)
    print(f"OK {os.path.basename(path)}")


def build_pdf_cv():
    S = _styles()

    def rule():
        return HRFlowable(width="100%", thickness=0.6, color=PDF_BLUE, spaceBefore=1, spaceAfter=3)

    def heading(text):
        return [Paragraph(_esc(text.upper()), S["heading"]), rule()]

    def bullets(items):
        return ListFlowable(
            [ListItem(Paragraph(_bold_md(t), S["bullet"]), leftIndent=10, value="•") for t in items],
            bulletType="bullet", bulletColor=PDF_BLUE, leftIndent=12, bulletFontSize=8,
        )

    def entry(dates, title, org, location, bl):
        loc = f", {_esc(location)}" if location else ""
        head = (f'<b>{_esc(title)}</b>'
                f'<font color="#555555">&nbsp;&nbsp;—&nbsp;&nbsp;{_esc(org)}{loc}</font>')
        header = [Paragraph(head, S["entry_title"]), Paragraph(_esc(dates), S["entry_dates"])]
        if not bl:
            return [KeepTogether(header)]
        out = [KeepTogether(header + [bullets(bl[:1])])]
        if len(bl) > 1:
            out.append(bullets(bl[1:]))
        return out

    story = [Paragraph(_esc(NAME), S["name"]),
             Paragraph(_esc(CV_CONTACT), S["contact"]),
             Paragraph(_esc(PROFILE), S["body"])]
    story += heading("Core Competencies")
    story.append(bullets([f"**{l}**{r}" for l, r in COMPETENCIES]))
    story += heading("Relevant Experience")
    for e in EXPERIENCE:
        story += entry(*e)
    story += heading("Additional Experience")
    story.append(bullets(ADDITIONAL))
    story += heading("Education")
    for e in EDUCATION:
        story += entry(*e)
    story += heading("Languages")
    story.append(Paragraph("English (native), French (bilingual, professional working proficiency).", S["body"]))
    story += heading("Selected Publications")
    story.append(bullets(PUBLICATIONS))
    story.append(Paragraph(_esc(PUB_FOOT), S["small_italic"]))

    _render("docx/cv_bnp_datasolutions.pdf", story)


def build_pdf_cover():
    S = _styles()
    story = [Paragraph(_esc(NAME), S["name"]),
             Paragraph(_esc(COVER_CONTACT), S["contact"]),
             Spacer(1, 6),
             Paragraph(_esc(COVER_DATE), S["body"]),
             Spacer(1, 8),
             Paragraph(f"<b>{_esc(COVER_SUBJECT)}</b>", S["body"]),
             Spacer(1, 4),
             Paragraph(_esc(COVER_SALUTATION), S["letter"])]
    for para in COVER_PARAS:
        story.append(Paragraph(_esc(para), S["letter"]))
    story += [Spacer(1, 6),
              KeepTogether([Paragraph(_esc(COVER_CLOSING), S["body"]),
                            Spacer(1, 12),
                            Paragraph(_esc(NAME), S["body"])])]
    _render("docx/cover_bnp_datasolutions.pdf", story, top=2.0, bottom=1.3, left=2.5, right=2.5)


if __name__ == "__main__":
    build_docx_cv()
    build_docx_cover()
    build_pdf_cv()
    build_pdf_cover()
    print("\nDone. Files in docx/")
