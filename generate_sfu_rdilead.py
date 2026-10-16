"""
Generate application docs for Simon Fraser University, Stephens Family School of
Medicine — Research Data and Infrastructure Lead. Surrey BC. Continuing, APSA
grade 12, $99,865 - $119,148. Hybrid-eligible.
Outputs (in docx/):
  - cv_sfu_rdilead.docx     / .pdf
  - cover_sfu_rdilead.docx  / .pdf

Content is defined ONCE below, then rendered to both docx (python-docx) and
pdf (reportlab) so the two formats never drift.

Targeting note: founding-stage research data governance + infrastructure +
analytics for a brand-new medical school. CV anchors on The Neuro (governance
frameworks, portal/repository build, pipelines), national ecosystem work, and the
training/mentoring dimension the posting emphasizes. CCNA dementia study appears
for the health-data angle; Kitikmeot Heritage Society for Indigenous-context
collaboration. OCAP acknowledged as a gap in the letter, never claimed.
"""
import os

os.makedirs("docx", exist_ok=True)

# ------------------------------------------------------------------
# SHARED CONTENT
# ------------------------------------------------------------------

NAME = "Ted Strauss"
CV_CONTACT = "Montreal, Quebec, Canada  |  514-475-2374  |  ted.strauss@gmail.com  |  linkedin.com/in/tedstrauss-63b01821"
COVER_CONTACT = "514-475-2374  |  ted.strauss@gmail.com  |  linkedin.com/in/tedstrauss-63b01821"

PROFILE = (
    "Research data leader with five years building data governance, infrastructure, and stewardship "
    "for a major academic research centre, McGill's McConnell Brain Imaging Centre, and a published "
    "contributor to Canada's national research data platforms (Canadian Open Neuroscience Platform, "
    "DataLad). Designed governance frameworks and metadata standards, launched self-serve data "
    "systems, built acquisition-to-archive pipelines, and paired every system with the training and "
    "mentoring that made researchers adopt it. Committee experience with Portage, CONP, and RDA. "
    "Active research profile: 9 publications, 630 citations. Bilingual (EN/FR), prepared to relocate "
    "to Metro Vancouver."
)

COMPETENCIES = [
    ("Research data governance",
     ": developed and implemented data governance frameworks, metadata standards, data quality "
     "practices, data sharing policy, and stewardship workflows for a leading brain imaging centre "
     "handling sensitive human research data under research ethics and privacy requirements"),
    ("Research data infrastructure and repositories",
     ": designed and launched an OwnCloud self-serve data management portal replacing manual "
     "request handling; built data pipelines spanning the research data lifecycle from acquisition "
     "through storage and archiving; contributor to the Canadian Open Neuroscience Platform, a "
     "national platform integrating distributed neuroscience data repositories"),
    ("Analytics and reporting",
     ": Python and R for data analysis (pandas, plotly), Tableau dashboards, SQL through a decade "
     "of LAMP-stack (MySQL) development; software developer on the national CCNA dementia study, "
     "a pan-Canadian linked health research dataset"),
    ("Mentoring, training, and data literacy",
     ": designed and delivered research data management training for students and faculty at "
     "McGill's Tanenbaum Open Science Institute; built a Moodle-based self-serve education portal "
     "and a public knowledge base on open science; certified in Scientific Community Engagement "
     "Fundamentals (CSCCE)"),
    ("Cross-institutional collaboration",
     ": represented a research centre on national bodies (Portage Network, integrated into the "
     "Digital Research Alliance of Canada in 2021; CONP; Research Data Alliance), sustaining "
     "multi-institutional relationships and contributing to national research data policy"),
    ("Platform engineering and AI adoption",
     ": LAMP-stack development, AWS and Azure, Docker and Kubernetes, SSO/LDAP, Git and devops; "
     "evaluated and adopted AI tooling including Claude Code in enterprise learning contexts"),
]

EXPERIENCE = [
    ("09/2016 – 10/2021", "Manager of Data Resources", "McConnell Brain Imaging Centre, The Neuro, McGill University", "Montreal, QC", [
        "Led data governance, metadata standards, and stewardship workflows for one of Canada's leading brain imaging centres, supporting dozens of principal investigators and hundreds of researchers working with sensitive human data",
        "Designed and launched an OwnCloud portal that moved researchers to self-serve data management, replacing manual data-request handling and removing an administrative bottleneck",
        "Built data pipelines supporting the centre's daily acquisition workflows, from instrument output through storage and archiving",
        "Represented the centre on national research data bodies (Portage, CONP, RDA), sustaining multi-institutional relationships and contributing to national research data policy",
        "Designed a Moodle-based self-serve education portal for researcher onboarding, and delivered hands-on training to build data literacy",
    ]),
    ("08/2020 – 10/2021", "Data Management & Sharing Coordinator (p/t)", "Tanenbaum Open Science Institute, The Neuro, McGill University", "Montreal, QC", [
        "Designed and administered research data management and open science training workshops for students and faculty, supporting the launch of a new academic unit",
        "Built a public-facing online knowledge base explaining open science concepts to non-specialist audiences",
    ]),
    ("07/2025 – 01/2026", "Moodle Analyst", "Premergency Inc.", "Canada", [
        "Coordinated between government clients and development teams to deliver learning-platform solutions, maintaining regulatory compliance for government deployments",
    ]),
    ("01/2024 – 10/2024", "Full Stack Developer", "KnowledgeOne", "Montreal, QC", [
        "Developed learning platform features for active clients (LAMP++, Moodle, LXD); evaluated and adopted AI tools including Claude Code in enterprise learning contexts",
    ]),
    ("2015–2016", "Software Developer, national CCNA dementia study", "McGill Centre for Integrative Neuroscience", "Montreal, QC", [
        "Built research software for the Canadian Consortium on Neurodegeneration in Aging, a pan-Canadian multi-site health research study",
    ]),
    ("04/2017 – 08/2021", "Founder and Director", "Sur Place Media", "Montreal, QC", [
        "Founded and directed a neighbourhood arts school; hired, supervised, and coordinated 20 instructors across 150 workshop groups, and led the organization's pivot to virtual delivery during the pandemic",
    ]),
]

ADDITIONAL = [
    "Kitikmeot Heritage Society — Researcher, writer, and producer (2010–2012): co-authored The Art of Inuinnait Drum Dancing (2012), published by the Kitikmeot Heritage Society",
    "Photonic Knowledge — Database administrator and software developer (2012–2014)",
]

EDUCATION = [
    ("1998–2003", "Bachelor of Arts (Honours), Cognitive Science", "McGill University", "Montreal, QC",
     ["Minor: Music Technology. Topics: psycholinguistics, cognitive science, computational modeling of language."]),
    ("2021", "Scientific Community Engagement Fundamentals", "The Center for Scientific Collaboration and Community Engagement (CSCCE)", "", []),
]

PUBLICATIONS = [
    "Poline, J.-B., et al. (2023). Data and tools integration in the Canadian open neuroscience platform. Scientific Data, 10(1), 189.",
    "Halchenko, Y. O., et al. (2021). DataLad: distributed system for joint management of code, data, and their relationship. Journal of Open Source Software, 6(63).",
    "Strauss, T. (2021). Data Management Plan Template: Neuroimaging in the Neurosciences. Zenodo. doi:10.5281/zenodo.4673558",
]
PUB_FOOT = "9 publications | citation count 630 | h-index 9 | Full list: scholar.google.com/citations"

COVER_DATE = "August 31, 2026"
COVER_SUBJECT = "Subject: Application: Research Data and Infrastructure Lead, Stephens Family School of Medicine"
COVER_SALUTATION = "Dear Hiring Committee,"
COVER_PARAS = [
    "I am applying for the Research Data and Infrastructure Lead position in the School of Medicine's "
    "Research Office. I spent five years building the data governance and infrastructure of a major "
    "academic research centre, McGill's McConnell Brain Imaging Centre, and pairing that "
    "infrastructure with the training and mentoring that made researchers actually use it. Building "
    "a research data foundation for a school that is just opening its doors is the version of that "
    "work I most want to do next.",

    "The Stephens Family School of Medicine welcomed its first cohort in August 2026 with a mandate "
    "built around primary care and community, and it needs a research data foundation built to grow "
    "with the portfolio. That founding-stage condition is where my record is strongest. At the Brain Imaging "
    "Centre I developed the governance framework, metadata standards, and data sharing policies for "
    "sensitive human research data, launched an OwnCloud self-serve portal that replaced manual "
    "data-request handling, and built pipelines covering the lifecycle from acquisition through "
    "archiving. Researchers followed the standards because training and adoption were designed in "
    "from the start, through workshops I ran at McGill's Tanenbaum Open Science Institute and a "
    "Moodle-based onboarding portal I built.",

    "On the analytics side, I work in Python and R (pandas, plotly), build Tableau dashboards, and "
    "have used SQL throughout a decade of LAMP-stack development. My health research data "
    "experience includes software development for the Canadian Consortium on Neurodegeneration in "
    "Aging, a pan-Canadian multi-site dementia study. For cross-institutional work, I represented "
    "my centre on the Portage Network, whose operations were integrated into the Digital Research "
    "Alliance of Canada in 2021, and on CONP and RDA committees, and I co-authored the Canadian "
    "Open Neuroscience Platform's data and tools integration paper in Scientific Data (2023).",

    "Two honest notes. I have not held formal OCAP training; my respect for Indigenous data "
    "contexts comes from two years co-authoring The Art of Inuinnait Drum Dancing with the "
    "Kitikmeot Heritage Society in Nunavut, and I would pursue OCAP certification early in the "
    "role. And my degree is an Honours B.A. rather than a Master's; I "
    "am applying under the equivalent-combination clause with over a decade of research data "
    "practice, a publication record, and national committee work. I am based in Montreal and "
    "prepared to relocate to Metro Vancouver, and the posted range fits what I am looking for.",

    "I would welcome the opportunity to discuss how I can help the School build a research data "
    "foundation its researchers trust. Thank you for your consideration.",
]
COVER_CLOSING = "Kind regards,"

# ==================================================================
# DOCX RENDERING
# ==================================================================
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BLUE = RGBColor(0x21, 0x6F, 0xAA)
BLACK = RGBColor(0x00, 0x00, 0x00)
GRAY = RGBColor(0x55, 0x55, 0x55)


def _margins(doc, top=1.5, bottom=1.5, left=2.0, right=2.0):
    for s in doc.sections:
        s.top_margin, s.bottom_margin = Cm(top), Cm(bottom)
        s.left_margin, s.right_margin = Cm(left), Cm(right)


def _name_header(doc, name, contact):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(name); r.bold = True; r.font.size = Pt(28); r.font.color.rgb = BLUE
    p2 = doc.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(contact); r2.font.size = Pt(10); r2.font.color.rgb = GRAY
    p2.paragraph_format.space_after = Pt(6)


def _heading(doc, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title.upper()); r.bold = True; r.font.size = Pt(11); r.font.color.rgb = BLUE
    pPr = p._p.get_or_add_pPr(); pBdr = OxmlElement('w:pBdr'); bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single'); bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1'); bottom.set(qn('w:color'), '216FAA')
    pBdr.append(bottom); pPr.append(pBdr)


def _entry(doc, dates, title, org, location, bullets):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    r1 = p.add_run(title); r1.bold = True; r1.font.size = Pt(10.5); r1.font.color.rgb = BLACK
    r2 = p.add_run(f"  —  {org}"); r2.font.size = Pt(10.5); r2.font.color.rgb = GRAY
    if location:
        r3 = p.add_run(f", {location}"); r3.font.size = Pt(10.5); r3.font.color.rgb = GRAY
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_before = Pt(0); p2.paragraph_format.space_after = Pt(2)
    p2.paragraph_format.keep_with_next = True
    r4 = p2.add_run(dates); r4.italic = True; r4.font.size = Pt(9.5); r4.font.color.rgb = GRAY
    for b in bullets:
        _bullet(doc, b)


def _bullet(doc, text, indent=0.25):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(1); p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.left_indent = Inches(indent)
    for i, part in enumerate(text.split('**')):
        if not part:
            continue
        r = p.add_run(part); r.font.size = Pt(10); r.bold = (i % 2 == 1)


def _para(doc, text, size=10, italic=False, color=None, space_after=4):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(text); r.font.size = Pt(size); r.italic = italic
    if color:
        r.font.color.rgb = color
    return p


def build_docx_cv():
    doc = Document(); _margins(doc)
    _name_header(doc, NAME, CV_CONTACT)
    _para(doc, PROFILE, size=10, space_after=6)
    _heading(doc, "Core Competencies")
    for label, rest in COMPETENCIES:
        _bullet(doc, f"**{label}**{rest}")
    _heading(doc, "Relevant Experience")
    for e in EXPERIENCE:
        _entry(doc, *e)
    _heading(doc, "Additional Experience")
    for a in ADDITIONAL:
        _bullet(doc, a)
    _heading(doc, "Education")
    for e in EDUCATION:
        _entry(doc, *e)
    _heading(doc, "Languages")
    _para(doc, "English (native), French (bilingual, professional working proficiency).", size=10)
    _heading(doc, "Selected Publications")
    for pub in PUBLICATIONS:
        _bullet(doc, pub)
    _para(doc, PUB_FOOT, size=9, italic=True, color=GRAY)
    doc.save("docx/cv_sfu_rdilead.docx")
    print("OK cv_sfu_rdilead.docx")


def build_docx_cover():
    doc = Document(); _margins(doc, top=2.0, bottom=2.0, left=2.5, right=2.5)
    _name_header(doc, NAME, COVER_CONTACT)
    _para(doc, COVER_DATE, size=10, space_after=12)
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(6)
    r = p.add_run(COVER_SUBJECT); r.bold = True; r.font.size = Pt(10)
    _para(doc, COVER_SALUTATION, size=10, space_after=6)
    for para in COVER_PARAS:
        pp = doc.add_paragraph(); pp.paragraph_format.space_after = Pt(8)
        rr = pp.add_run(para); rr.font.size = Pt(10)
    _para(doc, COVER_CLOSING, size=10, space_after=24)
    _para(doc, NAME, size=10)
    doc.save("docx/cover_sfu_rdilead.docx")
    print("OK cover_sfu_rdilead.docx")


# ==================================================================
# PDF RENDERING
# ==================================================================
from reportlab.lib.pagesizes import letter
from reportlab.lib.units import cm
from reportlab.lib.colors import HexColor
from reportlab.lib.enums import TA_CENTER
from reportlab.platypus import (
    BaseDocTemplate, PageTemplate, Frame, Paragraph, Spacer,
    HRFlowable, ListFlowable, ListItem, KeepTogether,
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

PDF_BLUE = HexColor("#216FAA")
PDF_GRAY = HexColor("#555555")


def _esc(t):
    return t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _bold_md(t):
    t = _esc(t)
    parts = t.split("**")
    out = ""
    for i, p in enumerate(parts):
        out += f"<b>{p}</b>" if i % 2 == 1 else p
    return out


def _styles():
    styles = getSampleStyleSheet()
    return {
        "name": ParagraphStyle("Name", parent=styles["Normal"], fontName="Helvetica-Bold",
                               fontSize=22, textColor=PDF_BLUE, alignment=TA_CENTER, spaceAfter=2, leading=26),
        "contact": ParagraphStyle("Contact", parent=styles["Normal"], fontName="Helvetica",
                                  fontSize=9, textColor=PDF_GRAY, alignment=TA_CENTER, spaceAfter=8, leading=12),
        "heading": ParagraphStyle("Heading", parent=styles["Normal"], fontName="Helvetica-Bold",
                                  fontSize=11, textColor=PDF_BLUE, spaceBefore=8, spaceAfter=2, leading=13),
        "body": ParagraphStyle("Body", parent=styles["Normal"], fontName="Helvetica", fontSize=10,
                               leading=13, spaceAfter=3),
        "bullet": ParagraphStyle("Bullet", parent=styles["Normal"], fontName="Helvetica",
                                 fontSize=9.5, leading=12, spaceAfter=1),
        "entry_title": ParagraphStyle("EntryTitle", parent=styles["Normal"], fontName="Helvetica",
                                      fontSize=10.5, spaceBefore=5, spaceAfter=0, leading=12),
        "entry_dates": ParagraphStyle("EntryDates", parent=styles["Normal"], fontName="Helvetica",
                                      fontSize=9, textColor=PDF_GRAY, spaceAfter=2, leading=11),
        "letter": ParagraphStyle("Letter", parent=styles["Normal"], fontName="Helvetica",
                                 fontSize=10.5, leading=15, spaceAfter=8),
        "small_italic": ParagraphStyle("SmallItalic", parent=styles["Normal"],
                                       fontName="Helvetica-Oblique", fontSize=8.5,
                                       textColor=PDF_GRAY, spaceAfter=2, leading=11),
    }


def _render(path, story, top=1.5, bottom=1.3, left=2.0, right=2.0):
    doc = BaseDocTemplate(path, pagesize=letter,
                          topMargin=top * cm, bottomMargin=bottom * cm,
                          leftMargin=left * cm, rightMargin=right * cm)
    frame = Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id="main")
    doc.addPageTemplates([PageTemplate(id="all", frames=[frame])])
    doc.build(story)
    print(f"OK {os.path.basename(path)}")


def build_pdf_cv():
    S = _styles()

    def rule():
        return HRFlowable(width="100%", thickness=0.6, color=PDF_BLUE, spaceBefore=1, spaceAfter=3)

    def heading(text):
        return [Paragraph(_esc(text.upper()), S["heading"]), rule()]

    def bullets(items):
        return ListFlowable(
            [ListItem(Paragraph(_bold_md(t), S["bullet"]), leftIndent=10, value="•") for t in items],
            bulletType="bullet", bulletColor=PDF_BLUE, leftIndent=12, bulletFontSize=8,
        )

    def entry(dates, title, org, location, bl):
        loc = f", {_esc(location)}" if location else ""
        head = (f'<b>{_esc(title)}</b>'
                f'<font color="#555555">&nbsp;&nbsp;—&nbsp;&nbsp;{_esc(org)}{loc}</font>')
        header = [Paragraph(head, S["entry_title"]), Paragraph(_esc(dates), S["entry_dates"])]
        if not bl:
            return [KeepTogether(header)]
        out = [KeepTogether(header + [bullets(bl[:1])])]
        if len(bl) > 1:
            out.append(bullets(bl[1:]))
        return out

    story = [Paragraph(_esc(NAME), S["name"]),
             Paragraph(_esc(CV_CONTACT), S["contact"]),
             Paragraph(_esc(PROFILE), S["body"])]
    story += heading("Core Competencies")
    story.append(bullets([f"**{l}**{r}" for l, r in COMPETENCIES]))
    story += heading("Relevant Experience")
    for e in EXPERIENCE:
        story += entry(*e)
    story += heading("Additional Experience")
    story.append(bullets(ADDITIONAL))
    story += heading("Education")
    for e in EDUCATION:
        story += entry(*e)
    story += heading("Languages")
    story.append(Paragraph("English (native), French (bilingual, professional working proficiency).", S["body"]))
    story += heading("Selected Publications")
    story.append(bullets(PUBLICATIONS))
    story.append(Paragraph(_esc(PUB_FOOT), S["small_italic"]))

    _render("docx/cv_sfu_rdilead.pdf", story)


def build_pdf_cover():
    S = _styles()
    story = [Paragraph(_esc(NAME), S["name"]),
             Paragraph(_esc(COVER_CONTACT), S["contact"]),
             Spacer(1, 6),
             Paragraph(_esc(COVER_DATE), S["body"]),
             Spacer(1, 8),
             Paragraph(f"<b>{_esc(COVER_SUBJECT)}</b>", S["body"]),
             Spacer(1, 4),
             Paragraph(_esc(COVER_SALUTATION), S["letter"])]
    for para in COVER_PARAS:
        story.append(Paragraph(_esc(para), S["letter"]))
    story += [Spacer(1, 6),
              KeepTogether([Paragraph(_esc(COVER_CLOSING), S["body"]),
                            Spacer(1, 12),
                            Paragraph(_esc(NAME), S["body"])])]
    _render("docx/cover_sfu_rdilead.pdf", story, top=1.6, bottom=1.2, left=2.5, right=2.5)


if __name__ == "__main__":
    build_docx_cv()
    build_docx_cover()
    build_pdf_cv()
    build_pdf_cover()
    print("\nDone. Files in docx/")
